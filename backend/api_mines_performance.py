"""
"Mines Production Performance" -- standalone report (mines_performance.py),
pages 5-6 layout with mines in place of plants. Not part of the monthly PDF.

  GET /api/mines-performance?month=YYYY-MM                 -> JSON
  GET /api/mines-performance/xlsx?month=YYYY-MM&decimals=N -> .xlsx
  GET /api/mines-performance/pdf?month=YYYY-MM&decimals=N  -> .pdf
"""
import io
import re

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response, StreamingResponse

import mines_performance
import mines_performance_pdf
from colors_loader import load_colors_config

router = APIRouter(prefix="/api/mines-performance", tags=["mines-performance"])


def _check_month(month: str):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month or ""):
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")


def _check_decimals(decimals: int):
    if not 0 <= decimals <= 3:
        raise HTTPException(status_code=400, detail="decimals must be 0, 1, 2 or 3")


def _data(month: str) -> dict:
    return mines_performance.generate(month)


@router.get("")
def mines_performance_json(month: str = Query(...)):
    _check_month(month)
    return {"month": month, **_data(month)}


def _qty_fmt(decimals: int) -> str:
    return "#,##0" if decimals == 0 else "#,##0." + "0" * decimals


def write_sheet(ws, data: dict, decimals: int):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    L, pct = data["labels"], set(data["pct_idx"])
    ncols = 13
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="DBEAFE")
    total_fill = PatternFill("solid", fgColor="FEF3C7")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(1, 1, L["title"]).font = Font(bold=True, size=13, color="1E3A8A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.cell(2, ncols, "Unit: '000 T").alignment = Alignment(horizontal="right")

    h1, h2 = 3, 4
    top = [("", 1, 1), ("Item", 2, 2), (f"{L['fy']}\nPlan", 3, 3), (L["month"], 4, 6),
           (f"{L['cply']}\nAct.", 7, 7), (f"%Gr.\n{L['cply']}", 8, 8), (L["ytd"], 9, 11),
           (f"{L['ytd_cply']}\nAct.", 12, 12), (f"%Gr.\n{L['ytd_cply']}", 13, 13)]
    for text, c1, c2 in top:
        ws.cell(h1, c1, text)
        if c1 != c2:
            ws.merge_cells(start_row=h1, start_column=c1, end_row=h1, end_column=c2)
        else:
            ws.merge_cells(start_row=h1, start_column=c1, end_row=h2, end_column=c1)
    for i, t in enumerate(["Plan", "Actual", "%Ful"]):
        ws.cell(h2, 4 + i, t)
        ws.cell(h2, 9 + i, t)
    for r in (h1, h2):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.alignment, cell.border = bold, hdr_fill, center, border

    qty = _qty_fmt(decimals)
    indent = {"heading": 0, "subheading": 1, "total": 1, "item": 2}
    r = h2 + 1
    for b in data["blocks"]:
        start = r
        for row in b["rows"]:
            kind = row["kind"]
            lbl = ws.cell(r, 2, row["label"])
            lbl.alignment = Alignment(indent=indent[kind])
            for vi, v in enumerate(row["values"]):
                c = ws.cell(r, 3 + vi, None if v is None else round(v, 3))
                if vi in pct:
                    c.number_format = "0"
                elif row["decimals"] is not None:
                    c.number_format = "0." + "0" * row["decimals"]
                else:
                    c.number_format = qty
            for col in range(1, ncols + 1):
                cell = ws.cell(r, col)
                cell.border = border
                if col >= 2 and kind in ("heading", "total"):
                    cell.font = bold
                if col >= 2 and kind == "total":
                    cell.fill = total_fill
            if kind == "subheading":
                lbl.font = Font(bold=True, italic=True)
            r += 1
        ws.cell(start, 1, b["name"])
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(start, 1).font = bold
        ws.cell(start, 1).alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")

    for c, w in enumerate([6, 32] + [10] * 11, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[h1].height = 30
    ws.freeze_panes = ws.cell(h2 + 1, 3)


@router.get("/xlsx")
def mines_performance_xlsx(month: str = Query(...), decimals: int = Query(0)):
    _check_month(month)
    _check_decimals(decimals)
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Mines Performance"
    write_sheet(ws, _data(month), decimals)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Mines_Performance_{month}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/pdf")
def mines_performance_pdf_download(month: str = Query(...), decimals: int = Query(0)):
    # A plain def, so FastAPI runs it in its threadpool: the DB queries and
    # the synchronous Playwright render must not block the event loop.
    _check_month(month)
    _check_decimals(decimals)
    content = mines_performance_pdf.render_pdf(_data(month), decimals, load_colors_config())
    fname = f"Mines_Performance_{month}.pdf"
    return Response(content, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})
