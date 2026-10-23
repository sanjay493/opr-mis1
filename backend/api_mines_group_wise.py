"""
"Iron Ore Mines -- Mine & Group-wise (FY)" -- workbook-style reshape of the
existing mine-level Production & Despatch tables (mines_group_wise.py).

  GET /api/mines-group-wise?fy_start=YYYY         -> JSON
  GET /api/mines-group-wise/xlsx?fy_start=YYYY     -> .xlsx
"""
import io

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse

import mines_group_wise

router = APIRouter(prefix="/api/mines-group-wise", tags=["mines-group-wise"])


def _check_fy_start(fy_start: int):
    if not 2000 <= fy_start <= 2100:
        raise HTTPException(status_code=400, detail="fy_start out of range")


def _data(fy_start: int) -> dict:
    return mines_group_wise.generate(fy_start)


@router.get("")
def mines_group_wise_json(fy_start: int = Query(...)):
    _check_fy_start(fy_start)
    return _data(fy_start)


_MONTH_LABEL = {
    "04": "Apr", "05": "May", "06": "Jun", "07": "Jul", "08": "Aug", "09": "Sep",
    "10": "Oct", "11": "Nov", "12": "Dec", "01": "Jan", "02": "Feb", "03": "Mar",
}


def _month_header(month: str) -> str:
    y, m = month.split("-")
    return f"{_MONTH_LABEL[m]}'{y[2:]}"


def _row_label_values(row, row_cols_keys):
    return [row[k] for k in row_cols_keys]


def _write_sheet(ws, data: dict, section: dict, row_cols_keys, row_cols_headers, title: str):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    months = data["months"]
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="DBEAFE")
    mine_fill = PatternFill("solid", fgColor="F1F5F9")
    total_fill = PatternFill("solid", fgColor="FEF3C7")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    label_cols = len(row_cols_headers)  # 1 for Production (Commodity), 2 for Despatch (Commodity, Destination)
    first_label_col = 3  # col 1 = Group, col 2 = Mine, col 3.. = Commodity[, Destination]
    qty_start_col = first_label_col + label_cols
    ncols = qty_start_col - 1 + len(months) * 2 + 2  # +FY Total(Plan,Actual)

    ws.cell(1, 1, f"{title} -- FY {data['fy_label']}").font = Font(bold=True, size=13, color="1E3A8A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(2, ncols, "Unit: '000 T").alignment = Alignment(horizontal="right")

    h1, h2 = 3, 4
    ws.cell(h1, 1, "Group")
    ws.merge_cells(start_row=h1, start_column=1, end_row=h2, end_column=1)
    ws.cell(h1, 2, "Mine")
    ws.merge_cells(start_row=h1, start_column=2, end_row=h2, end_column=2)
    for i, label in enumerate(row_cols_headers):
        c = 3 + i
        ws.cell(h1, c, label)
        ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
    col = qty_start_col
    for month in months:
        ws.cell(h1, col, _month_header(month))
        ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=col + 1)
        ws.cell(h2, col, "Plan")
        ws.cell(h2, col + 1, "Actual")
        col += 2
    ws.cell(h1, col, "FY Total")
    ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=col + 1)
    ws.cell(h2, col, "Plan")
    ws.cell(h2, col + 1, "Actual")

    for r in (h1, h2):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.alignment, cell.border = bold, hdr_fill, center, border

    def write_values(r, months_cells, fy_total, fill=None):
        c = qty_start_col
        for month in months:
            cell = months_cells[month]
            ws.cell(r, c, None if cell["plan"] is None else round(cell["plan"], 3))
            ws.cell(r, c + 1, None if cell["actual"] is None else round(cell["actual"], 3))
            c += 2
        ws.cell(r, c, None if fy_total["plan"] is None else round(fy_total["plan"], 3))
        ws.cell(r, c + 1, None if fy_total["actual"] is None else round(fy_total["actual"], 3))
        for col_ in range(1, ncols + 1):
            cell = ws.cell(r, col_)
            cell.border = border
            cell.number_format = "#,##0.0"
            if fill:
                cell.fill = fill

    r = h2 + 1
    for g in section["groups"]:
        group_start = r
        for mine in g["mines"]:
            mine_start = r
            for row in mine["rows"]:
                for i, val in enumerate(_row_label_values(row, row_cols_keys)):
                    ws.cell(r, first_label_col + i, val)
                write_values(r, row["months"], row["fy_total"])
                r += 1
            if r > mine_start:
                ws.cell(mine_start, 2, mine["mine_name"])
                if r - 1 > mine_start:
                    ws.merge_cells(start_row=mine_start, start_column=2, end_row=r - 1, end_column=2)
                ws.cell(mine_start, 2).font = bold
                ws.cell(mine_start, 2).alignment = Alignment(vertical="center")
                for rr in range(mine_start, r):
                    ws.cell(rr, 2).fill = mine_fill
            ws.cell(r, 2, f"{mine['mine_name']} Total")
            ws.cell(r, 2).font = bold
            write_values(r, mine["total"]["months"], mine["total"]["fy_total"], fill=mine_fill)
            r += 1
        ws.cell(group_start, 1, g["group_name"])
        if r - 1 > group_start:
            ws.merge_cells(start_row=group_start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(group_start, 1).font = bold
        ws.cell(group_start, 1).alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        ws.cell(r, 2, f"{g['group_name']} Total")
        ws.cell(r, 2).font = bold
        write_values(r, g["total"]["months"], g["total"]["fy_total"], fill=total_fill)
        r += 1

    ws.cell(r, 2, "SAIL TOTAL")
    ws.cell(r, 2).font = bold
    write_values(r, section["sail_total"]["months"], section["sail_total"]["fy_total"], fill=total_fill)

    widths = [5, 16] + [14] * label_cols + [9] * (len(months) * 2) + [10, 10]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[h1].height = 26
    ws.freeze_panes = ws.cell(h2 + 1, qty_start_col)


@router.get("/xlsx")
def mines_group_wise_xlsx(fy_start: int = Query(...)):
    _check_fy_start(fy_start)
    from openpyxl import Workbook

    data = _data(fy_start)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Production"
    _write_sheet(ws1, data, data["production"], ["material_name"], ["Commodity"], "Iron Ore Mines Production")

    ws2 = wb.create_sheet("Despatch")
    _write_sheet(ws2, data, data["despatch"], ["material_name", "end_use_name"],
                 ["Commodity", "Destination"], "Iron Ore Mines Despatch")

    ws3 = wb.create_sheet("Reported Production")
    _write_sheet(ws3, data, data["reported_production"], ["material_name"],
                 ["Component"], "Iron Ore Mines Reported Production")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"IronOreMines_GroupWise_FY{data['fy_label']}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
