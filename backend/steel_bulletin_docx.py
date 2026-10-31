"""
"Inputs for Steel Bulletin" (3.1 Performance) as a Word document.

The figures are computed by the frontend page (reports/steel-bulletin/View.js)
from /api/production-fy; it posts the already-formatted table here so the
.docx matches what is on screen cell for cell. No DB access.
"""

import io


def build_steel_bulletin_docx_bytes(payload: dict) -> bytes:
    """payload: {title, heading, intro, unit, header: [str], rows: [[str]]}."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    header = [str(h) for h in payload.get("header") or []]
    rows = [[("" if v is None else str(v)) for v in r] for r in payload.get("rows") or []]

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(15))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def para(text, bold=False, size=None, align=None, space_after=6):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    if payload.get("title"):
        para(payload["title"], bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    if payload.get("heading"):
        para(payload["heading"], bold=True)
    if payload.get("intro"):
        para(payload["intro"])
    if payload.get("unit"):
        para(payload["unit"], size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)

    def shade(cell, hex_fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def fill(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.RIGHT):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        lines = text.split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            run.bold = bold
            run.font.size = Pt(10)
            if color:
                run.font.color.rgb = color
            if i < len(lines) - 1:
                run.add_break()
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    white = RGBColor(0xFF, 0xFF, 0xFF)
    for c, h in enumerate(header):
        cell = table.rows[0].cells[c]
        fill(cell, h, bold=True, color=white,
             align=WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, "1E3A5F")

    for r, row in enumerate(rows, start=1):
        for c in range(len(header)):
            text = row[c] if c < len(row) else ""
            fill(table.rows[r].cells[c], text, bold=(c == 0),
                 align=WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
