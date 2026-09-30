"""
"Plant Wise Performance of Main Items" — the PDF report's page 4 as an
External Report: same rows/columns (page4.generate_page4_rows), but with the
unrounded figures (raw=True) so the Excel download carries actual tonnage to
3 decimal places instead of the PDF's whole numbers.

  GET /api/plant-performance-main-items?month=YYYY-MM        -> JSON rows
  GET /api/plant-performance-main-items/xlsx?month=YYYY-MM   -> .xlsx download
"""
import calendar
import io
import re

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse

import db
from page4 import generate_page4_rows

router = APIRouter(prefix="/api/plant-performance-main-items", tags=["plant-performance"])

# Index into a row's 15 values that are percentages (shown as whole numbers);
# every other value is a quantity ('000 T, or nos/day for Oven Pushing).
PCT_IDX = {4, 6, 7, 11, 13, 14}

_MON = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _labels(month: str) -> dict:
    y, m = int(month[:4]), int(month[5:7])
    fy = y if m >= 4 else y - 1
    sy, py = f"{y % 100:02d}", f"{(y - 1) % 100:02d}"
    return {
        "title": f"Plant Wise Performance of Main Items during {calendar.month_name[m]}'{sy} "
                 f"and Apr-{_MON[m]}'{sy}",
        "fy": f"{fy % 100:02d}-{(fy + 1) % 100:02d}",
        "month": f"{_MON[m]}'{sy}",
        "cply": f"{_MON[m]}'{py}",
        "ytd": f"Apr-{_MON[m]}'{sy}",
        "ytd_cply": f"Apr-{_MON[m]}'{py}",
    }


def _rows(month: str) -> list:
    out = []
    for r in generate_page4_rows(month, raw=True):
        display = r.get("display_name", "")
        plant = r["label"][len(display):].strip() if r["label"].startswith(display) else ""
        out.append({
            "item": display,
            "plant": plant,
            "capacity": r.get("capacity"),
            "values": r["values"],
            "is_conversion": bool(r.get("is_conversion")),
            "is_sail_incl_conv": bool(r.get("is_sail_incl_conv")),
        })
    return out


def _check_month(month: str):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month or ""):
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")


@router.get("")
def plant_performance(month: str = Query(...)):
    _check_month(month)
    db.init_db()
    return {"month": month, "labels": _labels(month), "pct_idx": sorted(PCT_IDX), "rows": _rows(month)}


@router.get("/xlsx")
def plant_performance_xlsx(month: str = Query(...)):
    _check_month(month)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    db.init_db()
    L = _labels(month)
    rows = _rows(month)

    wb = Workbook()
    ws = wb.active
    ws.title = "Plant Performance"
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="DBEAFE")
    act_fill = PatternFill("solid", fgColor="ECFDF5")
    sail_fill = PatternFill("solid", fgColor="FEF3C7")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ncols = 18

    ws.cell(1, 1, L["title"]).font = Font(bold=True, size=13, color="1E3A8A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols - 3)
    ws.cell(1, ncols - 2, "w.r.t APP").font = bold
    ws.cell(2, 1, "Tentative").font = Font(italic=True)
    ws.cell(2, ncols, "Unit: '000 T").font = Font(italic=True)
    ws.cell(2, ncols).alignment = Alignment(horizontal="right")

    h1, h2 = 3, 4
    top = [("Items", 1, 1), ("Plant", 2, 2), ("Ann.\nCap.", 3, 3), (f"APP\n{L['fy']}", 4, 4),
           (L["month"], 5, 8), (f"{L['cply']}\nAct.", 9, 9), (f"%Gr.\n{L['cply']}", 10, 10), ("CU%", 11, 11),
           (L["ytd"], 12, 15), (f"{L['ytd_cply']}\nAct.", 16, 16), (f"%Gr.\n{L['ytd_cply']}", 17, 17),
           ("CU%", 18, 18)]
    for text, c1, c2 in top:
        ws.cell(h1, c1, text)
        if c1 != c2:
            ws.merge_cells(start_row=h1, start_column=c1, end_row=h1, end_column=c2)
        else:
            ws.merge_cells(start_row=h1, start_column=c1, end_row=h2, end_column=c1)
    for i, t in enumerate(["APP", "Act.", "Var", "%Ful."]):
        ws.cell(h2, 5 + i, t)
        ws.cell(h2, 12 + i, t)
    for r in (h1, h2):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.alignment, cell.border = bold, hdr_fill, center, border

    qty_fmt, pct_fmt = "#,##0.000", "0"
    r = h2 + 1
    i = 0
    while i < len(rows):
        # item group: consecutive rows with the same item name (not the two
        # conversion rows, which span the Items/Plant columns themselves)
        j = i
        while (j + 1 < len(rows) and rows[j + 1]["item"] == rows[i]["item"]
               and not rows[i]["is_conversion"] and not rows[i]["is_sail_incl_conv"]):
            j += 1
        for k in range(i, j + 1):
            row = rows[k]
            rr = r + (k - i)
            if row["is_conversion"]:
                ws.cell(rr, 1, "Conversion")
                ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
            elif row["is_sail_incl_conv"]:
                ws.cell(rr, 1, "SAIL incl. conversion")
                ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
            else:
                if k == i:
                    ws.cell(rr, 1, row["item"])
                ws.cell(rr, 2, row["plant"])
            if not row["is_conversion"] and row["capacity"] is not None:
                ws.cell(rr, 3, row["capacity"]).number_format = qty_fmt
            for vi, v in enumerate(row["values"]):
                c = ws.cell(rr, 4 + vi, None if v is None else round(v, 3))
                c.number_format = pct_fmt if vi in PCT_IDX else qty_fmt
            is_sail = row["plant"] in ("SAIL", "5 Plants") or row["is_sail_incl_conv"] or row["is_conversion"]
            for c in range(1, ncols + 1):
                cell = ws.cell(rr, c)
                cell.border = border
                if c >= 2 and is_sail:
                    cell.font = bold
                    cell.fill = sail_fill
                elif c in (6, 13):
                    cell.fill = act_fill
            ws.cell(rr, 2).alignment = Alignment(horizontal="center")
        if j > i:
            ws.merge_cells(start_row=r, start_column=1, end_row=r + (j - i), end_column=1)
        ws.cell(r, 1).font = bold
        ws.cell(r, 1).alignment = Alignment(vertical="center", wrap_text=True)
        r += j - i + 1
        i = j + 1

    widths = [16, 9, 10, 11] + [11, 11, 10, 7] + [11, 7, 7] + [11, 11, 10, 7] + [11, 7, 7]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[h1].height = 30
    ws.freeze_panes = ws.cell(h2 + 1, 3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Plant_Performance_Main_Items_{month}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
