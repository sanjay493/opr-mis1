import io
import re
from pathlib import Path

import pytest
from pptx import Presentation

import page_secretary_review as psr
import page_secretary_review_texts as pst
import secretary_review_layout as L
import secretary_review_pptx as sp
from test_secretary_review_context import live

REF = Path(__file__).parent / "fixtures" / "secretary_review_sep26_reference.pptx"

# (slide, row, col) -> reason, for reference-deck cells the DB no longer reproduces.
# Every entry needs a reason checked by hand.
KNOWN_TABLE_DIFFS = {}


def _num(s):
    s = s.replace(",", "").strip()
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def _perf_tables(slide):
    return [sh.table for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table
            and len(sh.table.rows) > 1 and len(sh.table.columns) == 10]


@pytest.fixture(scope="module")
def rendered():
    texts = {k: "" for k, _, _ in L.BLOCKS}
    texts["hl_SAIL"] = "Line one:\nLine two"
    res = psr.render_pptx("2026-09", texts)
    return res, Presentation(io.BytesIO(res.content))


@live
def test_tables_match_reference(rendered):
    _, out = rendered
    ref = Presentation(str(REF))
    diffs = []
    for idx in (2, 3, 5, 8, 10, 12, 14, 16):
        for t_out, t_ref in zip(_perf_tables(out.slides[idx - 1]), _perf_tables(ref.slides[idx - 1])):
            for r, (row_o, row_r) in enumerate(zip(t_out.rows, t_ref.rows)):
                for c, (co, cr) in enumerate(zip(row_o.cells, row_r.cells)):
                    a, b = _num(co.text), _num(cr.text)
                    if b is not None and (a is None or abs(a - b) > 1) and (idx, r, c) not in KNOWN_TABLE_DIFFS:
                        diffs.append((idx, r, c, co.text, cr.text))
    assert not diffs, diffs


@live
def test_charts_filled(rendered):
    _, out = rendered
    shapes = sp.named_shapes(out)
    ch = shapes["ch_SAIL_coke"].chart
    assert list(ch.plots[0].categories) == ["FY25", "FY26", "FY'27 Target", "Sep", "Apr-Sep"]
    assert ch.plots[0].series[0].values == (421.0, 419.0, 400.0, 432.0, 424.0)
    tr = shapes["trend_BSP_fuel"].chart
    assert list(tr.plots[0].categories)[:3] == ["FY25", "FY26", "FY'27 Target"]
    assert list(tr.plots[0].categories)[-1] == "Sep"


@live
def test_labels_and_narrative(rendered):
    res, out = rendered
    assert res.filename == "SECRETARY REVIEW Operations Inputs Sep26.pptx"
    all_text = " ".join(sh.text_frame.text for s in out.slides for sh in s.shapes if sh.has_text_frame)
    assert "{" not in all_text and "Sep’26" in all_text
    shapes = sp.named_shapes(out)
    hl = shapes["tbl_sail_hl"].table.cell(0, 1).text_frame
    assert [p.text for p in hl.paragraphs] == ["Line one:", "Line two"]
    assert "tbl_BSP_hl" not in shapes                     # empty highlights -> table removed
    bd = shapes["tbl_bd_BSP"].table
    assert len(bd.rows) == 1 and bd.cell(0, 1).text == "No major breakdowns"


def test_render_month_without_data():
    res = psr.render_pptx("2030-01", {k: "" for k, _, _ in L.BLOCKS})
    out = Presentation(io.BytesIO(res.content))
    assert len(out.slides) == 30
    assert res.filename.endswith("Jan30.pptx")


def test_missing_texts_fall_back_to_effective(monkeypatch):
    monkeypatch.setattr(pst, "effective_texts", lambda m: {k: {"text": "", "saved": False} for k, _, _ in L.BLOCKS})
    res = psr.render_pptx("2030-01", {"delay_fs": "MILL: x – 3 days"})
    out = Presentation(io.BytesIO(res.content))
    assert sp.named_shapes(out)["tbl_delay_fs"].table.cell(0, 1).text == "MILL: x – 3 days"


def _cr_run_sizes(out):
    t = sp.named_shapes(out)["tbl_cr_1"].table
    return {int(r.font._element.get("sz")) for row in t.rows if row.cells[0].text.strip().upper() in L.PLANTS
            for c in list(row.cells)[1:3]
            for p in c.text_frame.paragraphs for r in p.runs if r.font._element.get("sz")}


def _render_cr(n):
    texts = {k: "" for k, _, _ in L.BLOCKS}
    texts["cr_BSP_cur"] = "\n".join(f"Item {i}: repair work" for i in range(n))
    res = psr.render_pptx("2030-01", texts)
    return res, Presentation(io.BytesIO(res.content))


def test_long_capital_repairs_text_is_scaled():
    res, out = _render_cr(25)
    sizes = _cr_run_sizes(out)
    assert sizes and max(sizes) < 1200 and min(sizes) >= 800
    if min(sizes) == 800:
        assert any("Capital repairs (slide 19)" in w for w in res.warnings)


def test_overlong_capital_repairs_text_warns():
    res, out = _render_cr(45)
    assert _cr_run_sizes(out) == {800}
    assert any("Capital repairs (slide 19)" in w and "may overflow" in w for w in res.warnings)


def test_short_capital_repairs_text_keeps_template_size():
    res, out = _render_cr(3)
    assert _cr_run_sizes(out) == {1200}
    assert not any("Capital repairs" in w for w in res.warnings)
