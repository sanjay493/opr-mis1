import docx

import board_note_docx_utils as u


def _doc_with_paragraphs(*texts):
    d = docx.Document()
    paras = [d.add_paragraph(t) for t in texts]
    return d, paras


def test_set_para_keeps_first_run_formatting_and_clears_rest():
    d, [p] = _doc_with_paragraphs("old")
    p.add_run(" extra")
    u.set_para(p, "new")
    assert p.text == "new"


def test_replace_in_para_substring_swap():
    d, [p] = _doc_with_paragraphs("Q-2'25-26:")
    u.replace_in_para(p, [("Q-2'25-26", "Q-2'26-27")])
    assert p.text == "Q-2'26-27:"


def test_set_cell_clears_extra_paragraphs():
    d = docx.Document()
    t = d.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    cell.add_paragraph("second line")
    u.set_cell(cell, "only line")
    assert [p.text for p in cell.paragraphs] == ["only line"]


def test_fill_variable_bullets_exact_fit():
    d, slots = _doc_with_paragraphs("a", "b", "c")
    u.fill_variable_bullets(slots, ["x", "y", "z"])
    assert [p.text for p in d.paragraphs] == ["x", "y", "z"]


def test_fill_variable_bullets_fewer_lines_removes_unused_slots():
    d, slots = _doc_with_paragraphs("a", "b", "c")
    u.fill_variable_bullets(slots, ["x"])
    assert [p.text for p in d.paragraphs] == ["x"]


def test_fill_variable_bullets_more_lines_clones_extra_paragraphs():
    d, slots = _doc_with_paragraphs("a")
    u.fill_variable_bullets(slots, ["x", "y", "z"])
    assert [p.text for p in d.paragraphs] == ["x", "y", "z"]


def test_fill_variable_bullets_empty_lines_removes_all_slots():
    d, slots = _doc_with_paragraphs("a", "b")
    u.fill_variable_bullets(slots, [])
    assert [p.text for p in d.paragraphs] == []


def test_clear_strike_and_colour_keeps_bold():
    from docx.shared import RGBColor
    d = docx.Document()
    cell = d.add_table(rows=1, cols=1).rows[0].cells[0]
    r = cell.paragraphs[0].add_run("5377")
    r.font.strike, r.font.bold, r.font.color.rgb = True, True, RGBColor(0xFF, 0, 0)
    u.clear_strike_and_colour(cell)
    r = cell.paragraphs[0].runs[0]
    assert r.font.strike is None and r.font.color.rgb is None and r.font.bold is True


def _q2_template():
    return docx.Document("board_note_templates/board_note_q2_template.docx")


def test_fs_charts_found_by_title():
    charts = u.fs_charts(_q2_template())
    assert sorted(charts) == ["long", "quarter"]


def _chart_vals(part):
    import re
    x = part.blob.decode("utf-8")
    title = re.search(r"<a:t>([^<]*)</a:t>", x).group(1)
    vals = {}
    for ser in re.findall(r"<c:ser>.*?</c:ser>", x, flags=re.S):
        name = re.search(r"<c:tx>.*?<c:v>([^<]*)</c:v>", ser, flags=re.S).group(1)
        cache = re.search(r"<c:val>.*?</c:numCache>", ser, flags=re.S).group(0)
        v = re.search(r"<c:v>([^<]*)</c:v>", cache)
        vals[name] = None if v is None else float(v.group(1))
    return title, vals


def test_fill_fs_chart_sets_title_series_and_workbook():
    import io
    import openpyxl
    d = _q2_template()
    part = u.fs_charts(d)["quarter"]
    u.fill_fs_chart(part, "Q-2'26-27", mou=4858.5, act=4360.1, cply=4233.1)
    title, vals = _chart_vals(part)
    assert title == "Q-2'26-27"
    assert vals == {"MoU": 4858.5, "Actual": 4360.1, "%ful. MoU": 90, "CPLY": 4233.1, "%Growth": 3.0}
    wb_part = next(r.target_part for r in part.rels.values() if r.reltype.endswith("/package"))
    ws = openpyxl.load_workbook(io.BytesIO(wb_part.blob)).active
    assert [ws.cell(2, c).value for c in range(2, 7)] == [4858.5, 4360.1, 90, 4233.1, 3.0]


def test_fill_fs_chart_blank_mou_and_raised_axis():
    import re
    d = _q2_template()
    part = u.fs_charts(d)["quarter"]
    u.fill_fs_chart(part, "Q-2'26-27", mou=None, act=6500.0, cply=6100.0)
    _, vals = _chart_vals(part)
    assert vals["MoU"] is None and vals["%ful. MoU"] is None and vals["%Growth"] == 6.6
    assert re.search(r'<c:max val="([^"]+)"/>', part.blob.decode()).group(1) == "8000"
