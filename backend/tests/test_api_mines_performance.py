"""Mines performance API: validation, JSON, xlsx layout, pdf bytes."""
import io

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import api_mines_performance as app_mod

LABELS = {"title": "Mines Production Performance :Sep'26 and Apr-Sep'26", "fy": "26-27",
          "month": "Sep'26", "cply": "Sep'25", "ytd": "Apr-Sep'26", "ytd_cply": "Apr-Sep'25"}


def _row(label, kind, vals=None, decimals=None):
    return {"label": label, "kind": kind, "values": vals or [None] * 11, "decimals": decimals}


V = [1200.4, 100.0, 95.25, 95.25, 90.0, 5.8, 600.0, 570.0, 95.0, 550.0, 3.6]
DATA = {"labels": LABELS, "pct_idx": [3, 5, 8, 10], "blocks": [
    {"name": "Kiriburu", "kind": "mine", "group": "JGoM", "note": None, "rows": [
        _row("I/ORE FRESH PRODUCTION", "heading"), _row("Lump", "item", V),
        _row("Total Fresh Production", "total", V)]},
    {"name": "SAIL", "kind": "sail", "group": None, "note": "note", "rows": [
        _row("Yield (%)", "item", V, decimals=1)]},
]}


@pytest.fixture
def tc(monkeypatch):
    monkeypatch.setattr(app_mod.mines_performance, "generate", lambda month: DATA)
    monkeypatch.setattr(app_mod.db, "init_db", lambda: None)
    app = FastAPI()
    app.include_router(app_mod.router)
    return TestClient(app)


def test_bad_month_and_decimals(tc):
    assert tc.get("/api/mines-performance?month=2026-13").status_code == 400
    assert tc.get("/api/mines-performance/xlsx?month=2026-09&decimals=5").status_code == 400
    assert tc.get("/api/mines-performance/pdf?month=26-09").status_code == 400


def test_json(tc):
    d = tc.get("/api/mines-performance?month=2026-09").json()
    assert d["month"] == "2026-09" and len(d["blocks"]) == 2 and d["pct_idx"] == [3, 5, 8, 10]


def test_xlsx_layout(tc):
    r = tc.get("/api/mines-performance/xlsx?month=2026-09&decimals=2")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Mines Performance"
    assert ws.cell(1, 1).value == LABELS["title"]
    assert ws.cell(3, 4).value == "Sep'26" and ws.cell(4, 4).value == "Plan" and ws.cell(4, 5).value == "Actual"
    assert "A5:A7" in [str(m) for m in ws.merged_cells.ranges]        # Kiriburu block, 3 rows
    assert ws.cell(5, 1).value == "Kiriburu" and ws.cell(5, 2).value == "I/ORE FRESH PRODUCTION"
    assert ws.cell(6, 3).value == 1200.4 and ws.cell(6, 3).number_format == "#,##0.00"
    assert ws.cell(6, 6).number_format == "0"                        # %Ful
    assert ws.cell(8, 4).number_format == "0.0"                      # Yield quantity, 1 decimal
    assert ws.cell(8, 6).number_format == "0"                        # Yield %Ful stays whole
    assert ws.cell(7, 2).font.bold


def test_pdf(tc, monkeypatch):
    seen = {}
    monkeypatch.setattr(app_mod.mines_performance_pdf, "render",
                        lambda html: seen.setdefault("html", html) and b"%PDF-fake")
    r = tc.get("/api/mines-performance/pdf?month=2026-09&decimals=1")
    assert r.status_code == 200 and r.content == b"%PDF-fake"
    assert r.headers["content-type"] == "application/pdf"
    assert "Kiriburu" in seen["html"]
