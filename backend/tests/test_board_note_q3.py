"""Q-3 + 9M Board Note adapter: template index map and title text (no DB needed)."""

import docx

import board_note_q3 as bq3


def _doc():
    return docx.Document(bq3.TEMPLATE_PATH)


def test_title_text_names_the_quarter_and_its_months():
    assert bq3._title_text(2026) == (
        "PRODUCTION PERFORMANCE DURING QUARTER-3’26-27 "
        "(OCTOBER’26, NOVEMBER’26, DECEMBER’26 and APRIL-DECEMBER’26) ")


def test_index_map_points_at_the_headings_it_claims():
    paras = _doc().paragraphs
    assert paras[bq3.P_SAIL_Q3_HEADING].text.startswith("SAIL: Production performance during Q-3")
    assert paras[bq3.P_SAIL_9M_HEADING].text.startswith("SAIL: Production performance during 9M")
    assert paras[bq3.P_PLANT_SECTION_TITLE].text.startswith("Plant wise: Production performance during Q-3")
    assert paras[bq3.P_TECHNO_SECTION_TITLE].text.startswith("SAIL and Plant-wise Major Techno-economic")
    assert paras[bq3.P_MOU_HEADING].text.startswith("SAIL: Production performance for Finished Steel")
    for plant, idx in bq3.WHY_HEADING.items():
        assert paras[idx].text.startswith("Production Performance during Q-3"), plant
    for key, idx in bq3.HIGHLIGHTS_HEADING.items():
        assert paras[idx].text == "Highlights:", key
    for key, idx in bq3.HIGHLIGHTS_LEADIN.items():
        if idx is not None:
            assert paras[idx].text.endswith("for following:"), key
    for plant, cfg in bq3.PLANT_CONFIG.items():
        assert paras[cfg["label_q3"]].text.startswith("Q-3"), plant
        assert paras[cfg["label_9m"]].text.startswith("9M"), plant


def test_every_slot_is_a_non_empty_template_paragraph():
    paras = _doc().paragraphs
    for key, idxs in bq3.HIGHLIGHTS_SLOTS.items():
        for i in idxs:
            assert paras[i].text.strip(), (key, i)
    for plant, idxs in bq3.WHY_SLOTS.items():
        for i in idxs:
            assert paras[i].text.strip(), (plant, i)


def test_highlight_gaps_are_only_paragraphs_the_adapter_removes():
    paras = _doc().paragraphs
    for key, idxs in bq3.HIGHLIGHTS_SLOTS.items():
        removable = set(bq3.HIGHLIGHTS_BLANKS_REMOVE.get(key, []))
        if key in bq3.SECOND_LEADIN_REMOVE:
            removable.add(bq3.SECOND_LEADIN_REMOVE[key])
        gaps = [i for i in range(idxs[0], idxs[-1] + 1) if i not in idxs]
        assert all(i in removable for i in gaps), (key, gaps)
    for key, idx in bq3.SECOND_LEADIN_REMOVE.items():
        assert paras[idx].text.endswith("production for following:"), key
    for key, idxs in bq3.HIGHLIGHTS_BLANKS_REMOVE.items():
        for i in idxs:
            assert paras[i].text == "", (key, i)


def test_unit_tables_match_their_item_rows():
    tables = _doc().tables
    for plant, cfg in bq3.PLANT_CONFIG.items():
        for table_idx in (cfg["table_q3"], cfg["table_9m"]):
            table = tables[table_idx]
            assert len(table.rows) == 2 + len(cfg["item_rows"]), (plant, table_idx)
            for row_idx, display, _db_item, _add_conv in cfg["item_rows"]:
                assert table.rows[row_idx].cells[0].text.strip() == display, (plant, row_idx)
    for table_idx in (bq3.TABLE_SAIL_Q3, bq3.TABLE_SAIL_9M):
        sail = tables[table_idx]
        assert [r.cells[0].text.strip() for r in sail.rows[2:]] == [
            "Hot Metal", "Crude Steel", "Saleable Steel", "Finished Steel"]


def test_mou_techno_and_annexure_tables_have_the_expected_shape():
    tables = _doc().tables
    mou = tables[bq3.TABLE_MOU]
    assert len(mou.rows) == 9
    assert len(mou.columns) == 8
    assert len(tables[bq3.TABLE_TECHNO].rows) == 26
    assert len(tables[bq3.TABLE_ANNEXURE_INDEX].rows) == 6


def test_techno_parameter_blocks_run_in_the_order_the_adapter_expects():
    t = _doc().tables[bq3.TABLE_TECHNO]
    starts = [t.rows[2 + 6 * k].cells[0].text for k in range(4)]
    assert starts[0].startswith("Coke Rate")
    assert starts[1].startswith("Overall CDI")
    assert starts[2].startswith("Fuel Rate")
    assert starts[3].startswith("BF Productivity")
