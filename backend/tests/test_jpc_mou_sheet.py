"""JPC report, MoU sheet: MoU / Var / %Ful. filled from mou_plan_table for
the report month, with the same plant rules as the actuals (5 Plants =
BSP..ISP; SAIL adds VISL for Hot Metal / Pig Iron, ASP+SSP+VISL otherwise).
Left blank when the month has no MoU loaded."""
import openpyxl
import pytest

import dbengine
import page_jpc_report as jpc
from v2dual import put_production

ACT = {("BSP", "Hot Metal"): 370.0, ("DSP", "Hot Metal"): 205.0, ("VISL", "Hot Metal"): 1.0,
       ("BSP", "Finished Steel"): 374.0, ("SSP", "Saleable Steel"): 13.9}
MOU = {("BSP", "Hot Metal"): 549.24, ("DSP", "Hot Metal"): 232.06, ("VISL", "Hot Metal"): 0.0,
       ("BSP", "Finished Steel"): 456.22, ("SSP", "Finished Steel"): 25.33}


@pytest.fixture
def cur(mysql_db):
    for t, data in (("production_table", ACT), ("mou_plan_table", MOU)):
        put_production(t, [("2026-09", p, i, v) for (p, i), v in data.items()])
    c = dbengine.connect()
    yield c.cursor()
    c.close()


def _rows(cur, month="2026-09"):
    ws = openpyxl.Workbook().active
    jpc._build_mou_sheet(ws, cur, month, "2025-09")
    out, item = {}, None
    for r in range(6, ws.max_row + 1):
        item = ws.cell(r, 1).value or item
        out[(item, ws.cell(r, 2).value)] = [ws.cell(r, c).value for c in range(3, 7)]
    return ws, out


def test_mou_var_ful_filled(cur):
    _, rows = _rows(cur)
    mou, act, var, ful = rows[("HOT METAL", "BSP")]
    assert (mou, act) == (pytest.approx(549.24), 370.0)
    assert var == pytest.approx(370.0 - 549.24)
    assert ful == pytest.approx(370.0 / 549.24 * 100, abs=0.05)


def test_mou_aggregates_follow_plant_rules(cur):
    _, rows = _rows(cur)
    assert rows[("HOT METAL", "5 Plants")][0] == pytest.approx(549.24 + 232.06)
    assert rows[("HOT METAL", "SAIL")][0] == pytest.approx(549.24 + 232.06)   # + VISL 0
    assert rows[("Finished Steel", "SAIL")][0] == pytest.approx(456.22 + 25.33)


def test_no_mou_loaded_leaves_columns_blank(cur):
    _, rows = _rows(cur, "2026-10")
    assert rows[("HOT METAL", "BSP")][0] is None and rows[("HOT METAL", "BSP")][2] is None


def test_month_header_sits_over_mou_block(cur):
    ws, _ = _rows(cur)
    assert ws.cell(4, 3).value == "Sep'26" and ws.cell(5, 3).value == "MoU"
