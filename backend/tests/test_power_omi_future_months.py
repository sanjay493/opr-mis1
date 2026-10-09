"""Power-OIS extractor: months the file hasn't reported yet keep only their
pre-filled PLAN and LAST-YEAR figures. The workbook also pre-fills other
cells for those months (SAIL's Sp. Power Cons target, CFP's formula zeros),
and none of them are actuals."""
import datetime
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "excel_extractors"))
from excel_extractor_power_omi import _FIXED_COLS, extract_power_omi  # noqa: E402

_LY_START = 23  # W


def _build(path, reported_through=9):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, label in enumerate(["OWN CPP", "JV CPP", "DRAWAL PP3", "TOTAL GEN"]):
        ws.cell(row=5, column=_LY_START + i, value=label)
    r = 7
    for plant in ["BSP", "DSP", "RSP", "BSL", "ISP", "SSP", "VISP", "CFP", "SAIL"]:
        ws.cell(row=r, column=1, value=plant)
        for k in range(12):
            y, m = (2026, 4 + k) if k < 9 else (2027, k - 8)
            row = r + k
            ws.cell(row=row, column=2, value=datetime.datetime(y, m, 1))
            ws.cell(row=row, column=_FIXED_COLS["plan_total"], value=100.0)
            ws.cell(row=row, column=_LY_START + 4, value=300.0)
            reported = (y, m) <= (2026, reported_through)
            if plant in ("SSP", "VISP"):
                # no generation columns at all, only grid drawal
                if reported:
                    ws.cell(row=row, column=_FIXED_COLS["total_power_consump"], value=20.0)
            elif plant == "CFP":
                # formula columns read 0 for months not reported yet
                ws.cell(row=row, column=_FIXED_COLS["actual_total"], value=0.0)
                if reported:
                    ws.cell(row=row, column=_FIXED_COLS["total_power_consump"], value=39.0)
            else:
                if reported:
                    ws.cell(row=row, column=_FIXED_COLS["actual_total"], value=90.0)
                    ws.cell(row=row, column=_FIXED_COLS["total_power_consump"], value=200.0)
                if plant == "SAIL":
                    # pre-filled for the whole FY
                    ws.cell(row=row, column=_FIXED_COLS["specific_power_cons"], value=436.0)
        cum = r + 12
        ws.cell(row=cum, column=2, value="Cum.")
        ws.cell(row=cum, column=_FIXED_COLS["total_power_consump"], value=55.0)
        r = cum + 1
    wb.save(path)


def test_unreported_months_keep_only_plan_and_last_year(tmp_path):
    f = tmp_path / "power.xlsx"
    _build(f)
    recs = extract_power_omi(str(f))["records"]
    future = [x for x in recs if x["report_month"] > "2026-09"]
    assert future, "plan / last-year rows for future months should still be extracted"
    bad = [x for x in future
           if not x["item_name"].startswith(("plan_", "last_year_"))]
    assert bad == []
    sail_sep = {x["item_name"]: x["value"] for x in recs
                if x["plant_name"] == "SAIL" and x["report_month"] == "2026-09"}
    assert sail_sep["specific_power_cons"] == 436.0


def test_cum_row_goes_to_latest_reported_month_for_every_plant(tmp_path):
    f = tmp_path / "power.xlsx"
    _build(f)
    recs = extract_power_omi(str(f))["records"]
    cum_months = {(x["plant_name"], x["report_month"]) for x in recs
                  if x["item_name"].endswith("_cum")}
    assert {m for _, m in cum_months} == {"2026-09"}
    assert {p for p, _ in cum_months} >= {"SSP", "VISP", "CFP"}


def test_double_dot_typo_is_read_and_flagged(tmp_path):
    f = tmp_path / "power.xlsx"
    _build(f)
    wb = openpyxl.load_workbook(f)
    ws = wb.active
    rsp_sep = next(r for r in range(7, ws.max_row + 1) if ws.cell(r, 1).value == "RSP") + 5
    ws.cell(row=rsp_sep, column=_FIXED_COLS["actual_own"], value="26..4")
    wb.save(f)
    out = extract_power_omi(str(f))
    own = [x["value"] for x in out["records"] if x["plant_name"] == "RSP"
           and x["report_month"] == "2026-09" and x["item_name"] == "actual_own"]
    assert own == [26.4]
    assert any("'26..4'" in w and "RSP 2026-09 actual_own" in w for w in out["warnings"])
