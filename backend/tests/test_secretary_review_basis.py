"""Secretary Review deck: production plan is MoU when mou_plan_table has the
report month for every deck item, else APP; the template's 'APP' header
cells follow the basis."""
import pytest
from pptx import Presentation
from pptx.util import Inches

import dbengine
import page_secretary_review as psr
from repo.conn import tx
from v2dual import put_production

_ITEMS = [db_item for _, db_item in psr.ITEMS]


def _cursor(mou_rows):
    with tx() as seed:
        put_production("mou_plan_table", mou_rows)
    return dbengine.connect().cursor()


def test_basis_mou_when_every_item_has_mou_for_month(mysql_db):
    cur = _cursor([("2026-08", "BSP", it, 1.0) for it in _ITEMS])
    assert psr.plan_basis(cur, "2026-08") == "mou"


def test_basis_app_when_no_mou(mysql_db):
    assert psr.plan_basis(_cursor([]), "2026-08") == "app"


def test_basis_app_when_mou_missing_an_item_or_month(mysql_db):
    cur = _cursor([("2026-08", "BSP", it, 1.0) for it in _ITEMS[:-1]] + [("2026-09", "BSP", _ITEMS[-1], 1.0)])
    assert psr.plan_basis(cur, "2026-08") == "app"


def test_basis_app_when_table_missing(mysql_db):
    # Drop the (empty) table for this test only, then restore it from its own
    # SHOW CREATE TABLE so later tests still see the schema (v2: MoU lives in production_monthly).
    with tx() as ddl:
        ddl.execute("SHOW CREATE TABLE production_monthly")
        create_sql = ddl.fetchone()[1]
        ddl.execute("DROP TABLE production_monthly")
    try:
        cur = dbengine.connect().cursor()
        assert psr.plan_basis(cur, "2026-08") == "app"
    finally:
        with tx() as ddl:
            ddl.execute(create_sql)


def test_plan_headers_relabelled():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_table(2, 3, Inches(0), Inches(0), Inches(3), Inches(1)).table
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(0, 2).text = "Item", "APP", "Actual"
    t.cell(1, 0).text = "Hot Metal APP note"
    psr._relabel_plan_headers(t, "MoU")
    assert [t.cell(0, i).text for i in range(3)] == ["Item", "MoU", "Actual"]
    assert t.cell(1, 0).text == "Hot Metal APP note"


@pytest.mark.parametrize("basis,label", [("app", "APP"), ("mou", "MoU")])
def test_basis_label(basis, label):
    assert psr.BASIS_LABEL[basis] == label
