"""MoU month-wise plan extractor (excel_extractor_mou_plan.py)."""
import datetime
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "excel_extractors"))
from excel_extractor_mou_plan import extract_mou_plan  # noqa: E402

from conftest import find_sample

MONTHS = [datetime.datetime(2026 if m >= 4 else 2027, m, 1) for m in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)]


def _build(path, sail_hm_apr=30.0):
    """Hot Metal: BSP 10/month, DSP 20/month, then 'SAIL - 5PL' and 'SAIL'
    totals (item name only on the first row, as in the merged original).
    Crude Steel: BSP 5, ASP 1, SAIL 6."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MoU 26-27"
    ws.cell(1, 1, "Monthwise Production MoU Target : 2026-27")
    ws.cell(3, 1, "Items")
    ws.cell(3, 2, "Plant")
    for i, d in enumerate(MONTHS):
        ws.cell(3, 3 + i, d)
    ws.cell(3, 15, "1st Qtr")
    rows = [
        ("Hot Metal", "BSP", 10.0), (None, "DSP", 20.0), (None, "SAIL - 5PL", 30.0), (None, "SAIL", 30.0),
        ("Crude Steel", "BSP", 5.0), (None, "ASP", 1.0), (None, "SAIL ", 6.0),
    ]
    for r, (item, plant, v) in enumerate(rows, start=4):
        if item:
            ws.cell(r, 1, item)
        ws.cell(r, 2, plant)
        for i in range(12):
            ws.cell(r, 3 + i, v)
        ws.cell(r, 15, v * 3)
    ws.cell(7, 3, sail_hm_apr)  # Hot Metal SAIL, Apr
    wb.save(path)


def test_plant_rows_only_with_page4_item_names(tmp_path):
    f = tmp_path / "mou.xlsx"
    _build(f)
    out = extract_mou_plan(str(f))
    keys = {(r["item_name"], r["plant_name"]) for r in out["records"]}
    assert keys == {("Hot Metal", "BSP"), ("Hot Metal", "DSP"),
                    ("Total Crude Steel", "BSP"), ("Total Crude Steel", "ASP")}
    assert len(out["months"]) == 12 and out["months"][0] == "2026-04" and out["months"][-1] == "2027-03"
    assert len(out["records"]) == 4 * 12
    assert out["items_found"] == ["Hot Metal", "Total Crude Steel"]
    assert out["warnings"] == []


def test_sail_row_mismatch_is_warned(tmp_path):
    f = tmp_path / "mou.xlsx"
    _build(f, sail_hm_apr=31.0)
    out = extract_mou_plan(str(f))
    assert any("Hot Metal" in w and "2026-04" in w for w in out["warnings"])


def test_no_header_raises(tmp_path):
    f = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(f)
    with pytest.raises(ValueError):
        extract_mou_plan(str(f))


def test_real_workbook():
    sample = find_sample("ABP/MoU 26-27.xlsx")
    if sample is None:
        pytest.skip("MoU 26-27.xlsx not available")
    out = extract_mou_plan(str(sample))
    assert out["items_found"] == ["Hot Metal", "Total Crude Steel", "Saleable Steel", "Pig Iron", "Finished Steel"]
    assert out["warnings"] == []
    bsp_hm_apr = [r["value"] for r in out["records"] if r["item_name"] == "Hot Metal"
                  and r["plant_name"] == "BSP" and r["report_month"] == "2026-04"]
    assert bsp_hm_apr == [pytest.approx(565.452483, abs=1e-5)]
