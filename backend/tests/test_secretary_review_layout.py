from pptx import Presentation

import secretary_review_layout as L
import secretary_review_pptx as sp


def test_period_labels_september():
    d = L.period_labels("2026-09")
    assert d["fy_start"] == 2026 and d["n"] == 6
    assert (d["mon"], d["cply"], d["ytd"], d["ytd_prev"]) == ("Sep’26", "Sep’25", "Apr-Sep’26", "Apr-Sep’25")
    assert (d["period_short"], d["period_hdr"]) == ("H1", "H-1")
    assert (d["fy_m2"], d["fy_m1"], d["fy_tgt"]) == ("FY25", "FY26", "FY'27 Target")
    assert (d["fy_label"], d["fy_prev_label"]) == ("2026-27", "2025-26")
    assert d["ytd_cat"] == "Apr-Sep"
    assert d["filename"] == "SECRETARY REVIEW Operations Inputs Sep26.pptx"


def test_period_labels_april():
    d = L.period_labels("2026-04")
    assert d["n"] == 1
    assert d["period_short"] == "Apr" and d["period_hdr"] == "Apr"
    assert d["ytd_cat"] == "Apr (YTD)"


def test_period_labels_january():
    d = L.period_labels("2027-01")
    assert d["fy_start"] == 2026 and d["n"] == 10
    assert (d["mon"], d["cply"], d["ytd"]) == ("Jan’27", "Jan’26", "Apr-Jan’27")
    assert d["period_short"] == "Apr-Jan"
    assert d["fy_tgt"] == "FY'27 Target"


def test_period_labels_quarters():
    assert L.period_labels("2026-06")["period_hdr"] == "Q-1"
    assert L.period_labels("2026-12")["period_hdr"] == "9M"
    assert L.period_labels("2027-03")["period_hdr"] == "FY"


def test_validate_month():
    assert L.validate_month("2026-09")
    assert not L.validate_month("2026-13")
    assert not L.validate_month("26-09")


def test_blocks_are_unique_and_complete():
    keys = [k for k, _, _ in L.BLOCKS]
    assert len(keys) == len(set(keys)) == 30
    assert "hl_SAIL" in keys and "cr_ISP_prev" in keys and "bd_RSP_MILL" in keys


def test_template_has_all_named_shapes_and_placeholders():
    prs = Presentation(str(L.TEMPLATE_PATH))
    names = set(sp.named_shapes(prs))
    missing = L.required_shape_names() - names
    assert not missing, sorted(missing)
    text = " ".join(
        cell.text for sh in sp.named_shapes(prs).values() if getattr(sh, "has_table", False) and sh.has_table
        for row in sh.table.rows for cell in row.cells
    )
    for ph in ("{MON}", "{YTD}", "{CPLY}", "{YTD_PREV}", "{PERIOD}"):
        assert ph in text, ph
    assert len(prs.slides) == 30
