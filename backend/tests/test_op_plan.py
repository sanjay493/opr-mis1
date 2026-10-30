"""Operating Plan (OP): repo.production's production_op_monthly functions."""
from repo import conn as rc
from repo import production as rp


def _n():
    with rc.cursor() as cur:
        cur.execute("SELECT n FROM data_version WHERE id = 1")
        return cur.fetchone()[0]


def test_save_read_update_clear(mysql_db):
    n0 = _n()
    assert rp.save_op_plan("2026-11", "BSP", [("Hot Metal", 520.5), ("Total Crude Steel", 470.0),
                                              ("  ", 1.0)], "a@sail.in") == (2, 0)
    assert _n() > n0
    assert [(n, v, by) for n, v, by, _ in rp.op_plan_items("BSP", "2026-11")] == [
        ("Hot Metal", 520.5, "a@sail.in"), ("Total Crude Steel", 470.0, "a@sail.in")]
    assert rp.op_plan_items("BSP", "2026-12") == [] and rp.op_plan_items("DSP", "2026-11") == []

    # update one, clear the other, clear an item that never had an OP
    assert rp.save_op_plan("2026-11", "BSP", [("Hot Metal", 530.0), ("Total Crude Steel", None),
                                              ("Never Entered", None)], "b@sail.in") == (1, 1)
    assert [(n, v, by) for n, v, by, _ in rp.op_plan_items("BSP", "2026-11")] == [("Hot Metal", 530.0, "b@sail.in")]

    rp.save_op_plan("2026-11", "RSP", [("Hot Metal", 300.0)], None)
    assert rp.op_plan_plant_counts("2026-11") == [("BSP", 1), ("RSP", 1)]


def test_op_does_not_touch_the_app(mysql_db):
    rp.save_production_plan("2026-11", "BSP", "Hot Metal", 500.0)
    rp.save_op_plan("2026-11", "BSP", [("Hot Metal", 530.0)], "a@sail.in")
    with rc.cursor() as cur:
        assert [tuple(r) for r in rp.plant_month_plan_items(cur, "BSP", "2026-11")] == [("Hot Metal", 500.0)]


def test_op_comparison(mysql_db):
    import asyncio
    import main

    for m, app, act, op in (("2026-04", 100.0, 90.0, 95.0), ("2026-05", 100.0, 110.0, None)):
        rp.save_production_plan(m, "BSP", "Hot Metal", app)
        rp.save_production_actual(m, "BSP", "Hot Metal", act)
        if op is not None:
            rp.save_op_plan(m, "BSP", [("Hot Metal", op)], "a@sail.in")
    rp.save_production_plan("2026-04", "BSP", "Oven Pushing (nos/day)", 300.0)   # 30 days
    rp.save_production_plan("2026-05", "BSP", "Oven Pushing (nos/day)", 362.0)   # 31 days
    rp.save_production_plan("2026-05", "DSP", "Hot Metal", 50.0)
    rp.save_op_plan("2026-05", "DSP", [("Hot Metal", 55.0)], None)

    def get(**kw):
        return asyncio.run(main.get_op_comparison(**kw))

    month = {r["item"]: r for r in get(month="2026-04", plant="BSP")["rows"]}
    assert month["Hot Metal"]["op"] == 95.0 and month["Hot Metal"]["act_vs_op"] == 94.7
    assert month["Hot Metal"]["op_vs_app"] == 95.0 and month["Hot Metal"]["act_vs_app"] == 90.0

    ytd = {r["item"]: r for r in get(month="2026-05", plant="BSP", period="ytd")["rows"]}
    hm = ytd["Hot Metal"]
    assert (hm["app"], hm["actual"], hm["op"], hm["op_months"], hm["months"]) == (200.0, 200.0, 95.0, 1, 2)
    assert hm["act_vs_op"] == 94.7          # April's actual against April's OP only
    assert hm["op_vs_app"] is None          # OP covers 1 of 2 months
    assert ytd["Oven Pushing (nos/day)"]["app"] == (300 * 30 + 362 * 31) / 61   # day-weighted, not summed

    sail = {g["item"]: g for g in get(month="2026-05", plant="SAIL")["groups"]}["Hot Metal"]
    assert [r["plant"] for r in sail["rows"]] == ["BSP", "DSP"]
    assert (sail["total"]["app"], sail["total"]["op"]) == (150.0, 55.0)
    assert sail["total"]["op_complete"] is False and sail["total"]["op_vs_app"] is None   # BSP has no May OP


def test_op_comparison_xlsx(mysql_db):
    import io
    import openpyxl
    import main

    rp.save_production_plan("2026-05", "BSP", "Hot Metal", 100.0)
    rp.save_production_actual("2026-05", "BSP", "Hot Metal", 110.0)
    rp.save_op_plan("2026-05", "BSP", [("Hot Metal", 105.0)], None)
    res = main.op_comparison_xlsx(month="2026-05")
    assert 'filename="OP_vs_APP_vs_Actual_2026-05.xlsx"' in res.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(res.body))
    assert wb.sheetnames == ["SAIL", "BSP", "DSP", "RSP", "BSL", "ISP", "ASP", "SSP", "VISL"]
    sail = [r for r in wb["SAIL"].iter_rows(min_row=5, values_only=True)]
    assert sail[:3] == [("Hot Metal",) + (None,) * 6, ("   BSP", 100.0, 105.0, 110.0, 110.0, 104.8, 105.0),
                        ("   SAIL", 100.0, 105.0, 110.0, 110.0, 104.8, 105.0)]
    assert [c.value for c in wb["BSP"][5]][:4] == ["Hot Metal", 100.0, 105.0, 110.0]
