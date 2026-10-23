"""Plant performance API: basis/decimals validation and the two-sheet xlsx."""
import io

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import api_plant_performance as app_mod

ROW = {"item": "HOT METAL", "plant": "BSP", "capacity": 7000.0,
       "values": [6688.1476, 565.4525, 500.12345, -65.33, 88.4, 520.0, -3.8, 85.0,
                  2800.0, 2700.5, -99.5, 96.4, 2600.0, 3.9, 90.0],
       "is_conversion": False, "is_sail_incl_conv": False}


@pytest.fixture
def tc(monkeypatch):
    monkeypatch.setattr(app_mod, "_rows", lambda month, basis: [dict(ROW)])
    app = FastAPI()
    app.include_router(app_mod.router)
    return TestClient(app)


def test_json_basis(tc):
    d = tc.get("/api/plant-performance-main-items?month=2026-08&basis=mou").json()
    assert d["labels"]["basis"] == "MoU" and d["has_plan"] is True


@pytest.mark.parametrize("q", ["basis=abp", "decimals=4", "decimals=x", "decimals=-1"])
def test_bad_params_400(tc, q):
    path = "" if q.startswith("basis") else "/xlsx"
    assert tc.get(f"/api/plant-performance-main-items{path}?month=2026-08&{q}").status_code in (400, 422)


def test_xlsx_two_sheets_with_decimals(tc):
    r = tc.get("/api/plant-performance-main-items/xlsx?month=2026-08&decimals=1")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["w.r.t APP", "w.r.t MoU"]
    ws = wb["w.r.t MoU"]
    assert ws.cell(3, 4).value.startswith("MoU")
    assert ws.cell(5, 6).number_format == "#,##0.0"
    assert ws.cell(5, 8).number_format == "0"


def test_pdf_html_has_both_pages():
    import plant_performance_pdf as ppdf
    sec = {"labels": {**app_mod._labels("2026-08", "app")}, "pct_idx": sorted(app_mod.PCT_IDX),
           "rows": [dict(ROW)], "has_plan": True}
    sec2 = {**sec, "labels": app_mod._labels("2026-08", "mou")}
    html = ppdf.build_html([sec, sec2], 2, {"perf_app_bg": "#111111"})
    assert html.count('class="sheet"') == 2 and "w.r.t MoU" in html and "500.12" in html and "#111111" in html


def test_pdf_small_negative_pct_is_zero():
    import plant_performance_pdf as ppdf
    assert ppdf._p(-0.3) == "0" and ppdf._p(-0.6) == "-1"


def test_pdf_uses_only_config_colours():
    import re
    import plant_performance_pdf as ppdf
    sec = {"labels": app_mod._labels("2026-08", "app"), "pct_idx": sorted(app_mod.PCT_IDX),
           "rows": [dict(ROW)], "has_plan": True}
    colors = {k: "#123456" for k in ("perf_app_bg", "perf_app_text", "perf_act_bg", "perf_act_text", "perf_gr_bg",
                                     "perf_cap_bg", "perf_cap_text", "perf_app_border", "border_medium",
                                     "border_light", "border_black", "text_primary", "text_black",
                                     "table_header_bg")}
    html = ppdf.build_html([sec], 3, colors)
    assert set(re.findall(r"#[0-9a-fA-F]{6}\b", html)) == {"#123456"}


def test_pdf_endpoint_runs_off_the_event_loop():
    import inspect
    assert not inspect.iscoroutinefunction(app_mod.plant_performance_pdf_download)


def _sec(basis="app"):
    return {"labels": app_mod._labels("2026-08", basis), "pct_idx": sorted(app_mod.PCT_IDX),
            "rows": [dict(ROW)], "has_plan": True}


def test_pdf_matches_report_page4_portrait_layout():
    # Same sheet as the monthly report's page 4: A4 portrait with page 4's
    # margins, its 18 fixed column widths and blue heading.
    import plant_performance_pdf as ppdf
    html = ppdf.build_html([_sec()], 2, {})
    assert "size: A4 portrait" in html and "margin: 7mm 10mm 5mm 10mm" in html
    assert html.count("<col ") == 18
    assert "rgb(0,0,255)" in html


def test_pdf_numbers_have_no_thousands_separator():
    # Page 4 prints 6688, not 6,688.
    import plant_performance_pdf as ppdf
    html = ppdf.build_html([_sec()], 2, {})
    assert ">6688.15<" in html and "6,688" not in html


def test_pdf_rounded_negative_zero_prints_zero():
    import plant_performance_pdf as ppdf
    assert ppdf._q(-0.246, 0) == "0" and ppdf._q(-0.0004, 3) == "0.000" and ppdf._q(-0.6, 0) == "-1"


def test_pdf_decimals_use_auto_width_table():
    # 0 decimals: page 4's fixed column widths. With decimals the widths
    # become minimums so the longer numbers aren't clipped.
    import plant_performance_pdf as ppdf
    assert '<table class="auto">' not in ppdf.build_html([_sec()], 0, {})
    assert '<table class="auto">' in ppdf.build_html([_sec()], 3, {})
