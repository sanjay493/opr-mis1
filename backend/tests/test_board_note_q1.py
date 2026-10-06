"""Q-1 Board Note adapter: template index map and title text (no DB needed)."""

import docx

import board_note_q1 as bq1


def _doc():
    return docx.Document(bq1.TEMPLATE_PATH)


def test_title_text_names_the_quarter_and_its_months():
    assert bq1._title_text(2026) == "PRODUCTION PERFORMANCE DURING QUARTER-1’26-27 (APRIL’26, MAY’26, JUNE’26) "


def test_index_map_points_at_the_headings_it_claims():
    paras = _doc().paragraphs
    assert paras[bq1.P_SAIL_HEADING].text.startswith("SAIL: Production performance during Q-1")
    assert paras[bq1.P_PLANT_SECTION_TITLE].text.startswith("Plant wise: Production performance during Q-1")
    assert paras[bq1.P_TECHNO_SECTION_TITLE].text.startswith("SAIL and Plant-wise Major Techno-economic")
    for plant, idx in bq1.WHY_HEADING.items():
        assert paras[idx].text.startswith("Production Performance during Q-1"), plant
    for key, idx in bq1.HIGHLIGHTS_HEADING.items():
        assert paras[idx].text == "Highlights:", key
    for key, idx in bq1.HIGHLIGHTS_LEADIN.items():
        if idx is not None:
            assert paras[idx].text.endswith("for following:"), key
    for plant, cfg in bq1.PLANT_CONFIG.items():
        assert paras[cfg["label"]].text.startswith("Q-1"), plant
        assert paras[cfg["summary"]].text.startswith(("Hot Metal production", "Crude Steel production",
                                                      "Saleable Steel production")), plant


def test_every_slot_is_a_non_empty_template_paragraph():
    paras = _doc().paragraphs
    for key, idxs in bq1.HIGHLIGHTS_SLOTS.items():
        for i in idxs:
            # SAIL's highlights slot is a deliberately empty bullet paragraph
            # that the generated bullets fill; every other slot carries text.
            if key == ("SAIL", "q1"):
                continue
            assert paras[i].text.strip(), (key, i)
    assert paras[bq1.HIGHLIGHTS_SLOTS[("SAIL", "q1")][0]].text == ""
    for plant, idxs in bq1.WHY_SLOTS.items():
        for i in idxs:
            assert paras[i].text.strip(), (plant, i)


def test_highlight_slots_are_contiguous_so_no_blank_line_splits_the_bullets():
    paras = _doc().paragraphs
    for key, idxs in bq1.HIGHLIGHTS_SLOTS.items():
        if key == ("SAIL", "q1"):
            continue
        gaps = [i for i in range(idxs[0], idxs[-1] + 1)
                if i not in idxs and paras[i].text.strip() == ""]
        assert gaps == [] or key in bq1.HIGHLIGHTS_BLANKS_REMOVE, (key, gaps)
    assert bq1.HIGHLIGHTS_BLANKS_REMOVE[("BSL", "q1")] == [139]
    for key, idxs in bq1.HIGHLIGHTS_SLOTS.items():
        for i in bq1.HIGHLIGHTS_BLANKS_REMOVE.get(key, []):
            assert paras[i].text == "" and i not in idxs


def test_unit_tables_match_their_item_rows():
    tables = _doc().tables
    for plant, cfg in bq1.PLANT_CONFIG.items():
        table = tables[cfg["table"]]
        assert len(table.rows) == 2 + len(cfg["item_rows"]), plant
        for row_idx, display, _db_item, _add_conv in cfg["item_rows"]:
            assert table.rows[row_idx].cells[0].text.strip() == display, (plant, row_idx)
    sail = tables[bq1.TABLE_SAIL]
    assert [r.cells[0].text.strip() for r in sail.rows[2:]] == [
        "Hot Metal", "Crude Steel", "Saleable Steel", "Finished Steel"]


def test_techno_and_annexure_tables_have_the_expected_shape():
    tables = _doc().tables
    assert len(tables[bq1.TABLE_TECHNO].rows) == 26
    assert len(tables[bq1.TABLE_ANNEXURE_INDEX].rows) == 6
