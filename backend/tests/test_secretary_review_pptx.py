import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

import secretary_review_pptx as sp


def _deck_with_table(rows=3, cols=2):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(6), Inches(2))
    gf.name = "tbl_x"
    return prs, slide, gf


def test_norm_collapses_whitespace():
    assert sp.norm("  Crude\n  Steel ") == "crude steel"


def test_named_shapes_finds_by_name():
    prs, _, _ = _deck_with_table()
    assert "tbl_x" in sp.named_shapes(prs)


def test_set_text_lines_escapes_and_drops_blanks():
    prs, _, gf = _deck_with_table()
    tf = gf.table.cell(0, 1).text_frame
    sp.set_text_lines(tf, ["BSP (Sep'26):", "", "BF-4: R&D <test> – 15 days", "   "], header_bold=True)
    paras = tf.paragraphs
    assert [p.text for p in paras] == ["BSP (Sep'26):", "BF-4: R&D <test> – 15 days"]
    assert paras[0].runs[0].font.bold is True
    assert paras[1].runs[0].font.bold is False
    buf = io.BytesIO()
    prs.save(buf)                      # must serialise (XML-escaped)
    Presentation(io.BytesIO(buf.getvalue()))


def test_set_text_lines_header_bold_accepts_colon_dash():
    prs, _, gf = _deck_with_table()
    tf = gf.table.cell(0, 1).text_frame
    sp.set_text_lines(tf, ["Highlights for following:-  ", "Plan met:–", "body - text"], header_bold=True)
    assert [p.runs[0].font.bold for p in tf.paragraphs] == [True, True, False]


def test_set_text_lines_empty_list_leaves_one_empty_paragraph():
    prs, _, gf = _deck_with_table()
    tf = gf.table.cell(0, 0).text_frame
    sp.set_text_lines(tf, [])
    assert [p.text for p in tf.paragraphs] == [""]


def test_replace_placeholders_in_tables_and_text():
    prs, slide, gf = _deck_with_table()
    gf.table.cell(0, 0).text = "% Growth Over {CPLY}"
    tb = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    tb.text_frame.text = "SAIL Performance : {MON} & {YTD}"
    sp.replace_placeholders(prs, {"CPLY": "Sep'25", "MON": "Sep'26", "YTD": "Apr-Sep'26"})
    assert gf.table.cell(0, 0).text == "% Growth Over Sep'25"
    assert tb.text_frame.text == "SAIL Performance : Sep'26 & Apr-Sep'26"


def test_remove_row_and_shape():
    prs, slide, gf = _deck_with_table(rows=3)
    sp.remove_row(gf.table, 1)
    assert len(gf.table.rows) == 2
    sp.remove_shape(gf)
    assert "tbl_x" not in sp.named_shapes(prs)


def test_replace_chart_data_keeps_point_formatting():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["FY25", "FY26", "FY'27 Target", "Sep", "Apr-Sep"]
    cd.add_series("Series 1", (1, 2, 3, 4, 5))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), cd).chart
    pt = chart.plots[0].series[0].points[2]
    pt.format.fill.solid()
    before = len(chart._chartSpace.xpath(".//c:dPt"))
    sp.replace_chart_data(chart, ["FY26", "FY27", "FY'28 Target", "Oct", "Apr-Oct"], [421, None, 400, 432, 424], "0")
    assert len(chart._chartSpace.xpath(".//c:dPt")) == before == 1
    assert list(chart.plots[0].categories) == ["FY26", "FY27", "FY'28 Target", "Oct", "Apr-Oct"]
    assert chart.plots[0].series[0].values == (421.0, None, 400.0, 432.0, 424.0)


def test_set_text_lines_drops_xml_illegal_characters():
    prs, _, gf = _deck_with_table()
    tf = gf.table.cell(0, 1).text_frame
    bad = "a\x00b\x07c\x0bd\x0c\x1fe￾￿\tf"
    sp.set_text_lines(tf, [bad, "\x01\x02", "ok"])
    assert [p.text for p in tf.paragraphs] == ["abcde\tf", "ok"]
    buf = io.BytesIO()
    prs.save(buf)                                   # lxml would raise on the raw characters
    out = Presentation(io.BytesIO(buf.getvalue()))
    assert sp.named_shapes(out)["tbl_x"].table.cell(0, 1).text_frame.paragraphs[0].text == "abcde\tf"
