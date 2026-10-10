"""Q-3 + 9M Board Note docx generator.

Fills `board_note_templates/board_note_q3_template.docx` (a verbatim copy of
`Report_format/BN_ Production Q-3 and 9M'25-26.docx`) for any financial year,
using the same technique and shared helpers as `board_note_q2.py`.

MoU: the Finished-Steel-vs-MoU opening and table 0 take the annual /
quarterly / 9M Finished Steel MoU from mou_plan_table (bnc.fs_mou) when it
has the whole FY. Otherwise the template's MoU cells (FY2025-26 targets) are
blanked rather than carried over, every %Ful that depends on them is blank,
and the opening drops its MoU clause. Actuals are always filled from the DB.

Known limitations (documented, not bugs):
  - ASP, SSP and VISL have no "Highlights:" section in this template for
    either period, so no highlights are written for them.
  - The "This was achieved through..." paragraphs under the techno section
    describe FY2025-26 drivers and are removed rather than shipped stale.
  - No plant has a "why production fell short" block for 9M, only Q-3.
"""

import io

import docx

import board_note_common as bnc
import board_note_docx_utils as bdu
import board_note_manual_text as bnm
import board_note_q2 as bq2
import db
import page_techno as pt
import techno_period as tp

TEMPLATE_PATH = "board_note_templates/board_note_q3_template.docx"

# -- Title / MoU opening (P3, P7, P9, P11) -----------------------------------
P_TITLE = 3
P_MOU_HEADING = 7            # "SAIL: Production performance for Finished Steel w.r.t. MoU "
P_FS_Q3_NARRATIVE = 9
P_FS_9M_NARRATIVE = 11
TABLE_MOU = 0                # Plant | MoU | Q-3 MoU/Act/%Ful | Apr-Dec MoU/Act/%Ful
MOU_ROW_PLANT = {2: "BSP", 3: "DSP", 4: "RSP", 5: "BSL", 6: "ISP"}
MOU_ROW_SPECIAL_STEEL = 7
MOU_ROW_SAIL = 8

# -- Intro -------------------------------------------------------------------
P_INTRO_Q3 = 18              # "SAIL and Plant-wise Production performance during Q-3’25-26 and "
P_INTRO_9M = 19              # "9M’25-26 (Apr-Dec’25) w.r.t. ABP"
P_PLANT_SECTION_TITLE = 49

# -- SAIL --------------------------------------------------------------------
TABLE_SAIL_Q3 = 1
TABLE_SAIL_9M = 2
P_SAIL_Q3_HEADING = 21
P_SAIL_Q3_LABEL = 23
P_SAIL_Q3_SUMMARY = 24
P_SAIL_9M_HEADING = 35
P_SAIL_9M_LABEL = 37
P_SAIL_9M_SUMMARY = 38

# -- Techno and annexure -----------------------------------------------------
P_TECHNO_SECTION_TITLE = 318
TABLE_TECHNO = 19            # Parameters | Plant | Target | Q-3cur | Q-3cply | 9Mcur | 9Mcply
P_TECHNO_PERF_HEADING = 321  # "Techno-economic Performance during Q-3’25-26:"
P_TECHNO_COKE = 322
P_TECHNO_FUEL = 323
P_TECHNO_BFP = 324
TECHNO_STALE_REMOVE = [326, 327, 328]   # "This was achieved through..." FY2025-26 drivers
TABLE_ANNEXURE_INDEX = 20

TECHNO_PLANTS = ["BSP", "DSP", "RSP", "BSL", "ISP", "SAIL"]
TECHNO_PARAMS = ["Coke Rate", "CDI Rate", "Fuel Rate", "BF Productivity"]

BIG5_ITEMS = bq2.BIG5_ITEMS
SAIL_ITEMS = bq2.SAIL_ITEMS
SPECIAL2_ITEMS = bq2.SPECIAL2_ITEMS
VISL_ITEMS = bq2.VISL_ITEMS
SPECIAL_STEEL_PLANTS = ["ASP", "SSP", "VISL"]

# -- Per-plant index map -----------------------------------------------------
PLANT_CONFIG = {
    "BSP": dict(table_q3=3, table_9m=4, label_q3=53, summary_q3=54, label_9m=77, summary_9m=78,
                item_rows=BIG5_ITEMS, kind="big"),
    "DSP": dict(table_q3=5, table_9m=6, label_q3=101, summary_q3=102, label_9m=124, summary_9m=125,
                item_rows=BIG5_ITEMS, kind="big"),
    "RSP": dict(table_q3=7, table_9m=8, label_q3=135, summary_q3=136, label_9m=157, summary_9m=158,
                item_rows=BIG5_ITEMS, kind="big"),
    "BSL": dict(table_q3=9, table_9m=10, label_q3=171, summary_q3=172, label_9m=205, summary_9m=206,
                item_rows=BIG5_ITEMS, kind="big"),
    "ISP": dict(table_q3=11, table_9m=12, label_q3=217, summary_q3=218, label_9m=252, summary_9m=253,
                item_rows=BIG5_ITEMS, kind="big"),
    "ASP": dict(table_q3=13, table_9m=14, label_q3=267, summary_q3=268, label_9m=283, summary_9m=284,
                item_rows=SPECIAL2_ITEMS, kind="small"),
    "SSP": dict(table_q3=15, table_9m=16, label_q3=289, summary_q3=290, label_9m=298, summary_9m=299,
                item_rows=SPECIAL2_ITEMS, kind="small"),
    "VISL": dict(table_q3=17, table_9m=18, label_q3=303, summary_q3=305, label_9m=312, summary_9m=314,
                 item_rows=VISL_ITEMS, kind="visl"),
}
PLANT_ORDER = ["BSP", "DSP", "RSP", "BSL", "ISP", "ASP", "SSP", "VISL"]

# -- Highlights bullet slots -------------------------------------------------
HIGHLIGHTS_SLOTS = {
    ("SAIL", "q3"): [30, 31, 32, 33],
    ("SAIL", "9m"): [43, 44, 45],
    ("BSP", "q3"): [59, 60],
    ("BSP", "9m"): [83, 84, 85, 86, 87, 88, 89, 90, 91, 93, 94, 95, 96, 97],
    ("DSP", "q3"): [106, 107],
    ("DSP", "9m"): [129, 130, 131],
    ("RSP", "q3"): [140, 141],
    ("RSP", "9m"): [164, 165, 166, 167],
    ("BSL", "q3"): [177, 178, 179, 180, 181, 184, 185],
    ("BSL", "9m"): [211, 212, 213],
    ("ISP", "q3"): [223, 224, 225, 226, 227],
    ("ISP", "9m"): [258, 259, 260, 261, 262, 263],
}
HIGHLIGHTS_HEADING = {
    ("SAIL", "q3"): 28, ("SAIL", "9m"): 41,
    ("BSP", "q3"): 57, ("BSP", "9m"): 81,
    ("DSP", "q3"): 105, ("DSP", "9m"): 128,
    ("RSP", "q3"): 139, ("RSP", "9m"): 162,
    ("BSL", "q3"): 175, ("BSL", "9m"): 209,
    ("ISP", "q3"): 221, ("ISP", "9m"): 256,
}
HIGHLIGHTS_LEADIN = {
    ("SAIL", "q3"): 29, ("SAIL", "9m"): 42,
    ("BSP", "q3"): 58, ("BSP", "9m"): 82,
    ("DSP", "q3"): None, ("DSP", "9m"): None,   # no lead-in sentence in the template
    ("RSP", "q3"): None, ("RSP", "9m"): 163,
    ("BSL", "q3"): 176, ("BSL", "9m"): 210,
    ("ISP", "q3"): 222, ("ISP", "9m"): 257,
}
# Lead-in rewritten to the narrower claim that best_ever_bullets checks.
LEADIN_TEXT_FIX = {
    ("BSL", "q3"): (176, "BSL achieved best ever Q-3 production for following:"),
}
# A second "...for following:" lead-in in the middle of a bullet block; removed.
SECOND_LEADIN_REMOVE = {
    ("BSL", "q3"): 183,
}
# Empty spacer paragraphs that would otherwise show as a blank line between
# generated bullets. Removed before the bullets are filled.
HIGHLIGHTS_BLANKS_REMOVE = {
    ("BSL", "q3"): [182],
    ("BSP", "9m"): [92],
}

# -- "Production Performance during Q-3’25-26:" why blocks (Q-3 only) --------
WHY_HEADING = {
    "BSP": 62, "DSP": 109, "RSP": 143, "BSL": 187, "ISP": 229,
    "ASP": 271, "SSP": 293, "VISL": 308,
}
WHY_SLOTS = {
    "BSP": [64, 65, 66, 67, 68, 70, 71, 72, 73, 75],
    "DSP": [111, 112, 113, 114, 115, 116, 117, 118, 119, 121, 122],
    "RSP": [145, 146, 147, 148, 150, 151, 152, 153, 155],
    "BSL": [189, 191, 192, 193, 195, 196, 197, 198, 199, 200, 201, 202, 203],
    "ISP": [231, 232, 233, 234, 236, 237, 239, 240, 241, 242, 243, 244, 245, 246, 247, 248, 249, 250],
    "ASP": [273, 276, 277, 278, 279, 281],
    "SSP": [295, 296],
    "VISL": [310],
}


def _title_text(fy_start):
    yy = str(fy_start)[2:]
    zz = str(fy_start + 1)[2:]
    m1, m2, m3 = (bq2._month_name(m) for m in bnc.quarter_months(fy_start, 3))
    return (f"PRODUCTION PERFORMANCE DURING QUARTER-3’{yy}-{zz} "
            f"({m1}’{yy}, {m2}’{yy}, {m3}’{yy} and APRIL-DECEMBER’{yy}) ")


def _improvement(techno_by_param, param, higher_is_better=False):
    sec = techno_by_param[param]
    row = next(r for r in sec["rows"] if r["plant"] == "SAIL")
    return bnc.improvement_pct(row["values"]["Q3cur"]["value"], row["values"]["Q3cply"]["value"],
                               higher_is_better)


def _fmt_techno_target(param, v):
    if v is None:
        return None
    return f"{v:.2f}" if param == "BF Productivity" else str(round(v))


def _special_steel_total(cur, months):
    total, found = 0.0, False
    for p in SPECIAL_STEEL_PLANTS:
        v = bnc.period_sum(cur, "act", months, p, "Finished Steel")
        if v is not None:
            total += v
            found = True
    return total if found else None


def _sail_finished_with_conv(cur, months):
    act = bnc.period_sum(cur, "act", months, "SAIL", "Finished Steel")
    if act is None:
        return None
    cv = bnc.conv_sum(cur, months)
    return act + cv if cv is not None else None


def generate(fy: str) -> bytes:
    fy_start = int(fy[:4])
    yy = str(fy_start)[2:]

    Q3_CUR = bnc.quarter_months(fy_start, 3)
    Q3_CPLY = bnc.quarter_months(fy_start - 1, 3)
    L_CUR = bnc.long_period_months(fy_start, 3)
    L_CPLY = bnc.long_period_months(fy_start - 1, 3)

    q3_label = bnc.quarter_label(fy_start, 3)
    q3_cply_label = bq2._cply_label_nl(bnc.quarter_label, fy_start, 3)
    l_label = bnc.long_period_label(fy_start, 3)
    l_cply_label = bq2._cply_label_nl(bnc.long_period_label, fy_start, 3)

    manual = bnm.get_manual_text(fy, 3)

    conn = db.connect()
    cur = conn.cursor()
    try:
        doc = docx.Document(TEMPLATE_PATH)
        tables = doc.tables
        paras = doc.paragraphs

        def para(i):
            return paras[i]

        def apply_highlights(plant, period, bullets):
            key = (plant, period)
            slots_idx = HIGHLIGHTS_SLOTS.get(key)
            if not slots_idx:
                return
            fix = LEADIN_TEXT_FIX.get(key)
            if fix:
                idx, text = fix
                bdu.set_para(para(idx), text)
            second = SECOND_LEADIN_REMOVE.get(key)
            if second is not None:
                bdu.remove_paragraph(para(second))
            for blank_idx in HIGHLIGHTS_BLANKS_REMOVE.get(key, []):
                bdu.remove_paragraph(para(blank_idx))
            bdu.fill_variable_bullets([para(i) for i in slots_idx], bullets)
            if not bullets:
                heading_idx = HIGHLIGHTS_HEADING.get(key)
                leadin_idx = HIGHLIGHTS_LEADIN.get(key)
                if heading_idx is not None:
                    bdu.remove_paragraph(para(heading_idx))
                if leadin_idx is not None:
                    bdu.remove_paragraph(para(leadin_idx))

        # ---- unit tables -----------------------------------------------
        bq2.set_abp_header(tables[TABLE_SAIL_Q3], fy_start)
        bq2.set_period_header(tables[TABLE_SAIL_Q3], q3_label, q3_cply_label)
        bq2.fill_period_table(cur, tables[TABLE_SAIL_Q3], "SAIL", SAIL_ITEMS, Q3_CUR, Q3_CPLY, fy_start)

        bq2.set_abp_header(tables[TABLE_SAIL_9M], fy_start)
        bq2.set_period_header(tables[TABLE_SAIL_9M], l_label, l_cply_label)
        bq2.fill_period_table(cur, tables[TABLE_SAIL_9M], "SAIL", SAIL_ITEMS, L_CUR, L_CPLY, fy_start)

        for plant in PLANT_ORDER:
            cfg = PLANT_CONFIG[plant]
            bq2.set_abp_header(tables[cfg["table_q3"]], fy_start)
            bq2.set_period_header(tables[cfg["table_q3"]], q3_label, q3_cply_label)
            bq2.fill_period_table(cur, tables[cfg["table_q3"]], plant, cfg["item_rows"], Q3_CUR, Q3_CPLY, fy_start)

            bq2.set_abp_header(tables[cfg["table_9m"]], fy_start)
            bq2.set_period_header(tables[cfg["table_9m"]], l_label, l_cply_label)
            bq2.fill_period_table(cur, tables[cfg["table_9m"]], plant, cfg["item_rows"], L_CUR, L_CPLY, fy_start)

        # ---- table 0: Finished Steel vs MoU (MoU blank if not in the DB) ----
        FY_CUR = bnc.fy_months(fy_start)
        mou_in_db = bnc.fs_mou(cur, FY_CUR, bnc.SAIL_MOU_PLANTS) is not None

        def _mou(months, plants):
            return bnc.fs_mou(cur, months, plants) if mou_in_db else None

        t0 = tables[TABLE_MOU]
        mou_col = f"MoU \n{fy_start}-{str(fy_start + 1)[2:]}"
        for row in (t0.rows[0].cells, t0.rows[1].cells):
            bdu.set_cell(row[1], mou_col)
        for c in (2, 3, 4):
            bdu.set_cell(t0.rows[0].cells[c], q3_label)
        for c in (5, 6, 7):
            bdu.set_cell(t0.rows[0].cells[c], l_label)

        def _fill_mou_row(ridx, q3_act, l_act, plants):
            cells = t0.rows[ridx].cells
            ann, q3_mou, l_mou = _mou(FY_CUR, plants), _mou(Q3_CUR, plants), _mou(L_CUR, plants)
            bdu.set_cell(cells[1], bnc.fmt_ann(ann))
            bdu.set_cell(cells[2], bnc.fmt_ann(q3_mou))
            bdu.set_cell(cells[3], bnc.fmt_ann(q3_act))
            bdu.set_cell(cells[4], bnc.pct_ful(q3_act, q3_mou))
            bdu.set_cell(cells[5], bnc.fmt_ann(l_mou))
            bdu.set_cell(cells[6], bnc.fmt_ann(l_act))
            bdu.set_cell(cells[7], bnc.pct_ful(l_act, l_mou))

        for ridx, plant in MOU_ROW_PLANT.items():
            _fill_mou_row(ridx,
                          bnc.period_sum(cur, "act", Q3_CUR, plant, "Finished Steel"),
                          bnc.period_sum(cur, "act", L_CUR, plant, "Finished Steel"), [plant])
        _fill_mou_row(MOU_ROW_SPECIAL_STEEL,
                      _special_steel_total(cur, Q3_CUR), _special_steel_total(cur, L_CUR), SPECIAL_STEEL_PLANTS)
        _fill_mou_row(MOU_ROW_SAIL,
                      _sail_finished_with_conv(cur, Q3_CUR), _sail_finished_with_conv(cur, L_CUR),
                      bnc.SAIL_MOU_PLANTS)
        for row in t0.rows:
            for cell in row.cells:
                bdu.clear_strike_and_colour(cell)

        # ---- Finished Steel opening (MoU clause only with MoU in the DB) ----
        _, _, fs_q3_act, _, fs_q3_cply, fs_q3_gr = bnc.row_values(
            cur, "SAIL", "Finished Steel", Q3_CUR, Q3_CPLY, fy_start, add_conv=True)
        _, _, fs_l_act, _, fs_l_cply, fs_l_gr = bnc.row_values(
            cur, "SAIL", "Finished Steel", L_CUR, L_CPLY, fy_start, add_conv=True)
        bdu.set_para(para(P_MOU_HEADING), "SAIL: Production performance for Finished Steel "
                     + ("w.r.t. MoU " if mou_in_db else ""))
        q3_clause = bnc.mou_clause(fs_q3_act, _mou(Q3_CUR, bnc.SAIL_MOU_PLANTS), "quarterly")
        l_clause = bnc.mou_clause(fs_l_act, _mou(L_CUR, bnc.SAIL_MOU_PLANTS), f"Apr-Dec’{yy}")
        bdu.set_para(
            para(P_FS_Q3_NARRATIVE),
            f"The Production of Finished Steel during {q3_label} was {bnc.fmt_mt(fs_q3_act)} MT{q3_clause} "
            f"with a growth of {bnc.fmt_pct(fs_q3_gr)}% over CPLY."
        )
        bdu.set_para(
            para(P_FS_9M_NARRATIVE),
            f"The Production of Finished Steel during April-December’{yy} was {bnc.fmt_mt(fs_l_act)} MT{l_clause}, "
            f"with a growth of {bnc.fmt_pct(fs_l_gr)}% over CPLY. "
        )
        charts = bdu.fs_charts(doc)
        if "quarter" in charts:
            bdu.fill_fs_chart(charts["quarter"], q3_label.replace("’", "'"),
                              _mou(Q3_CUR, bnc.SAIL_MOU_PLANTS), fs_q3_act, fs_q3_cply)
        if "long" in charts:
            bdu.fill_fs_chart(charts["long"], f"Apr-Dec'{yy}",
                              _mou(L_CUR, bnc.SAIL_MOU_PLANTS), fs_l_act, fs_l_cply)

        # ---- title, intro, plant-wise title ---------------------------------
        bdu.set_para(para(P_TITLE), _title_text(fy_start))
        bdu.set_para(para(P_INTRO_Q3), f"SAIL and Plant-wise Production performance during {q3_label} and ")
        bdu.set_para(para(P_INTRO_9M), f"{l_label} (Apr-Dec’{yy}) w.r.t. ABP")
        bdu.set_para(para(P_PLANT_SECTION_TITLE),
                     f"Plant wise: Production performance during {q3_label} and {l_label} w.r.t. ABP")

        # ---- SAIL Q-3 and 9M -------------------------------------------------
        bdu.set_para(para(P_SAIL_Q3_HEADING), f"SAIL: Production performance during {q3_label} ")
        bdu.set_para(para(P_SAIL_Q3_LABEL), f"{q3_label}:")
        bdu.set_para(para(P_SAIL_Q3_SUMMARY),
                     bq2.plant_summary_big(cur, "SAIL", Q3_CUR, Q3_CPLY, fy_start).replace("{P}", q3_label))
        sail_q3_bullets = bnc.best_ever_bullets(cur, "SAIL", lambda f: bnc.quarter_months(f, 3), fy_start, "q3")
        sail_q3_bullets += bnm.additional_highlight_lines(manual, "SAIL")
        apply_highlights("SAIL", "q3", sail_q3_bullets)

        bdu.set_para(para(P_SAIL_9M_HEADING), f"SAIL: Production performance during {l_label} ")
        bdu.set_para(para(P_SAIL_9M_LABEL), f"{l_label}:")
        bdu.set_para(para(P_SAIL_9M_SUMMARY),
                     bq2.plant_summary_big(cur, "SAIL", L_CUR, L_CPLY, fy_start).replace("{P}", l_label))
        sail_9m_bullets = bnc.best_ever_bullets(cur, "SAIL", lambda f: bnc.long_period_months(f, 3), fy_start, "9m")
        apply_highlights("SAIL", "9m", sail_9m_bullets)

        # ---- per-plant ---------------------------------------------------------
        for plant in PLANT_ORDER:
            cfg = PLANT_CONFIG[plant]
            summary_fn = bq2._SUMMARY_FN[cfg["kind"]]

            bdu.set_para(para(cfg["label_q3"]), f"{q3_label}:")
            bdu.set_para(para(cfg["summary_q3"]),
                         summary_fn(cur, plant, Q3_CUR, Q3_CPLY, fy_start).replace("{P}", q3_label))
            q3_bullets = bnc.best_ever_bullets(cur, plant, lambda f: bnc.quarter_months(f, 3), fy_start, "q3")
            q3_bullets += bnm.additional_highlight_lines(manual, plant)
            apply_highlights(plant, "q3", q3_bullets)

            bdu.set_para(para(cfg["label_9m"]), f"{l_label}:")
            bdu.set_para(para(cfg["summary_9m"]),
                         summary_fn(cur, plant, L_CUR, L_CPLY, fy_start).replace("{P}", l_label))
            nine_m_bullets = bnc.best_ever_bullets(cur, plant, lambda f: bnc.long_period_months(f, 3), fy_start, "9m")
            apply_highlights(plant, "9m", nine_m_bullets)

            why_heading_idx = WHY_HEADING.get(plant)
            if why_heading_idx is not None:
                bdu.set_para(para(why_heading_idx), f"Production Performance during {q3_label}:")
            placeholder = f"[Add production narrative for {plant} here]"
            lines = bnm.why_narrative_lines(manual, plant, placeholder)
            bdu.fill_variable_bullets([para(i) for i in WHY_SLOTS[plant]], lines)

        # ---- techno section ---------------------------------------------------
        periods = [
            {"label": "Q3cur", "months": Q3_CUR}, {"label": "Q3cply", "months": Q3_CPLY},
            {"label": "NMcur", "months": L_CUR}, {"label": "NMcply", "months": L_CPLY},
        ]
        techno_res = tp.build_period_report(TECHNO_PLANTS, TECHNO_PARAMS, periods)
        techno_by_param = {s["parameter"]: s for s in techno_res["sections"]}

        report_month = L_CUR[-1]
        targets = {}
        for plant in ["BSP", "DSP", "RSP", "BSL", "ISP"]:
            raw = pt._get_plant_techno_plan_targets(plant, report_month)
            targets[plant] = {k[1]: v for k, v in raw.items()}
        sail_targets_raw = pt.compute_sail_targets(fy)
        targets["SAIL"] = {k[1]: v for k, v in sail_targets_raw.items()}
        for plant in list(targets):
            targets[plant] = bnc.fuel_rate_fallback(targets[plant])

        t19 = tables[TABLE_TECHNO]
        zz = str(fy_start + 1)[2:]
        for r in (0, 1):
            bdu.set_cell(t19.rows[r].cells[3], q3_label)
            bdu.set_cell(t19.rows[r].cells[4], q3_cply_label)
            bdu.set_cell(t19.rows[r].cells[5], l_label)
            bdu.set_cell(t19.rows[r].cells[6], l_cply_label)
        bdu.set_cell(t19.rows[0].cells[2], f"Target\n{yy}-{zz}")
        bdu.set_cell(t19.rows[1].cells[2], f"{yy}-{zz}")

        ri = 2
        for param in TECHNO_PARAMS:
            sec = techno_by_param[param]
            row_by_plant = {r["plant"]: r["values"] for r in sec["rows"]}
            for plant in TECHNO_PLANTS:
                cells = t19.rows[ri].cells
                vals = row_by_plant.get(plant, {})

                def disp(label):
                    d = vals.get(label)
                    return d["display"] if d else ""

                tgt = _fmt_techno_target(param, (targets.get(plant) or {}).get(param))
                bdu.set_cell(cells[2], tgt if tgt is not None else "")
                bdu.set_cell(cells[3], disp("Q3cur"))
                bdu.set_cell(cells[4], disp("Q3cply"))
                bdu.set_cell(cells[5], disp("NMcur"))
                bdu.set_cell(cells[6], disp("NMcply"))
                ri += 1

        bdu.set_para(
            para(P_TECHNO_SECTION_TITLE),
            f"SAIL and Plant-wise Major Techno-economic parameters performance during {q3_label} and {l_label}:"
        )
        bdu.set_para(para(P_TECHNO_PERF_HEADING), f"Techno-economic Performance during {q3_label}:")
        bdu.set_para(para(P_TECHNO_COKE),
                     f"Coke Rate: SAIL registered {bnc.imp_phrase(_improvement(techno_by_param, 'Coke Rate'))} w.r.t. CPLY.")
        bdu.set_para(para(P_TECHNO_FUEL),
                     f"Fuel Rate: SAIL registered {bnc.imp_phrase(_improvement(techno_by_param, 'Fuel Rate'))} w.r.t. CPLY.")
        bdu.set_para(para(P_TECHNO_BFP),
                     f"BF Productivity: SAIL registered "
                     f"{bnc.imp_phrase(_improvement(techno_by_param, 'BF Productivity', higher_is_better=True))} "
                     f"w.r.t. CPLY.")
        for idx in TECHNO_STALE_REMOVE:
            bdu.remove_paragraph(para(idx))

        # ---- annexure index (table 20) -----------------------------------------
        t20 = tables[TABLE_ANNEXURE_INDEX]
        m1, m2, m3 = (bq2._month_name_title(m) for m in Q3_CUR)
        bdu.set_cell(t20.rows[0].cells[3],
                     f"SAIL: Production performance during {m1}’{yy}, {m2}’{yy} and {m3}’{yy}.")
        bdu.set_cell(t20.rows[1].cells[3],
                     f"SAIL: Production performance during {q3_label} and {l_label}")
        bdu.set_cell(t20.rows[2].cells[3],
                     f"SAIL: Techno-economic parameters during {q3_label} and {l_label}, best ever achieved, "
                     f"norms and benchmarks")
        bdu.set_cell(t20.rows[3].cells[3],
                     f"SAIL: Production performance :New Facilities during \n{q3_label} and {l_label}")
        bdu.set_cell(t20.rows[4].cells[3],
                     f"SAIL: Techno-economic parameters for New Blast Furnaces during {q3_label} and {l_label}")
        bdu.set_cell(t20.rows[5].cells[3],
                     f"SAIL: Product mix performance during {q3_label} and {l_label}")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    finally:
        conn.close()
