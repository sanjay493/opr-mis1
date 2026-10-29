"""
"Cost of Production" -- standalone report (cost_production.py): Rs/T cost of
production of BF Coke, Sinter, Hot Metal, Crude Steel and Saleable Steel,
all plants, monthly or yearly. Not part of the monthly PDF.

  GET /api/cost-production?view=monthly|yearly&fy=YYYY-YY&products=HM,CS -> JSON
  GET /api/cost-production/xlsx?...same...                             -> .xlsx
  GET /api/cost-production/pdf?...same...                              -> .pdf
"""
import io
import re
from typing import Optional

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response, StreamingResponse

import cost_production
import cost_production_pdf
from colors_loader import load_colors_config

router = APIRouter(prefix="/api/cost-production", tags=["cost-production"])

_CODES = {code for code, _ in cost_production.PRODUCTS}


def _data(view: str, fy: Optional[str], products: Optional[str]) -> dict:
    if view not in ("monthly", "yearly"):
        raise HTTPException(status_code=400, detail="view must be monthly or yearly")
    if fy is not None and not re.fullmatch(r"\d{4}-\d{2}", fy):
        raise HTTPException(status_code=400, detail="fy must be YYYY-YY")
    codes = [c for c in (products or "").split(",") if c]
    if any(c not in _CODES for c in codes):
        raise HTTPException(status_code=400, detail=f"products must be among {', '.join(sorted(_CODES))}")
    return cost_production.generate(view, fy, codes or None)


def _fname(data: dict, ext: str) -> str:
    return f"Cost_of_Production_{data['fy'] if data['view'] == 'monthly' else 'Yearly'}.{ext}"


@router.get("")
def cost_production_json(view: str = Query("monthly"), fy: Optional[str] = Query(None),
                         products: Optional[str] = Query(None)):
    return _data(view, fy, products)


def write_sheet(ws, data: dict):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    periods = data["periods"]
    ncols = len(periods) + 2
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="DBEAFE")
    blk_fill = PatternFill("solid", fgColor="E0E7FF")
    sail_fill = PatternFill("solid", fgColor="FEF3C7")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(1, 1, data["title"]).font = Font(bold=True, size=13, color="1E3A8A")
    ws.cell(2, ncols, f"Unit: {data['unit']}").alignment = Alignment(horizontal="right")
    r = 4
    for prod in data["products"]:
        ws.cell(r, 1, prod["name"]).font = Font(bold=True, size=12, color="0000FF")
        r += 1
        for c, text in enumerate(["Cost", "Plant"] + [p["label"] for p in periods], start=1):
            cell = ws.cell(r, c, text)
            cell.font, cell.fill, cell.alignment, cell.border = bold, hdr_fill, center, border
        r += 1
        for b in prod["blocks"]:
            start = r
            for row in b["rows"]:
                sail = row["plant"].startswith("SAIL")
                ws.cell(r, 2, row["plant"])
                for i, v in enumerate(row["values"]):
                    ws.cell(r, 3 + i, None if v is None else round(v, 2)).number_format = "#,##0"
                for c in range(1, ncols + 1):
                    cell = ws.cell(r, c)
                    cell.border = border
                    if sail and c >= 2:
                        cell.font, cell.fill = bold, sail_fill
                r += 1
            ws.cell(start, 1, b["label"])
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
            ws.cell(start, 1).font, ws.cell(start, 1).fill = bold, blk_fill
            ws.cell(start, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if prod.get("note"):
            ws.cell(r, 1, prod["note"]).font = Font(italic=True, size=9)
            r += 1
        r += 1
    for c, w in enumerate([14, 14] + [12] * len(periods), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


@router.get("/xlsx")
def cost_production_xlsx(view: str = Query("monthly"), fy: Optional[str] = Query(None),
                         products: Optional[str] = Query(None)):
    from openpyxl import Workbook

    data = _data(view, fy, products)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost of Production"
    write_sheet(ws, data)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_fname(data, "xlsx")}"'},
    )


@router.get("/pdf")
def cost_production_pdf_download(view: str = Query("monthly"), fy: Optional[str] = Query(None),
                                 products: Optional[str] = Query(None)):
    # A plain def, so FastAPI runs it in its threadpool: the DB queries and
    # the synchronous Playwright render must not block the event loop.
    data = _data(view, fy, products)
    content = cost_production_pdf.render_pdf(data, load_colors_config())
    return Response(content, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_fname(data, "pdf")}"'})
