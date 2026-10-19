"""Freeze what the app shows, so a schema change can be proven invisible.

    python -m tools.parity_snapshot --db mis_reports_baseline --out <dir> \
        [--months 2024-04..2026-09,2016-03,2021-03] [--pdf-months 2026-08,2026-09] \
        [--today 2026-10-19]

Runs the FastAPI app in-process (fastapi.testclient.TestClient) against the
MySQL schema `db_name` - never the live `mis_reports` - and writes, for every
month, the JSON of every month-relevant GET /api/ route plus xlsx/pptx/docx
exports, and for --pdf-months the full report PDF text (pdfplumber, per page).
Compare two snapshots with `python -m tools.parity_diff <a> <b>`.

Layout of out_dir:
    <YYYY-MM>/<endpoint-slug>.json        month-scoped endpoints
    <YYYY-MM>/report.pdf.txt              full report PDF text (pdf months only)
    _global/<endpoint-slug>.json          endpoints that do not depend on the month
                                          (FY-keyed ones once per FY, value in the slug)

--today YYYY-MM-DD pins the app's notion of "today" for the whole run. The
baseline snapshot was taken on 2026-10-19, so every parity run passes
--today 2026-10-19. Date-dependent output reached by snapshotted endpoints
(the capital-repair calendar's "N days so far" counters, the techno records'
last month, the default month of some GETs) then cannot drift when a run
crosses midnight. The patch is per module (TODAY_PINNED_MODULES):

    page_capital_repair.date          generate_capital_repair_calendar's today
    page_techno_records.date          generate_techno_records' last month
    main._dt                          _latest_report_month, _radar_default_month
    api_production_loss.date          production-loss-analysis default today
    production_loss_analysis.date     build_report's default today ...
    production_loss_analysis.datetime ... and its "now"

today()/now() return real date/datetime objects for the pinned day; now()
keeps the real time of day (and honours a tz argument). Constructors,
classmethods and isinstance behave as the real classes. Timestamps only
written to logs, or to the DB by write routes (never snapshotted), are left
alone: api_admin / api_auth / auth, cover_store, db.py writers, todo /
worklog, secretary_review_text saves, page_techno's calculated_at (POST),
the PDF job TTL (time.time()) and the upload-only extractors.

Auth: none needed. Only /api/admin/* and /api/auth/me read the session cookie,
and both are excluded; every other GET is open (the editor/admin gate in
main.py only covers write routes). DB selection: MYSQL_DB is set before
dbengine is imported, and the live connection is checked with SELECT DATABASE().
"""

import argparse
import hashlib
import io
import json
import os
import re
import sys
import time
from contextlib import contextmanager
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

LIVE_DB = "mis_reports"
DEFAULT_MONTHS = "2024-04..2026-09,2016-03,2021-03"
DEFAULT_PDF_MONTHS = "2026-08,2026-09"

PLANTS = ["BSP", "DSP", "RSP", "BSL", "ISP", "ASP", "SSP", "VISL"]
COST_TREND_PRODUCTS = ["HM", "CS", "SS", "COKE", "SINTER"]

# Routes never snapshotted, by prefix: auth/admin, write-adjacent and job plumbing.
EXCLUDED_PREFIXES = ("/api/admin", "/api/auth")

# GET /api/ routes that are not snapshotted -> why. Printed at the end of every
# run. Anything not listed here is driven (some by an explicit param map below).
SKIPPED_ENDPOINTS = {
    "/api/admin/*": "admin routes need an admin session; not report data",
    "/api/auth/me": "session introspection; not report data",
    "/api/generate-pdf/status/{job_id}": "PDF job plumbing; the job is driven by the report.pdf step",
    "/api/generate-pdf/result/{job_id}": "PDF job plumbing; the job is driven by the report.pdf step",
    "/api/todo/list": "operator to-do notes, not report data",
    "/api/worklog/list": "operator work log, not report data",
    "/api/cover/photos/{photo_id}": "binary image keyed by an uploaded photo id",
    "/api/breakdown/cr-candidates": "needs a breakdown start_ts picked from an existing entry",
    "/api/production-item-range": "needs a plant item name picked per plant",
    "/api/production-loss-analysis": "needs two comparison periods plus a plant item name",
    "/api/special-steel/grades": "needs a plant+product pair picked from /api/special-steel/products",
    "/api/techno-monthly-data": "needs a group_code from /api/techno-groups, which cannot list any (techno_param_group missing)",
    "/api/techno-data": "needs a free-text parameter list; the same figures are served by /api/techno/data",
    "/api/techno-page-targets": "needs a fractional report page number",
    "/api/techno/manual/cumulative-preview": "needs a param_key and live input values",
    "/api/techno/manual/param-history": "needs a unit and param_key",
    "/api/secretary-review/default-text": "needs a block id; covered by /api/secretary-review/context",
    "/api/production-query/excel": "POST; needs a user-chosen plant/unit list; the data is served by the GET production endpoints",
    "/api/techno-table-export/excel": "POST echoes a table the browser built; reads nothing from the DB",
    "/api/bd-cr-log/excel": "POST echoes a log the browser built; reads nothing from the DB",
    "/api/upload-excel, /api/upload-excel-plan, /api/rake-detention/extract-pdf": "POST uploads: write to the DB",
    "POST */pdf, /api/generate-pdf*": "POST PDF renders; the full report PDF is covered by report.pdf.txt for the pdf months",
    "*/pdf": "PDF-only twin of a JSON/xlsx endpoint (data covered there); rendering is covered by report.pdf",
    "/api/board-note/annexures (fmt=pdf)": "the xlsx variant is snapshotted; rendering is covered by report.pdf",
}

# Endpoints whose only difference from a sibling is the file format.
_PDF_PATH_SUFFIX = "/pdf"

# Optional query parameters that are always sent, in addition to required ones.
_PLANT_PARAMS = {"plant", "plant_code"}
_MONTH_PARAMS = {"month", "report_month", "stock_month", "end_mon", "end"}
_FY_START_MONTH_PARAMS = {"start_mon", "start"}
_FY_LABEL_PARAMS = {"fy", "financial_year"}
_FY_YEAR_PARAMS = {"fy_start"}
_QUARTER_PARAMS = {"quarter"}
_DYNAMIC_PARAMS = {"mine_code", "bf_id"}

# Optional params pinned to a single non-default value for specific routes.
# A value may be a callable(month) -> dict for month-dependent pins.
FIXED_PARAMS = {
    # start_mon/end_mon are calendar month numbers read within the Apr->Mar
    # financial year (page_records.generate_best_period): April up to the
    # report month is always a valid window, including Jan-Mar.
    "/api/best-period-records": lambda m: {"start_mon": 4, "end_mon": int(m[5:7])},
    "/api/board-note/annexures": {"fmt": "xlsx", "annex": "all"},
    "/api/production-query-items": {"plants": "BSP,DSP,RSP,BSL,ISP"},
}
_MONTH_FIXED = {"/api/best-period-records"}

# POST-only exports that only read the DB (checked: no INSERT/UPDATE/DELETE or
# commit on their render path; secretary_review_text save/unlink are separate
# routes). Each is driven once per month with a payload built from the month.
# (path, payload(month, ctx) -> dict)
POST_EXPORTS = [
    ("/api/secretary-review/pptx", lambda m, c: {"month": m}),
    ("/api/techno-custom-standard/excel", lambda m, c: {"month": m}),
    ("/api/techno-custom-period/excel",
     lambda m, c: {"plants": ["BSP", "DSP", "RSP", "BSL", "ISP", "SAIL"],
                   "periods": [{"label": m, "months": [m]}]}),
    ("/api/bf-benchmark/excel", lambda m, c: {"sail_bf_keys": c["sail_bf_keys"], "months": [m]}),
    ("/api/techno-bf-furnace/excel",
     lambda m, c: {"furnaces": c["furnace_keys"], "mode": "range", "start_month": m, "end_month": m}),
]

# Report-PDF polling gives up after this long.
PDF_TIMEOUT_SECONDS = 30 * 60

# FY-keyed routes that take about a minute per call (Word/annexure builds): run
# only for the FYs of the PDF months and the FY before, not for every FY seen.
SLOW_FY_ENDPOINTS = {"/api/board-note/docx", "/api/board-note/annexures"}


# ---------------------------------------------------------------------------
# --today: pin the app's date
# ---------------------------------------------------------------------------

# (module, attributes to replace). An attribute named date / datetime is that
# class; any other (main's `_dt`) is the datetime module itself.
TODAY_PINNED_MODULES = [
    ("page_capital_repair", ("date",)),
    ("page_techno_records", ("date",)),
    ("main", ("_dt",)),
    ("api_production_loss", ("date",)),
    ("production_loss_analysis", ("date", "datetime")),
]


def _pinned_classes(day):
    """(date, datetime) stand-ins whose today()/now() report `day`. Calls,
    classmethods and isinstance behave as the real classes, and every value
    they hand out is a real date / datetime."""
    import datetime as real

    class _Meta(type):
        def __instancecheck__(cls, obj):
            return isinstance(obj, cls._real)

        def __subclasscheck__(cls, sub):
            return issubclass(sub, cls._real)

    class PinnedDate(real.date, metaclass=_Meta):
        _real = real.date

        def __new__(cls, *a, **k):
            return real.date(*a, **k)

        @classmethod
        def today(cls):
            return real.date(day.year, day.month, day.day)

    class PinnedDatetime(real.datetime, metaclass=_Meta):
        _real = real.datetime

        def __new__(cls, *a, **k):
            return real.datetime(*a, **k)

        @classmethod
        def now(cls, tz=None):
            return real.datetime.now(tz).replace(year=day.year, month=day.month, day=day.day)

        @classmethod
        def today(cls):
            return cls.now()

    return PinnedDate, PinnedDatetime


def _pinned_datetime_module(day):
    """A copy of the datetime module with date / datetime pinned."""
    import datetime as real
    import types
    mod = types.ModuleType("datetime")
    mod.__dict__.update({k: v for k, v in vars(real).items() if not k.startswith("__")})
    mod.date, mod.datetime = _pinned_classes(day)
    return mod


@contextmanager
def pin_today(day):
    """Pin date.today() / datetime.now() / datetime.today() to `day` (a
    datetime.date) in every TODAY_PINNED_MODULES module; the real objects are
    put back on exit. `day` None is a no-op."""
    if day is None:
        yield
        return
    import importlib
    pdate, pdatetime = _pinned_classes(day)
    saved = []
    try:
        for name, attrs in TODAY_PINNED_MODULES:
            mod = importlib.import_module(name)
            for attr in attrs:
                new = (pdate if attr == "date" else pdatetime if attr == "datetime"
                       else _pinned_datetime_module(day))
                saved.append((mod, attr, getattr(mod, attr)))
                setattr(mod, attr, new)
        yield
    finally:
        for mod, attr, orig in reversed(saved):
            setattr(mod, attr, orig)


# ---------------------------------------------------------------------------
# months / financial years
# ---------------------------------------------------------------------------

def parse_months(spec: str) -> list:
    """'2024-04..2026-09,2016-03' -> sorted unique ['2016-03', '2024-04', ...]."""
    out = set()
    for part in filter(None, (s.strip() for s in spec.split(","))):
        if ".." in part:
            a, b = part.split("..")
            y, m = int(a[:4]), int(a[5:7])
            ey, em = int(b[:4]), int(b[5:7])
            while (y, m) <= (ey, em):
                out.add(f"{y:04d}-{m:02d}")
                y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        else:
            out.add(part)
    return sorted(out)


def fy_start_year(month: str) -> int:
    y, m = int(month[:4]), int(month[5:7])
    return y if m >= 4 else y - 1


def fy_label(start_year: int) -> str:
    return f"{start_year}-{(start_year + 1) % 100:02d}"


def quarter_of(month: str) -> int:
    return ((int(month[5:7]) - 4) % 12) // 3 + 1


# ---------------------------------------------------------------------------
# response -> JSON-able snapshot
# ---------------------------------------------------------------------------

def _xlsx_to_obj(content: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    out = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells[c.coordinate] = {"v": c.value if isinstance(c.value, (int, float, str, bool)) else str(c.value),
                                           "fmt": c.number_format}
        out[ws.title] = {"dims": ws.dimensions, "merged": sorted(str(r) for r in ws.merged_cells.ranges), "cells": cells}
    return out


def _pptx_to_obj(content: bytes) -> list:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = []
        for sh in slide.shapes:
            entry = {"name": sh.name}
            if sh.has_text_frame:
                entry["text"] = sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                entry["table"] = [[c.text for c in r.cells] for r in sh.table.rows]
            shapes.append(entry)
        slides.append({"slide": i, "shapes": shapes})
    return slides


def _docx_to_obj(content: bytes) -> dict:
    import docx
    d = docx.Document(io.BytesIO(content))
    return {
        "paragraphs": [p.text for p in d.paragraphs],
        "tables": [[[c.text for c in r.cells] for r in t.rows] for t in d.tables],
    }


def _pdf_text(content: bytes) -> str:
    import pdfplumber
    parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            parts.append(f"=== page {i} ===\n{page.extract_text() or ''}")
    return "\n".join(parts) + "\n"


def _normalise(resp):
    """-> (suffix, text) for a response. JSON is key-sorted; numbers untouched."""
    ctype = resp.headers.get("content-type", "").split(";")[0].strip().lower()
    if resp.status_code >= 400:
        try:
            detail = resp.json()
        except ValueError:
            detail = resp.text[:500]
        return ".json", _dump({"_status": resp.status_code, "_detail": detail})
    if ctype == "application/json":
        return ".json", _dump(resp.json())
    if "spreadsheetml" in ctype:
        return ".xlsx.json", _dump(_xlsx_to_obj(resp.content))
    if "presentationml" in ctype:
        return ".pptx.json", _dump(_pptx_to_obj(resp.content))
    if "wordprocessingml" in ctype:
        return ".docx.json", _dump(_docx_to_obj(resp.content))
    if ctype == "application/pdf":
        return ".pdf.txt", _pdf_text(resp.content)
    if ctype.startswith("text/"):
        return ".txt", resp.text
    # Unknown binary (e.g. an image): record type and a content hash.
    return ".json", _dump({"_content_type": ctype, "_bytes": len(resp.content),
                           "_sha256": hashlib.sha256(resp.content).hexdigest()})


def _dump(obj) -> str:
    return json.dumps(obj, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str)


# ---------------------------------------------------------------------------
# endpoint enumeration and param expansion
# ---------------------------------------------------------------------------

def enumerate_get_endpoints(app) -> list:
    """[(path, [(param_name, location, required)])] for every GET /api/ route.
    Read from the OpenAPI schema so routers mounted lazily are included."""
    out = []
    for path, item in sorted(app.openapi()["paths"].items()):
        get = item.get("get")
        if not get or not path.startswith("/api/"):
            continue
        params = [(p["name"], p["in"], bool(p.get("required"))) for p in get.get("parameters", [])
                  if p["in"] in ("query", "path")]
        out.append((path, params))
    return out


def _is_skipped(path: str) -> bool:
    return (path.startswith(EXCLUDED_PREFIXES) or path in SKIPPED_ENDPOINTS
            or path.endswith(_PDF_PATH_SUFFIX) or path.startswith("/api/generate-pdf/"))


def _dynamic_values(client) -> dict:
    """Values for params that name a row (a mine, an external BF)."""
    def ids(path, list_key, field):
        try:
            rows = client.get(path).json().get(list_key, [])
        except (ValueError, AttributeError):
            return []
        return sorted({str(r[field]) for r in rows if isinstance(r, dict) and r.get(field) not in (None, "")})
    return {
        "mine_code": ids("/api/mines-master", "mines", "mine_code"),
        "bf_id": ids("/api/bf-benchmark/external-bfs", "external_bfs", "id"),
    }


def expand_calls(path, params, month, fys, dynamic):
    """Param dicts to call `path` with. Returns (scope, [dict]) where scope is
    'month', 'fy' or 'global'; None when the endpoint needs a value we cannot make."""
    names = {n for n, _, _ in params}
    scope = "month" if names & (_MONTH_PARAMS | _FY_START_MONTH_PARAMS) else \
        "fy" if names & (_FY_LABEL_PARAMS | _FY_YEAR_PARAMS | _QUARTER_PARAMS) else "global"
    fixed = FIXED_PARAMS.get(path, {})
    fixed = fixed(month) if callable(fixed) else fixed
    if path in _MONTH_FIXED:
        scope = "month"
    variants = [dict(fixed)]
    for name, loc, required in params:
        if name in fixed:
            continue
        if name in _MONTH_PARAMS:
            choices = [month]
        elif name in _FY_START_MONTH_PARAMS:
            choices = [f"{fy_start_year(month)}-04"]
        elif name in _PLANT_PARAMS:
            choices = PLANTS if required else [None] + PLANTS
        elif name in _FY_LABEL_PARAMS and scope == "month":
            choices = [fy_label(fy_start_year(month))]
        elif name in _FY_YEAR_PARAMS and scope == "month":
            choices = [fy_start_year(month)]
        elif name in _FY_LABEL_PARAMS:
            choices = [fy_label(y) for y in fys]
        elif name in _FY_YEAR_PARAMS:
            choices = list(fys)
        elif name in _QUARTER_PARAMS:
            choices = [1, 2, 3, 4]
        elif name == "product" and required:
            choices = COST_TREND_PRODUCTS
        elif name in _DYNAMIC_PARAMS and required:
            choices = dynamic.get(name) or []
            if not choices:
                return scope, None
        elif required:
            return scope, None
        else:
            continue
        variants = [{**v, name: c} for v in variants for c in choices]
    return scope, variants


def slug(path: str, call: dict, drop=()) -> str:
    base = "api_" + path[len("/api/"):].replace("/", "__").replace("{", "").replace("}", "")
    extra = "".join(f"--{k}={v}" for k, v in sorted(call.items()) if k not in drop and v is not None)
    s = re.sub(r"[^A-Za-z0-9_.=,+-]", "_", base + extra)
    return s[:180] + ("-" + hashlib.sha1(s.encode()).hexdigest()[:8] if len(s) > 180 else "")


def _split_path_params(path, params, call):
    url = path
    q = {}
    for k, v in call.items():
        if v is None:
            continue
        if "{" + k + "}" in url:
            url = url.replace("{" + k + "}", str(v))
        else:
            q[k] = v
    return url, q


# ---------------------------------------------------------------------------
# the DB guard
# ---------------------------------------------------------------------------

def _select_database(db_name: str):
    if db_name == LIVE_DB:
        raise SystemExit(f"refusing to snapshot the live database {LIVE_DB!r}; "
                         f"use a restored copy such as mis_reports_baseline")
    os.environ["MYSQL_DB"] = db_name
    import dbengine
    if dbengine._MYSQL_CFG["database"] != db_name:
        dbengine._MYSQL_CFG["database"] = db_name
        dbengine._pool_clear()
    import db
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SELECT DATABASE()")
        actual = cur.fetchone()[0]
    finally:
        conn.close()
    if actual != db_name or actual == LIVE_DB:
        raise SystemExit(f"connected to {actual!r}, expected {db_name!r}")
    return actual


# ---------------------------------------------------------------------------
# snapshot
# ---------------------------------------------------------------------------

def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def _full_report_pdf(client, month: str) -> str:
    """The way the report page does it: /api/data pages -> /api/generate-pdf/start,
    poll status, fetch result. Returns the PDF text."""
    body = client.get("/api/data", params={"month": month}).json()
    pages = body["pages"] if isinstance(body, dict) and "pages" in body else body
    start = client.post("/api/generate-pdf/start", json={"month": month, "pages": pages, "full_export": True})
    start.raise_for_status()
    job = start.json()["job_id"]
    deadline = time.time() + PDF_TIMEOUT_SECONDS
    while True:
        if time.time() > deadline:
            raise TimeoutError(f"report PDF for {month} not finished after {PDF_TIMEOUT_SECONDS // 60} minutes (job {job})")
        time.sleep(4)
        st = client.get(f"/api/generate-pdf/status/{job}").json()
        if st["status"] == "error":
            raise RuntimeError(f"PDF job failed for {month}: {st.get('error')}")
        if st["status"] == "done":
            break
    res = client.get(f"/api/generate-pdf/result/{job}")
    res.raise_for_status()
    return _pdf_text(res.content)


_FILE_RE = re.compile(r"(\.(xlsx|pptx|docx|pdf|html))?\.(json|txt)$")


def summarise(out_dir: Path) -> dict:
    """Status counts and always-failing endpoints, read from the files on disk
    (so a partial regeneration still yields the summary of the whole snapshot).
    Written to out_dir/_summary.json; parity_diff ignores that file."""
    counts = {}
    per_endpoint = {}   # endpoint -> {variant: [statuses]}
    for f in sorted(Path(out_dir).rglob("*")):
        if not f.is_file() or f.name in ("_summary.json", "report.pdf.txt"):
            continue
        data = f.read_bytes()
        status = 200
        if data.startswith(b'{"_detail"'):
            m = re.search(rb'"_status":(\d+)', data[-400:])
            status = int(m.group(1)) if m else 200
        counts[str(status)] = counts.get(str(status), 0) + 1
        stem = _FILE_RE.sub("", f.name)
        per_endpoint.setdefault(stem.split("--")[0], {}).setdefault(stem, []).append(status)
    always_endpoints = {}
    always_variants = {}
    for ep, variants in per_endpoint.items():
        allst = [s for v in variants.values() for s in v]
        if all(s >= 400 for s in allst):
            always_endpoints[ep] = sorted(set(allst))
            continue
        for var, sts in variants.items():
            if all(s >= 400 for s in sts):
                always_variants[var] = sorted(set(sts))
    summary = {"status_counts": dict(sorted(counts.items())),
               "endpoints_failing_every_time": dict(sorted(always_endpoints.items())),
               "variants_failing_every_time": dict(sorted(always_variants.items()))}
    (Path(out_dir) / "_summary.json").write_text(json.dumps(summary, indent=1, sort_keys=True), encoding="utf-8")
    return summary


def snapshot(months: list, out_dir: Path, db_name: str, pdf_months=(), only=None, today=None) -> None:
    """`only`: restrict to these endpoint paths (GET or POST export), e.g. to
    regenerate a few files in an existing snapshot. `today`: a datetime.date
    pinned as the app's today for the whole run (see pin_today)."""
    out_dir = Path(out_dir)
    _select_database(db_name)
    with pin_today(today):
        _snapshot(months, out_dir, db_name, pdf_months, only)


def _snapshot(months, out_dir, db_name, pdf_months, only) -> None:
    from fastapi.testclient import TestClient
    import main

    t0 = time.time()
    endpoints = [(p, ps) for p, ps in enumerate_get_endpoints(main.app) if not _is_skipped(p) and (not only or p in only)]
    fys = sorted({fy_start_year(m) for m in months} | {fy_start_year(m) - 1 for m in months})
    slow_base = list(pdf_months) or [max(months)]
    slow_fys = sorted({fy_start_year(m) for m in slow_base} | {fy_start_year(m) - 1 for m in slow_base})
    written = 0
    calls = 0
    undriven = {}
    timings = {}
    with TestClient(main.app, raise_server_exceptions=False) as client:
        dynamic = _dynamic_values(client)
        ctx = {
            "sail_bf_keys": [f"{r['plant']}:{r['unit']}" for r in client.get("/api/bf-benchmark/sail-meta").json().get("sail_meta", [])],
            "furnace_keys": client.get("/api/techno-bf-furnace/meta").json().get("furnace_keys", []),
        }
        done_global = set()
        for month in months:
            for path, params in endpoints:
                scope, variants = expand_calls(path, params, month, slow_fys if path in SLOW_FY_ENDPOINTS else fys, dynamic)
                if variants is None:
                    undriven[path] = "required parameter has no value to drive it"
                    continue
                if scope != "month" and path in done_global:
                    continue
                if scope != "month":
                    done_global.add(path)
                target = out_dir / (month if scope == "month" else "_global")
                for call in variants:
                    url, q = _split_path_params(path, params, call)
                    t1 = time.time()
                    resp = client.get(url, params=q)
                    timings[path] = timings.get(path, 0.0) + time.time() - t1
                    calls += 1
                    suffix, text = _normalise(resp)
                    drop = {n for n in call if n in _MONTH_PARAMS | _FY_START_MONTH_PARAMS} if scope == "month" else ()
                    _write(target / (slug(path, call, drop) + suffix), text)
                    written += 1
        for month in months:
            for path, make in POST_EXPORTS:
                if only and path not in only:
                    continue
                resp = client.post(path, json=make(month, ctx))
                calls += 1
                suffix, text = _normalise(resp)
                _write(out_dir / month / (slug(path, {}) + suffix), text)
                written += 1
        for month in pdf_months:
            _write(out_dir / month / "report.pdf.txt", _full_report_pdf(client, month))
            written += 1

    print(f"snapshot of {db_name}: {len(months)} month(s), {calls} calls, {written} files, "
          f"{time.time() - t0:.0f}s -> {out_dir}")
    print("slowest endpoints (total seconds):")
    for p, secs in sorted(timings.items(), key=lambda kv: -kv[1])[:8]:
        print(f"  {secs:7.1f}s  {p}")
    summary = summarise(out_dir)
    print("status counts (all files in snapshot): " + ", ".join(f"{k}: {v}" for k, v in summary["status_counts"].items()))
    print("endpoints returning 4xx/5xx every time:")
    for ep, sts in summary["endpoints_failing_every_time"].items():
        print(f"  {ep}: {sts}")
    print("variants (plant etc.) returning 4xx/5xx every time:")
    for ep, sts in summary["variants_failing_every_time"].items():
        print(f"  {ep}: {sts}")
    skipped = dict(SKIPPED_ENDPOINTS)
    skipped.update({p: r for p, r in undriven.items() if p not in skipped})
    print("SKIPPED_ENDPOINTS:")
    for p, r in sorted(skipped.items()):
        print(f"  {p}: {r}")


def main_cli(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--db", required=True, help="MySQL schema to snapshot (never mis_reports)")
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--months", default=DEFAULT_MONTHS)
    ap.add_argument("--pdf-months", default=DEFAULT_PDF_MONTHS, help="months that also get the full report PDF text ('' for none)")
    ap.add_argument("--only", default="", help="comma-separated endpoint paths to (re)generate; other files are left alone")
    ap.add_argument("--today", default=None,
                    help="YYYY-MM-DD pinned as the app's today (parity runs: 2026-10-19)")
    a = ap.parse_args(argv)
    months = parse_months(a.months)
    pdf_months = parse_months(a.pdf_months)
    only = {x.strip() for x in a.only.split(",") if x.strip()} or None
    import datetime
    today = datetime.date.fromisoformat(a.today) if a.today else None
    snapshot(months, a.out, a.db, pdf_months=pdf_months, only=only, today=today)
    return 0


if __name__ == "__main__":
    sys.exit(main_cli())
