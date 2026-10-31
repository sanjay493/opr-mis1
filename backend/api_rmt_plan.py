"""
RMT "Production Plan for <month>" report (rmt_plan.py).

  GET  /api/rmt-plan?month=YYYY-MM        -> the report as JSON
  GET  /api/rmt-plan/form?month=YYYY-MM   -> the manual rows for the entry form
  POST /api/rmt-plan                      -> save the month's manual rows (editor)
  GET  /api/rmt-plan/xlsx?month=YYYY-MM   -> .xlsx (three sheets, as the PDF)
  GET  /api/rmt-plan/pdf?month=YYYY-MM    -> .pdf (A4 landscape, three pages)

The POST is gated by EditorAdminGateMiddleware via
constants.PAGE_MODULES["rmt_plan"].
"""
import io
import re
from datetime import date
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel

import rmt_plan
import rmt_plan_pdf
from auth import get_current_user_optional
from colors_loader import load_colors_config
from repo import misc as rmisc

router = APIRouter(prefix="/api/rmt-plan", tags=["rmt-plan"])


def _check_month(month: str):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month or ""):
        raise HTTPException(400, "month must be YYYY-MM, e.g. '2026-11'")


class RmtNote(BaseModel):
    section: str
    row_label: str
    remarks: Optional[str] = None
    coal_req_app: Optional[float] = None
    coal_req_op: Optional[float] = None


class RmtIptRow(BaseModel):
    item: str
    from_plant: Optional[str] = None
    to_plant: Optional[str] = None
    availability: Optional[str] = None
    requirement: Optional[str] = None
    remarks: Optional[str] = None


class RmtStockRow(BaseModel):
    plant: str
    bf_coke: Optional[str] = None
    slab: Optional[str] = None
    bloom: Optional[str] = None


class RmtSaveRequest(BaseModel):
    month: str
    stock_as_on: Optional[str] = None      # YYYY-MM-DD
    note: Optional[str] = None
    notes: List[RmtNote] = []
    ipt: List[RmtIptRow] = []
    stock: List[RmtStockRow] = []


def _user_email(request: Request) -> Optional[str]:
    try:
        from auth import COOKIE_NAME
        user = get_current_user_optional(request.cookies.get(COOKIE_NAME))
        return user["email"] if user else None
    except Exception:
        return None


def _clean(v: Optional[str]) -> Optional[str]:
    """Trim; blank -> None. Line breaks inside are kept."""
    v = (v or "").strip()
    return v or None


@router.get("")
def rmt_plan_json(month: str = Query(...)):
    _check_month(month)
    return rmt_plan.generate(month)


@router.get("/form")
def rmt_plan_form(month: str = Query(...)):
    _check_month(month)
    return {**rmt_plan.manual_form(month),
            "sections": [{"key": k, "label": label, "rows": rows, "coal": coal}
                         for k, label, rows, coal in rmt_plan.form_rows()]}


@router.post("")
def rmt_plan_save(body: RmtSaveRequest, request: Request):
    _check_month(body.month)
    stock_as_on = _clean(body.stock_as_on)
    if stock_as_on:
        try:
            date.fromisoformat(stock_as_on)
        except ValueError:
            raise HTTPException(400, "stock_as_on must be a date, YYYY-MM-DD")
    allowed = {(k, r) for k, _, rows, _ in rmt_plan.form_rows() for r in rows}
    notes = []
    for n in body.notes:
        if (n.section, n.row_label) not in allowed:
            raise HTTPException(400, f"Unknown row {n.section} / {n.row_label}")
        for v in (n.coal_req_app, n.coal_req_op):
            if v is not None and (v != v or v < 0):
                raise HTTPException(400, f"Coal requirement for {n.row_label} must be a number ≥ 0")
        d = {"section": n.section, "row_label": n.row_label, "remarks": _clean(n.remarks),
             "coal_req_app": n.coal_req_app, "coal_req_op": n.coal_req_op}
        if d["remarks"] or d["coal_req_app"] is not None or d["coal_req_op"] is not None:
            notes.append(d)
    ipt = []
    for r in body.ipt:
        d = {k: _clean(getattr(r, k)) for k in ("item", "from_plant", "to_plant", "availability", "requirement",
                                                  "remarks")}
        if not any(d.values()):
            continue
        if not d["item"]:
            raise HTTPException(400, "Every IPT row needs an item (Sinter, BF Coke, CC Slabs, …)")
        ipt.append(d)
    stock = []
    for r in body.stock:
        d = {k: _clean(getattr(r, k)) for k in ("plant", "bf_coke", "slab", "bloom")}
        if not any(d.values()):
            continue
        if not d["plant"]:
            raise HTTPException(400, "Every stock row needs a plant")
        stock.append(d)
    rmisc.save_rmt_plan_manual(body.month, stock_as_on, _clean(body.note), notes, ipt, stock, _user_email(request))
    return {"status": "ok", "notes": len(notes), "ipt": len(ipt), "stock": len(stock)}


# ── Excel ────────────────────────────────────────────────────────────────────

def _xl_num(v):
    """The value as the PDF prints it (rmt_plan_pdf.num), kept numeric."""
    if v is None:
        return None
    return round(v, 1) if abs(v) < 100 and round(v, 1) != round(v) else round(v)


def _plain(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text, flags=re.S) if text else text


def _workbook(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    colors = load_colors_config()
    fill = lambda key, dflt: PatternFill("solid", fgColor=colors.get(key, dflt).lstrip("#")[:6])
    hdr_fill = fill("rmt_header_bg", "#E2E8F0")
    app_fill, op_fill, act_fill = fill("perf_app_bg", "#FFF2CC"), fill("rmt_op_bg", "#DDEBF7"), fill("perf_act_bg", "#C6EFCE")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    L = data["labels"]

    def frame(ws, r1, r2, c1, c2):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = border

    def head(ws, r1, r2, ncols, cells):
        """cells: (row, col, text, rowspan, colspan, fill)."""
        for r, c, text, rs, cs, f in cells:
            ws.cell(r, c, text)
            if rs > 1 or cs > 1:
                ws.merge_cells(start_row=r, start_column=c, end_row=r + rs - 1, end_column=c + cs - 1)
            for rr in range(r, r + rs):
                for cc in range(c, c + cs):
                    cell = ws.cell(rr, cc)
                    cell.font, cell.alignment, cell.fill = bold, center, f or hdr_fill
        frame(ws, r1, r2, 1, ncols)

    def block(ws, r, label, rows, values, fills, remarks_col):
        """One item block: the vertical item name merged over its rows."""
        start = r
        for row in rows:
            ws.cell(r, 2, row["row"]).font = Font(bold=True)
            for i, (k, f) in enumerate(zip(values, fills)):
                cell = ws.cell(r, 3 + i, _xl_num(row[k]))
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if f:
                    cell.fill = f
                if row["total"]:
                    cell.font = bold
            rm = ws.cell(r, remarks_col, _plain(row["remarks"]))
            rm.alignment = wrap
            if row["remarks"]:
                ws.row_dimensions[r].height = max(15, 13 * (len(row["remarks"]) // 95 + row["remarks"].count("\n") + 1))
            r += 1
        ws.cell(start, 1, label)
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(start, 1).font = bold
        ws.cell(start, 1).alignment = Alignment(text_rotation=90, horizontal="center", vertical="center", wrap_text=True)
        frame(ws, start, r - 1, 1, remarks_col)
        return r

    wb = Workbook()

    # Sheet 1 -- HM / CS / SS
    ws = wb.active
    ws.title = "HM-CS-SS"
    ws.cell(1, 1, f"Production Plan for {L['month']}").font = Font(bold=True, size=14)
    ws.cell(2, 13, "Unit: '000T").alignment = Alignment(horizontal="right")
    head(ws, 3, 5, 13, [
        (3, 1, "Item", 3, 1, None), (3, 2, "Plant", 3, 1, None), (3, 3, L["ly_col"], 1, 1, None),
        (3, 4, L["prev_col"], 1, 5, None), (3, 9, L["month_col"], 1, 4, None),
        (3, 13, "Remarks regarding S/d plan", 3, 1, None),
        (4, 3, "Actual", 2, 1, act_fill), (4, 4, "APP", 2, 1, app_fill), (4, 5, "OP", 2, 1, op_fill),
        (4, 6, "M/rate", 2, 1, act_fill), (4, 7, "%ful. APP", 2, 1, None), (4, 8, "%ful. OP", 2, 1, None),
        (4, 9, "APP", 1, 2, app_fill), (4, 11, "OP", 1, 2, op_fill),
        (5, 9, "Total", 1, 1, app_fill), (5, 10, "TPD", 1, 1, app_fill),
        (5, 11, "Total", 1, 1, op_fill), (5, 12, "TPD", 1, 1, op_fill),
    ])
    r = 6
    keys = ["ly_actual", "prev_app", "prev_op", "prev_mrate", "ful_app", "ful_op", "app", "app_tpd", "op", "op_tpd"]
    fills = [act_fill, app_fill, op_fill, act_fill, None, None, app_fill, app_fill, op_fill, op_fill]
    for sec in data["page1"]:
        r = block(ws, r, sec["label"], sec["rows"], keys, fills, 13)
    if data.get("note"):
        ws.cell(r + 1, 1, data["note"]).font = bold
    for c, w in enumerate([6, 9] + [9] * 10 + [90], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "C6"

    # Sheet 2 -- Oven Pushing / Sinter / Pig Iron
    ws = wb.create_sheet("Coke-Sinter-Pig")
    ws.cell(1, 1, f"Production Plan for {L['month']}").font = Font(bold=True, size=14)
    ws.cell(2, 7, f"{L['month_col']} · Unit: '000T").alignment = Alignment(horizontal="right")
    r = 3
    for sec in data["page2"]:
        sub = (("Plan", "Coal Requirement per day (T)"), ("Plan", "Coal Requirement per day (T)")) if sec["rate"] \
            else (("APP", "APP (tpd)"), ("OP", "OP (tpd)"))
        head(ws, r, r + 1, 7, [
            (r, 1, "Item", 2, 1, None), (r, 2, "Plant", 2, 1, None), (r, 3, "APP", 1, 2, app_fill),
            (r, 5, "OP", 1, 2, op_fill), (r, 7, "Remarks", 2, 1, None),
            (r + 1, 3, sub[0][0], 1, 1, app_fill), (r + 1, 4, sub[0][1], 1, 1, app_fill),
            (r + 1, 5, sub[1][0], 1, 1, op_fill), (r + 1, 6, sub[1][1], 1, 1, op_fill),
        ])
        ws.row_dimensions[r + 1].height = 30
        r = block(ws, r + 2, sec["label"], sec["rows"], ["app", "app_day", "op", "op_day"],
                  [app_fill, app_fill, op_fill, op_fill], 7) + 1
    for c, w in enumerate([8, 9, 12, 16, 12, 16, 90], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # Sheet 3 -- IPT and stock
    ws = wb.create_sheet("IPT-Stock")
    ws.cell(1, 1, f"IPT Requirement / Availability Plant wise for {L['month']}").font = Font(bold=True, size=13)
    ws.cell(2, 6, "Unit: T").alignment = Alignment(horizontal="right")
    head(ws, 3, 3, 6, [(3, c, t, 1, 1, None) for c, t in
                       enumerate(["Item", "From", "To", "Availability", "Requirement", "Remarks"], start=1)])
    r = 4
    for item, rows in rmt_plan_pdf._ipt_groups(data["ipt"]):
        start = r
        for row in rows:
            for c, k in enumerate(["from_plant", "to_plant", "availability", "requirement"], start=2):
                ws.cell(r, c, row[k]).alignment = center
            ws.cell(r, 6, _plain(row["remarks"])).alignment = wrap
            r += 1
        ws.cell(start, 1, item).font = bold
        ws.cell(start, 1).alignment = center
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        frame(ws, start, r - 1, 1, 6)
    if data["stock"]:
        r += 1
        as_on = data.get("stock_as_on")
        ws.cell(r, 2, "Stock" + (f" as on {as_on[8:10]}.{as_on[5:7]}.{as_on[2:4]}" if as_on else "")).font = bold
        r += 1
        head(ws, r, r, 4, [(r, c, t, 1, 1, None) for c, t in enumerate(["", "BF Coke", "Slab", "Bloom"], start=1)])
        for row in data["stock"]:
            r += 1
            ws.cell(r, 1, row["plant"]).font = bold
            for c, k in enumerate(["bf_coke", "slab", "bloom"], start=2):
                ws.cell(r, c, row[k]).alignment = center
            frame(ws, r, r, 1, 4)
    for c, w in enumerate([14, 10, 10, 14, 14, 90], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    for s in wb.worksheets:
        s.page_setup.orientation = "landscape"
        s.page_setup.paperSize = s.PAPERSIZE_A4
        s.page_setup.fitToWidth, s.page_setup.fitToHeight = 1, 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fname(month: str, ext: str) -> str:
    return f"RMT_Production_Plan_{month}.{ext}"


@router.get("/xlsx")
def rmt_plan_xlsx(month: str = Query(...)):
    _check_month(month)
    return Response(_workbook(rmt_plan.generate(month)),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_fname(month, "xlsx")}"'})


@router.get("/pdf")
def rmt_plan_pdf_download(month: str = Query(...)):
    # A plain def, so FastAPI runs it in its threadpool: the DB queries and
    # the synchronous Playwright render must not block the event loop.
    _check_month(month)
    content = rmt_plan_pdf.render_pdf(rmt_plan.generate(month), load_colors_config())
    return Response(content, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_fname(month, "pdf")}"'})
