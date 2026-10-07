"""Secretary Review deck ("SECRETARY REVIEW Operations Inputs <Mon><YY>.pptx").

build_context(month) gathers the month's numbers; render_pptx(month, texts)
fills secretary_review_templates/secretary_review.pptx (see
scripts/prep_secretary_review_template.py for how the template was made).
Production figures come from page4._p4_row_values so they match the MIS
report; techno figures from techno_period.build_period_report."""

import calendar
import io
from dataclasses import dataclass, field

import board_note_common as bnc
import db
import page4
import page_secretary_review_texts as pst
import page_techno as pt
import secretary_review_pptx as sp
import techno_period as tp
from constants import FIVE_PLANTS
from pptx import Presentation
from pptx.oxml.ns import qn
from secretary_review_layout import (BD_GROUPS, BD_PLANTS, BLOCKS, KPIS, PLANTS, SCOPES,
                                     TEMPLATE_PATH, period_labels)

ITEMS = [("HM", "Hot Metal"), ("CS", "Total Crude Steel"), ("FS", "Finished Steel"), ("SS", "Saleable Steel")]
_P4 = {i["db_item"]: i for i in page4.PAGE4_ITEMS}


def latest_month():
    conn = db.connect()
    cur = conn.cursor()
    try:
        cur.execute("SELECT MAX(report_month) FROM production_table "
                    "WHERE item_name='Hot Metal' AND month_actual IS NOT NULL")
        row = cur.fetchone()
        return row[0] if row and row[0] else None
    finally:
        conn.close()


# --------------------------------------------------------------------------
# Production
# --------------------------------------------------------------------------

def _prod_row(cur, month, plant, db_item, members=None):
    """members: an explicit plant list summed as one row (the SSPs row)."""
    cfg = _P4[db_item]
    sail_set = members if members is not None else cfg["sail_set"]
    p = "SAIL" if members is not None else plant
    raw, cap = page4._p4_row_values(cur, month, p, db_item, False, cfg.get("five_plants", []),
                                    sail_set, has_capacity=True, raw=True)
    return {"cap": cap, "app_m": raw[1], "act_m": raw[2], "cply_m": raw[5], "gr_m": raw[6],
            "cu_m": raw[7], "app_ytd": raw[8], "act_ytd": raw[9], "cply_ytd": raw[12],
            "gr_ytd": raw[13], "cu_ytd": raw[14]}


def _plus(a, b):
    return None if a is None or b is None else a + b


def _with_conversion(cur, month, row):
    """SAIL Finished Steel plus Conversion (slide 5's Total row). CU% scales
    with the actual, since capacity is unchanged."""
    cm, cc, cy, cyc = page4._p4_conv_actuals(cur, month)
    r = dict(row)
    r["act_m"], r["cply_m"] = _plus(row["act_m"], cm), _plus(row["cply_m"], cc)
    r["act_ytd"], r["cply_ytd"] = _plus(row["act_ytd"], cy), _plus(row["cply_ytd"], cyc)
    r["gr_m"] = page4._raw_gr(r["act_m"], r["cply_m"])
    r["gr_ytd"] = page4._raw_gr(r["act_ytd"], r["cply_ytd"])
    for cu, act in (("cu_m", "act_m"), ("cu_ytd", "act_ytd")):
        r[cu] = (row[cu] * r[act] / row[act]) if (row[cu] is not None and r[act] is not None and row[act]) else None
    return r


def build_production(cur, month):
    out = {}
    for key, db_item in ITEMS:
        rows = {s: _prod_row(cur, month, s, db_item) for s in SCOPES}
        ssps = [p for p in _P4[db_item]["sail_set"] if p not in FIVE_PLANTS]
        rows["SSPs"] = _prod_row(cur, month, "SSPs", db_item, members=ssps) if ssps else None
        if key == "FS":
            rows["TOTAL_CONV"] = _with_conversion(cur, month, rows["SAIL"])
        out[key] = rows
    return out


# --------------------------------------------------------------------------
# Techno
# --------------------------------------------------------------------------

def _num(d):
    s = (d or {}).get("display") or ""
    s = s.replace(",", "").strip()
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _targets(labels, warnings):
    """{scope: {param: target}}; a scope whose targets can't be read (e.g. an
    FY with no techno_plan_fy rows yet) gets {} and shows as 'no FY target',
    with the error added to warnings."""
    out = {}
    fetch = {"SAIL": lambda: pt.compute_sail_targets(labels["fy_label"])}
    fetch.update({p: (lambda p=p: pt._get_plant_techno_plan_targets(p, labels["month"])) for p in PLANTS})
    for scope, fn in fetch.items():
        try:
            raw = fn() or {}
        except Exception as e:
            warnings.append(f"{scope} techno targets could not be read: {e}")
            raw = {}
        out[scope] = bnc.fuel_rate_fallback({k[1]: v for k, v in raw.items()})
    return out


def _round_fmt(v, fmt):
    if v is None:
        return None
    return round(v) if fmt == "0" else round(v, 2)


def build_techno(labels, warnings):
    fy_start, n, month = labels["fy_start"], labels["n"], labels["month"]
    months = bnc.fy_months(fy_start)[:n]
    periods = [{"label": "FYm2", "months": bnc.fy_months(fy_start - 2)},
               {"label": "FYm1", "months": bnc.fy_months(fy_start - 1)}]
    periods += [{"label": m, "months": [m]} for m in months]
    periods += [{"label": "ytd:" + m, "months": months[: i + 1]} for i, m in enumerate(months)]
    res = tp.build_period_report(SCOPES, [name for _, name, _ in KPIS], periods)
    display = {p["name"]: p["display_name"] for p in tp._build_major_params()}
    sections = {s["parameter"]: {r["plant"]: r["values"] for r in s["rows"]} for s in res["sections"]}
    targets = _targets(labels, warnings)

    def vals(name, scope):
        return sections.get(display.get(name, name), {}).get(scope, {})

    co2_name = dict((k, nm) for k, nm, _ in KPIS)["co2"]
    sail_co2 = vals(co2_name, "SAIL")
    co2_month = next((m for m in reversed(months) if _num(sail_co2.get(m)) is not None), month)
    if co2_month != month:
        warnings.append(f"Sp. CO₂ for {labels['mon']} not loaded yet – using "
                        f"{calendar.month_abbr[int(co2_month[5:])]}")

    out = {}
    for scope in SCOPES:
        out[scope] = {}
        for key, name, fmt in KPIS:
            v = vals(name, scope)
            used = co2_month if key == "co2" else month
            upto = months[: months.index(used) + 1]
            entry = {
                "fy_m2": _num(v.get("FYm2")), "fy_m1": _num(v.get("FYm1")),
                "target": _round_fmt((targets.get(scope) or {}).get(name), fmt),
                "month": _num(v.get(used)), "ytd": _num(v.get("ytd:" + used)),
                "trend": [(calendar.month_abbr[int(m[5:])], _num(v.get(m))) for m in upto],
                "month_used": used,
            }
            if entry["month"] is None:
                warnings.append(f"{scope} {name}: no value for {calendar.month_abbr[int(used[5:])]}")
            if entry["target"] is None:
                warnings.append(f"{scope} {name}: no FY target")
            out[scope][key] = entry
    return out


def build_context(month):
    labels = period_labels(month)
    warnings = []
    conn = db.connect()
    try:
        production = build_production(conn.cursor(), month)
    finally:
        conn.close()
    if production["HM"]["SAIL"]["act_m"] is None:
        warnings.append(f"No production actuals for {labels['mon']}")
    techno = build_techno(labels, warnings)
    return {"labels": labels, "production": production, "techno": techno, "warnings": warnings}


# --------------------------------------------------------------------------
# Render
# --------------------------------------------------------------------------

ITEM_BY_LABEL = {"hot metal": "HM", "crude steel": "CS", "finished steel": "FS", "saleable steel": "SS"}
PLANT_ROW = {"bsp": "BSP", "dsp": "DSP", "rsp": "RSP", "bsl": "BSL", "isp": "ISP", "ssps": "SSPs", "total": "SAIL"}
VALUE_COLS = ("cap", "app_m", "act_m", "gr_m", "cu_m", "app_ytd", "act_ytd", "gr_ytd", "cu_ytd")


@dataclass
class RenderResult:
    content: bytes
    filename: str
    warnings: list = field(default_factory=list)


def _fmt(v):
    return "" if v is None else str(round(v))


def _fill_values(row, values):
    for i, key in enumerate(VALUE_COLS, start=1):
        if i < len(row.cells):
            sp.set_text_lines(row.cells[i].text_frame, [_fmt((values or {}).get(key))])


def _row_scope(item, scope):
    """SAIL Finished Steel is shown with Conversion added (slides 2 and 5)."""
    return "TOTAL_CONV" if (item == "FS" and scope == "SAIL") else scope


def _fill_item_table(table, production, scope):
    """Rows labelled by item (slide 2 and the plant slides)."""
    for row in table.rows:
        item = ITEM_BY_LABEL.get(sp.norm(row.cells[0].text))
        if item:
            _fill_values(row, production[item].get(_row_scope(item, scope)))


def _fill_section_table(table, production):
    """Item header rows followed by plant rows (slides 3 and 5)."""
    item = None
    for row in table.rows:
        label = sp.norm(row.cells[0].text)
        if label in ITEM_BY_LABEL:
            item = ITEM_BY_LABEL[label]
        elif item and label in PLANT_ROW:
            _fill_values(row, production[item].get(_row_scope(item, PLANT_ROW[label])))


def _lines(text):
    return [ln for ln in (text or "").splitlines() if ln.strip()]


def _fill_hl(shape, text, header):
    """header: the left cell, the same '<period> Highlights' on every slide
    (the template's plant tables say just 'Highlights')."""
    lines = _lines(text)
    if not lines:
        sp.remove_shape(shape)
    else:
        sp.set_text_lines(shape.table.cell(0, 0).text_frame, [header])
        sp.set_text_lines(shape.table.cell(0, 1).text_frame, lines, header_bold=True)


def _fill_bd(shape, plant, texts):
    table = shape.table
    for i, (g, _) in enumerate(BD_GROUPS):
        if i < len(table.rows):
            sp.set_text_lines(table.cell(i, 1).text_frame, _lines(texts.get(f"bd_{plant}_{g}")), header_bold=True)
    empty = [i for i in range(len(table.rows)) if not table.cell(i, 1).text.strip()]
    if len(empty) == len(table.rows):
        for i in reversed(range(1, len(table.rows))):
            sp.remove_row(table, i)
        sp.set_text_lines(table.cell(0, 0).text_frame, [""])
        sp.set_text_lines(table.cell(0, 1).text_frame, ["No major breakdowns"])
    else:
        for i in reversed(empty):
            sp.remove_row(table, i)


def _fill_cr(shape, texts):
    for row in shape.table.rows:
        plant = row.cells[0].text.strip().upper()
        if plant in PLANTS and len(row.cells) >= 3:
            sp.set_text_lines(row.cells[1].text_frame, _lines(texts.get(f"cr_{plant}_cur")), header_bold=True)
            sp.set_text_lines(row.cells[2].text_frame, _lines(texts.get(f"cr_{plant}_prev")), header_bold=True)


LINE_FACTOR = 1.2          # line height = font size x this
CELL_MARGIN_EMU = 91440    # default top + bottom cell margins
DEFAULT_SZ = 1400
MIN_SZ = 800


def _base_sz(cells):
    for c in cells:
        tc = c._tc
        for tag in ("a:rPr", "a:endParaRPr"):
            el = tc.find(".//" + qn(tag))
            if el is not None and el.get("sz"):
                return int(el.get("sz"))
    return DEFAULT_SZ


def _wrapped_lines(cell, sz, col_w):
    """Lines the cell's paragraphs take at font size sz (hundredths of a point),
    with a rough wrap estimate (average glyph width 0.5 em)."""
    total = 0
    for p in cell.text_frame.paragraphs:
        ppr = p._p.find(qn("a:pPr"))
        indent = int(ppr.get("marL", 0)) if ppr is not None else 0
        usable = max(col_w - indent - 2 * 5355, 1)
        chars = max(int(usable / (sz / 100 * 0.5 * 12700)), 1)
        total += max(1, -(-len(p.text) // chars))
    return total


def _fit_table_text(shape, cells, slide_h, warn, label):
    """PowerPoint does not shrink table text, so estimate whether the written
    lines fit between the table's top and the slide bottom and, if not, shrink
    the narrative cells' font in 0.5 pt steps (floor MIN_SZ). Rows without
    narrative cells keep their template height. Warn when even the floor is
    too big."""
    table = shape.table
    budget = max(shape.height, slide_h - shape.top)
    narrative = {id(c._tc) for c in cells}
    base = _base_sz(cells)
    fixed, text_rows = 0, []
    for r in table.rows:
        idx = [i for i, c in enumerate(r.cells) if id(c._tc) in narrative]
        if idx:
            text_rows.append((r, idx))
        else:
            n = max(len(c.text_frame.paragraphs) for c in r.cells)
            fixed += max(r.height, int(n * DEFAULT_SZ / 100 * LINE_FACTOR * 12700) + CELL_MARGIN_EMU)

    def lines(sz):
        return [max(_wrapped_lines(r.cells[i], sz, table.columns[i].width) for i in idx)
                for r, idx in text_rows]

    def need(sz):
        return sum(n * sz / 100 * LINE_FACTOR * 12700 + CELL_MARGIN_EMU for n in lines(sz))

    avail = budget - fixed
    if need(base) <= avail:
        return
    sz = base
    while sz > MIN_SZ and need(sz) > avail:
        sz = max(MIN_SZ, sz - 50)
    for c in cells:
        for el in c._tc.iter(qn("a:rPr"), qn("a:endParaRPr")):
            el.set("sz", str(sz))
    for (r, _), n in zip(text_rows, lines(sz)):
        r.height = int(n * sz / 100 * LINE_FACTOR * 12700 + CELL_MARGIN_EMU)
    if need(sz) > avail:
        warn(f"{label}: {sum(lines(sz))} lines – may overflow; trim the text on the page")


def _fill_charts(shapes, techno, labels, warn):
    for scope in SCOPES:
        for key, _, fmt in KPIS:
            t = techno[scope][key]
            ab = calendar.month_abbr[int(t["month_used"][5:])]
            ytd_cat = labels["ytd_cat"] if t["month_used"] == labels["month"] else f"Apr-{ab}"
            head = [labels["fy_m2"], labels["fy_m1"], labels["fy_tgt"]]
            head_vals = [t["fy_m2"], t["fy_m1"], t["target"]]
            bar = shapes.get(f"ch_{scope}_{key}")
            if bar is None:
                warn(f"template shape ch_{scope}_{key} not found")
            else:
                sp.replace_chart_data(bar.chart, head + [ab, ytd_cat], head_vals + [t["month"], t["ytd"]], fmt)
            trend = shapes.get(f"trend_{scope}_{key}")
            if trend is None:
                warn(f"template shape trend_{scope}_{key} not found")
            else:
                sp.replace_chart_data(trend.chart, head + [a for a, _ in t["trend"]],
                                      head_vals + [v for _, v in t["trend"]], fmt)


def render_pptx(month, texts=None):
    ctx = build_context(month)
    labels, production = ctx["labels"], ctx["production"]
    warnings = list(ctx["warnings"])
    texts = texts or {}
    missing = [k for k, _, _ in BLOCKS if k not in texts]
    eff = {k: v["text"] for k, v in pst.effective_texts(month, missing).items()} if missing else {}
    merged = {k: (texts[k] if k in texts else eff.get(k, "")) for k, _, _ in BLOCKS}

    prs = Presentation(str(TEMPLATE_PATH))
    shapes = sp.named_shapes(prs)

    def get(name):
        sh = shapes.get(name)
        if sh is None:
            warnings.append(f"template shape {name} not found")
        return sh

    if (sh := get("tbl_sail")) is not None:
        _fill_item_table(sh.table, production, "SAIL")
    for p in PLANTS:
        if (sh := get(f"tbl_{p}")) is not None:
            _fill_item_table(sh.table, production, p)
    for name in ("tbl_plants_hm_cs", "tbl_plants_ss_fs"):
        if (sh := get(name)) is not None:
            _fill_section_table(sh.table, production)

    hl_header = f"{labels['period_hdr']} Highlights"
    if (sh := get("tbl_sail_hl")) is not None:
        _fill_hl(sh, merged["hl_SAIL"], hl_header)
    for p in PLANTS:
        if (sh := get(f"tbl_{p}_hl")) is not None:
            _fill_hl(sh, merged[f"hl_{p}"], hl_header)
    for name, key in (("tbl_delay_hmcs", "delay_hmcs"), ("tbl_delay_fs", "delay_fs")):
        if (sh := get(name)) is not None:
            sp.set_text_lines(sh.table.cell(0, 1).text_frame, _lines(merged[key]), header_bold=True)
    for name in ("tbl_cr_1", "tbl_cr_2"):
        if (sh := get(name)) is not None:
            _fill_cr(sh, merged)
    for p in BD_PLANTS:
        if (sh := get(f"tbl_bd_{p}")) is not None:
            _fill_bd(sh, p, merged)

    slide_of = {}
    for n, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            slide_of.setdefault(sh.name, n)
    kinds = {"hl": "Highlights", "delay": "Delay report", "cr": "Capital repairs", "bd": "Breakdowns"}
    for name, sh in shapes.items():
        kind = "hl" if name.endswith("_hl") else name.split("_")[1] if name.startswith("tbl_") else None
        if kind not in kinds or sh._element.getparent() is None or not getattr(sh, "has_table", False):
            continue
        t = sh.table
        if kind == "cr":
            cells = [r.cells[c] for r in t.rows if r.cells[0].text.strip().upper() in PLANTS
                     and len(r.cells) >= 3 for c in (1, 2)]
        elif kind == "bd":
            cells = [r.cells[1] for r in t.rows]
        else:
            cells = [t.cell(0, 1)]
        _fit_table_text(sh, cells, prs.slide_height, warnings.append,
                        f"{kinds[kind]} (slide {slide_of[name]})")

    _fill_charts(shapes, ctx["techno"], labels, warnings.append)
    sp.replace_placeholders(prs, {"MON": labels["mon"], "YTD": labels["ytd"], "CPLY": labels["cply"],
                                  "YTD_PREV": labels["ytd_prev"], "PERIOD": labels["period_hdr"]})
    buf = io.BytesIO()
    prs.save(buf)
    return RenderResult(content=buf.getvalue(), filename=labels["filename"], warnings=warnings)
