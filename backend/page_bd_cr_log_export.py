"""
Excel / PDF export for the Breakdown & Capital Repair log
(frontend: /reports/loss-breakdown?tab=bd-cr-log).

The log's unit grouping (free-text CR shop/equipment mapped onto breakdown
units, linked BD+CR shown once, process-order sort) is built in the browser
(components/reports/bd-cr-log/View.js buildUnits), so the frontend POSTs the
already-built tables and this module only lays them out -- the downloads
match the screen exactly without duplicating that logic here.

Payload: {plant, plant_title, fy, units: [{title, type_label, bd, cr,
events: [{month, kind, linked, dates, sub, details, duration, ongoing}]}]}
-- events in display order; consecutive events of one month form a group.

render_pdf_bytes comes from page_production_query_export (pure
HTML-string-in, PDF-bytes-out), as in page_capital_repair_calendar_export.
"""
import html
import io
from itertools import groupby

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

import page_production_query_export as _ppqe

render_pdf_bytes = _ppqe.render_pdf_bytes

_HDR_FILL = PatternFill("solid", fgColor="F8F9FA")
_HDR_FONT = Font(bold=True, color="5F6368", size=9)
_UNIT_FONT = Font(bold=True, size=11)
_MONTH_FILL = PatternFill("solid", fgColor="FAFAFA")
_KIND = {
    "BD": (PatternFill("solid", fgColor="FCE8E6"), Font(bold=True, size=8, color="C5221F")),
    "CR": (PatternFill("solid", fgColor="E8F0FE"), Font(bold=True, size=8, color="1A56C4")),
}
_THIN = Side(style="thin", color="DADCE0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_COLS = ["Month", "Dates", "Type", "Details", "Duration"]


def _unit_heading(u: dict) -> str:
    bd, cr = u.get("bd", 0), u.get("cr", 0)
    return (f"{u.get('type_label', '')} · {bd} breakdown{'' if bd == 1 else 's'}, "
            f"{cr} capital repair{'' if cr == 1 else 's'}")


def _month_groups(events: list) -> list:
    """[(month, [events])] for consecutive events of the same month."""
    return [(m, list(evs)) for m, evs in groupby(events, key=lambda e: e.get("month", ""))]


def _group_counts(evs: list) -> str:
    nbd = sum(1 for e in evs if e.get("kind") == "BD")
    ncr = len(evs) - nbd
    return " · ".join(p for p in (f"{nbd} BD" if nbd else "", f"{ncr} CR" if ncr else "") if p)


def _details(ev: dict) -> str:
    sub = ev.get("sub")
    return f"[{sub}] {ev.get('details') or ''}" if sub else (ev.get("details") or "")


def _filename(data: dict, ext: str) -> str:
    return f"BD_CR_Log_{data.get('plant', '')}_{data.get('fy', '')}.{ext}"


def build_excel_bytes(data: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"BD-CR Log {data.get('plant', '')}"[:31]

    plant_title = data.get("plant_title") or data.get("plant", "")
    ws.cell(row=1, column=1, value=f"Breakdown & Capital Repair Log — {plant_title}").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"FY {data.get('fy', '')}  ·  BD = Breakdown, CR = Capital Repair "
                                   "(⏱ = timed by the linked breakdown entry)").font = Font(italic=True, size=9)
    row = 4

    for u in data.get("units", []):
        ws.cell(row=row, column=1, value=u.get("title", "")).font = _UNIT_FONT
        ws.cell(row=row, column=3, value=_unit_heading(u)).font = Font(size=9, color="5F6368")
        row += 1
        for c, h in enumerate(_COLS, start=1):
            hc = ws.cell(row=row, column=c, value=h)
            hc.font, hc.fill, hc.border = _HDR_FONT, _HDR_FILL, _BORDER
            hc.alignment = Alignment(horizontal="right" if h == "Duration" else "left")
        row += 1

        for month, evs in _month_groups(u.get("events", [])):
            top = row
            for ev in evs:
                kind_fill, kind_font = _KIND.get(ev.get("kind"), _KIND["BD"])
                values = [None, ev.get("dates", ""),
                          ev.get("kind", "") + (" ⏱" if ev.get("linked") else ""),
                          _details(ev), ev.get("duration", "")]
                for c, v in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=c, value=v)
                    cell.border = _BORDER
                    cell.font = Font(size=9)
                    cell.alignment = Alignment(vertical="top", wrap_text=(c == 4),
                                               horizontal="right" if c == 5 else None)
                kc = ws.cell(row=row, column=3)
                kc.fill, kc.font = kind_fill, kind_font
                if ev.get("ongoing"):
                    ws.cell(row=row, column=5).font = Font(size=9, color="B06000")
                row += 1
            mc = ws.cell(row=top, column=1, value=f"{month}\n{_group_counts(evs)}")
            mc.font = Font(bold=True, size=9)
            mc.alignment = Alignment(vertical="top", wrap_text=True)
            for r in range(top, row):
                ws.cell(row=r, column=1).fill = _MONTH_FILL
                ws.cell(row=r, column=1).border = _BORDER
            if row - top > 1:
                ws.merge_cells(start_row=top, start_column=1, end_row=row - 1, end_column=1)
        row += 1

    for col, width in zip("ABCDE", (12, 26, 8, 80, 14)):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf_html(data: dict) -> str:
    def esc(s):
        return html.escape("" if s is None else str(s))

    plant_title = data.get("plant_title") or data.get("plant", "")
    header = "".join(f'<th class="{"r" if h == "Duration" else ""}">{h}</th>' for h in _COLS)

    sections = []
    for u in data.get("units", []):
        body = []
        for month, evs in _month_groups(u.get("events", [])):
            for i, ev in enumerate(evs):
                kind = ev.get("kind", "BD")
                cells = []
                if i == 0:
                    cells.append(f'<td class="month" rowspan="{len(evs)}">{esc(month)}'
                                 f'<div class="cnt">{esc(_group_counts(evs))}</div></td>')
                sub = f'<span class="sub">[{esc(ev.get("sub"))}]</span> ' if ev.get("sub") else ""
                linked = ' <span class="linked">⏱</span>' if ev.get("linked") else ""
                cells += [
                    f'<td class="nw">{esc(ev.get("dates", ""))}</td>',
                    f'<td class="nw"><span class="kind {kind.lower()}">{esc(kind)}</span>{linked}</td>',
                    f'<td>{sub}{esc(ev.get("details", ""))}</td>',
                    f'<td class="r nw{" ongoing" if ev.get("ongoing") else ""}">{esc(ev.get("duration", ""))}</td>',
                ]
                body.append(f'<tr class="{"first" if i == 0 else ""}">{"".join(cells)}</tr>')
        sections.append(
            f'<div class="unit"><h3>{esc(u.get("title", ""))} <span>{esc(_unit_heading(u))}</span></h3>'
            f'<table><colgroup><col style="width:10%"/><col style="width:24%"/><col style="width:7%"/>'
            f'<col/><col style="width:10%"/></colgroup>'
            f'<thead><tr>{header}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'
        )
    if not sections:
        sections.append('<p class="empty">No breakdowns or capital repairs recorded.</p>')

    return f"""<!doctype html>
<html><head><meta charset="utf-8">
<style>
  @page {{ size: A4 portrait; margin: 10mm 9mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: Arial, sans-serif; color: #202124; margin: 0; }}
  h1 {{ font-size: 14pt; margin: 0 0 2px 0; }}
  .subtitle {{ font-size: 8.5pt; color: #5f6368; margin: 0 0 8px 0; }}
  .unit {{ margin-bottom: 10px; }}
  h3 {{ font-size: 10.5pt; margin: 0 0 3px 0; page-break-after: avoid; }}
  h3 span {{ font-weight: 400; font-size: 8pt; color: #5f6368; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 8pt; table-layout: fixed; }}
  thead {{ display: table-header-group; }}
  tr {{ page-break-inside: avoid; }}
  th, td {{ border: 1px solid #dadce0; padding: 3px 5px; text-align: left; vertical-align: top; }}
  th {{ background: #f8f9fa; color: #5f6368; font-weight: 700; }}
  .r {{ text-align: right; }}
  .nw {{ white-space: nowrap; }}
  tr.first td {{ border-top: 1px solid #9aa0a6; }}
  td.month {{ background: #fafafa; font-weight: 700; }}
  .cnt {{ font-weight: 400; font-size: 7pt; color: #5f6368; }}
  .kind {{ padding: 0 5px; border-radius: 8px; font-size: 7pt; font-weight: 700; }}
  .kind.bd {{ background: #fce8e6; color: #c5221f; }}
  .kind.cr {{ background: #e8f0fe; color: #1a56c4; }}
  .linked {{ font-size: 7pt; color: #5f6368; }}
  .sub {{ color: #5f6368; }}
  .ongoing {{ color: #b06000; }}
  .empty {{ font-size: 9pt; color: #5f6368; }}
</style>
</head>
<body>
  <h1>Breakdown &amp; Capital Repair Log — {esc(plant_title)}</h1>
  <p class="subtitle">FY {esc(data.get("fy", ""))} · BD = Breakdown, CR = Capital Repair ·
    ⏱ = capital repair timed by its linked breakdown entry</p>
  {''.join(sections)}
</body></html>"""
