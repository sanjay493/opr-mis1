"""python-pptx helpers for filling the Secretary Review template
(page_secretary_review.py). Everything edits the existing XML in place so
the template's fonts, colours and chart styling survive."""

import copy
import re

from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip().lower()


def named_shapes(prs) -> dict:
    out = {}

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif sh.name:
                out.setdefault(sh.name, sh)

    for slide in prs.slides:
        walk(slide.shapes)
    return out


def _rpr_from_end(end):
    """A run's rPr copied from a paragraph's endParaRPr, so a run added to an
    empty template cell gets the cell's font."""
    rpr = OxmlElement("a:rPr")
    for k, v in end.attrib.items():
        rpr.set(k, v)
    for child in end:
        rpr.append(copy.deepcopy(child))
    return rpr


def set_text_lines(text_frame, lines, header_bold=False):
    """Replace the frame's paragraphs with one paragraph per non-blank line,
    each formatted like the frame's first paragraph/run. header_bold: a line
    ending in ':', ':-' or ':–' is bold, every other line not bold."""
    lines = [ln.rstrip() for ln in (lines or []) if ln and ln.strip()] or [""]
    tx = text_frame._txBody
    paras = tx.findall(qn("a:p"))
    proto = copy.deepcopy(paras[0]) if paras else OxmlElement("a:p")
    first_run = proto.find(qn("a:r"))
    run_proto = copy.deepcopy(first_run) if first_run is not None else None
    end_proto = proto.find(qn("a:endParaRPr"))
    for p in paras:
        tx.remove(p)
    for line in lines:
        p = copy.deepcopy(proto)
        for child in list(p):
            if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                p.remove(child)
        if run_proto is not None:
            r = copy.deepcopy(run_proto)
        else:
            r = OxmlElement("a:r")
            if end_proto is not None:
                r.append(_rpr_from_end(end_proto))
            r.append(OxmlElement("a:t"))
        r.find(qn("a:t")).text = line
        if header_bold:
            rpr = r.find(qn("a:rPr"))
            if rpr is None:
                rpr = OxmlElement("a:rPr")
                r.insert(0, rpr)
            rpr.set("b", "1" if line.endswith((":", ":-", ":–")) else "0")
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        tx.append(p)


def sub_paragraphs(text_frame, fn):
    """Apply fn to each paragraph's full text; when it changes, the new text
    goes into the first run and the other runs are emptied (placeholders may
    be split across runs)."""
    for p in text_frame.paragraphs:
        runs = p.runs
        if not runs:
            continue
        full = "".join(r.text for r in runs)
        new = fn(full)
        if new != full:
            runs[0].text = new
            for r in runs[1:]:
                r.text = ""


def _text_frames(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                yield sh.text_frame
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame


def replace_placeholders(prs, mapping):
    def fn(text):
        for k, v in mapping.items():
            text = text.replace("{" + k + "}", v)
        return text

    for tf in _text_frames(prs):
        sub_paragraphs(tf, fn)


def replace_chart_data(chart, categories, values, number_format):
    cd = CategoryChartData(number_format=number_format)
    cd.categories = categories
    cd.add_series(chart.plots[0].series[0].name, values)
    chart.replace_data(cd)


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def remove_row(table, idx):
    tbl = table._tbl
    tbl.remove(tbl.tr_lst[idx])
