"""
BSL Month-End Techno Extractor — tentative for-the-month BF techno values
from BSL's "DPR Mail" workbook, sheet 'BF Parameter'.

Sheet layout (one block per parameter, stacked vertically):
  Row 1:  B1 'On date' | C1 <report date> | D1 'Till Date' | E1.. one column
          per day of the month (E1 = 1st of the month).
  Each block: column A = parameter label on its first row, column B = furnace
  (' BF-I' .. ' BF-V', then ' TOTAL' / 'OVERALL' for the shop), column C =
  on-date value, column D = "Till Date" (month-to-date) value.

When generated at month end, column D holds the for-the-month value — that
is what this extractor reads, for:
    Coke Rate, Nut Coke Rate, CDI Rate, BF Productivity, % Sinter in burden,
    % Pellet in burden, HBT (hot blast temp), Oxygen Enrichment, Slag Rate
→ BF-I..BF-V stored under units 'BF-1'..'BF-5', TOTAL/OVERALL under
'BF_Shop' — the same units and param keys the BSL techno Excel / BF
Performance PDF uploads already use, so a later final upload cleanly
replaces these tentative values.

Blocks are located by their column-A label (not fixed rows), so an added or
removed block above doesn't shift anything; a label that can't be found is
reported as a warning rather than silently skipped.

Month verification: the month of the daily date columns (E1 onward) and of
C1 must match the report month selected in the UI (McrMonthMismatch
otherwise). A C1 date short of the month's last day is allowed but warned —
the Till Date column then covers only part of the month.

A furnace with zero Hot Metal till date (shut the whole month) has its rate
values left blank instead of stored as 0.
"""

import calendar
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

sys.path.insert(0, str(Path(__file__).parent))
from dsp_mcr_techno_extractor import McrMonthMismatch, _clean_val  # noqa: E402

_SHEET_NAME = "BF Parameter"

_LABEL_COL = 1     # column A — parameter label (first row of each block)
_FURNACE_COL = 2   # column B — BF-I .. BF-V / TOTAL
_VALUE_COL = 4     # column D — "Till Date" (month to date)

# Normalised column-A label -> canonical param key. Labels are matched
# exactly after normalising (upper case, spaces collapsed) — "CDI" (the
# consumption block, tonnes) must not match "CDI RATE", nor "COKE RATE"
# match "NUT COKE RATE".
_PARAMS = {
    "COKE RATE":           "coke_rate",
    "NUT COKE RATE":       "nut_coke_rate",
    "CDI RATE":            "cdi",
    "BF PRODUCTIVTY":      "bf_productivity",   # sic — the sheet's own spelling
    "BF PRODUCTIVITY":     "bf_productivity",
    "% SINTER IN BURDEN":  "sinter_in_burden",
    "% PELLET IN BURDEN":  "pellet_in_burden",
    "HBT(OC )":            "hot_blast_temp",
    "HBT(OC)":             "hot_blast_temp",
    "OXYGEN ENRICHMENT":   "o2_enrichment",
    "SLAG RATE":           "slag_rate",
}
_HOT_METAL_LABEL = "HOT METAL"

_FURNACE_UNIT = {
    "BF-I": "BF-1", "BF-II": "BF-2", "BF-III": "BF-3", "BF-IV": "BF-4", "BF-V": "BF-5",
    "TOTAL": "BF_Shop", "OVERALL": "BF_Shop",
}


def _norm(v) -> str:
    return " ".join(str(v or "").upper().split())


def _as_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _load_grid(file_path: str) -> List[List]:
    with open(file_path, "rb") as f:
        magic = f.read(4)
    if magic[:2] != b"PK":
        raise ValueError("Unrecognised file format — expected the BSL DPR Mail "
                         "workbook (.xlsx).")
    import openpyxl
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        if _SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Workbook has no '{_SHEET_NAME}' sheet — verify this "
                             "is the BSL DPR Mail workbook.")
        return [list(row) for row in wb[_SHEET_NAME].iter_rows(values_only=True)]
    finally:
        wb.close()


def _cell(grid, row_1b, col_1b):
    if row_1b < 1 or row_1b > len(grid):
        return None
    row = grid[row_1b - 1]
    if col_1b < 1 or col_1b > len(row):
        return None
    v = row[col_1b - 1]
    return None if v == "" else v


class BslMonthendTechnoExtractor:
    """
    extract() returns the standard techno record list:
        [{"plant": "BSL", "report_month": "YYYY-MM", "unit": str,
          "techno_json": {"month": {...}, "till_month": {}}}]

    After extract():  .report_date, .warnings
    """

    def __init__(self, file_path: str, report_month: str = ""):
        self.file_path = file_path
        self.report_month = report_month
        self.report_date: Optional[date] = None
        self.warnings: List[str] = []

    def _read_dates(self, grid):
        on_date = _as_date(_cell(grid, 1, 3))           # C1
        day_cols = [d for d in (_as_date(v) for v in (grid[0][4:] if grid else [])) if d]
        if not on_date and not day_cols:
            raise ValueError(
                f"Cannot read the report date from row 1 of '{_SHEET_NAME}' "
                "(C1 'On date' or the daily date columns from E1) — verify this "
                "is the BSL DPR Mail workbook.")
        sheet_month = (day_cols[0] if day_cols else on_date).strftime("%Y-%m")
        if on_date and on_date.strftime("%Y-%m") != sheet_month:
            raise McrMonthMismatch(
                f"'{_SHEET_NAME}' is inconsistent: C1 'On date' is "
                f"{on_date.strftime('%d.%m.%Y')} but the daily columns are for "
                f"{sheet_month}. Check the workbook.")
        self.report_date = on_date or day_cols[-1]

        if self.report_month and self.report_month != sheet_month:
            raise McrMonthMismatch(
                f"This BSL DPR Mail sheet is for {sheet_month} (On date "
                f"{self.report_date.strftime('%d.%m.%Y')}) but you selected "
                f"{self.report_month}. Select the matching month or upload the "
                "correct file.")
        self.report_month = sheet_month

        last_day = calendar.monthrange(self.report_date.year, self.report_date.month)[1]
        if self.report_date.day != last_day:
            self.warnings.append(
                f"On date is {self.report_date.strftime('%d.%m.%Y')}, not the month end "
                f"({last_day:02d}.{self.report_date.month:02d}.{self.report_date.year}) — "
                "the Till Date values cover only part of the month.")

    def _blocks(self, grid) -> Dict[str, Dict[str, object]]:
        """{normalised column-A label: {unit: raw column-D value}}."""
        blocks: Dict[str, Dict[str, object]] = {}
        current = None
        for r in range(2, len(grid) + 1):
            a = _cell(grid, r, _LABEL_COL)
            if isinstance(a, str) and a.strip():
                current = _norm(a)
                blocks.setdefault(current, {})
            unit = _FURNACE_UNIT.get(_norm(_cell(grid, r, _FURNACE_COL)))
            if current and unit:
                blocks[current][unit] = _cell(grid, r, _VALUE_COL)
            elif current and _cell(grid, r, _FURNACE_COL) is None and a is None:
                current = None   # blank row ends the block
        return blocks

    def extract(self) -> List[Dict]:
        grid = _load_grid(self.file_path)
        self._read_dates(grid)
        blocks = self._blocks(grid)

        hot_metal = blocks.get(_HOT_METAL_LABEL, {})
        idle = {u for u, v in hot_metal.items() if u != "BF_Shop" and _clean_val(v) == 0}
        if idle:
            self.warnings.append(
                f"No Hot Metal till date for {', '.join(sorted(idle))} — its rates are left blank.")

        units: Dict[str, Dict] = {}
        found = set()
        for label, key in _PARAMS.items():
            if label not in blocks:
                continue
            found.add(key)
            for unit, raw in blocks[label].items():
                v = None if unit in idle else _clean_val(raw)
                units.setdefault(unit, {"month": {}, "till_month": {}})
                if v is not None:
                    units[unit]["month"][key] = round(v, 4)
        missing = sorted(set(_PARAMS.values()) - found)
        if missing:
            self.warnings.append(
                f"Block not found in '{_SHEET_NAME}' column A for: {', '.join(missing)} — skipped.")

        records = [
            {"plant": "BSL", "report_month": self.report_month, "unit": unit, "techno_json": tj}
            for unit, tj in sorted(units.items())
            if tj["month"]
        ]
        if not records:
            raise ValueError(
                f"No techno values found in the '{_SHEET_NAME}' sheet — verify the file contents.")
        return records
