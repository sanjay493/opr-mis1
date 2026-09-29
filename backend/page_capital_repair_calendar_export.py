"""
Excel / PDF export for the Capital Repair Plan-vs-Actual calendar
(frontend: /reports/production-analysis?tab=capital-repair-calendar).

Source data: page_capital_repair.generate_capital_repair_calendar(plant, fy)
-- one row per unit (shop/equipment), 12 FY months (Apr-Mar), each month
marked Planned (from the free-text Period) and/or Actual (from the
structured/parsed Actual dates).

Excel can't split one cell's background in half, so each unit gets two
thin, merged-label rows (Plan / Actual) instead of the web UI's split
cell -- same information, different mechanics. The PDF keeps the true
split cell (two stacked half-height <div>s per month) since HTML/CSS can
do it directly, matching what's on screen.

render_pdf_bytes is imported from page_production_query_export directly
rather than duplicated, since it has zero business logic (pure
HTML-string-in, PDF-bytes-out).
"""
import html
import io

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import page_production_query_export as _ppqe

render_pdf_bytes = _ppqe.render_pdf_bytes

_HDR_FILL = PatternFill("solid", fgColor="1A73E8")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
_SHOP_FILL = PatternFill("solid", fgColor="1A73E8")
_SHOP_FONT = Font(bold=True, color="FFFFFF", size=10)
_PLAN_FILL = PatternFill("solid", fgColor="FCE8B2")
_ACTUAL_FILL = PatternFill("solid", fgColor="CFE2FF")
_THIN = Side(style="thin", color="DADCE0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_excel_bytes(data: dict) -> bytes:
    months = data.get("month_labels", [])
    ncols = 3 + len(months)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Capital Repair Calendar"[:31]

    plant_title = data.get("plant_title") or data.get("plant", "")
    ws.cell(row=1, column=1, value=f"Capital Repair — Plan vs Actual — {plant_title}").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"FY {data.get('fy', '')}").font = Font(italic=True, size=9)

    ws.cell(row=3, column=1, value="Legend:").font = Font(bold=True, size=9)
    lp = ws.cell(row=3, column=2, value="Plan")
    lp.fill = _PLAN_FILL
    lp.font = Font(size=9)
    la = ws.cell(row=3, column=3, value="Actual")
    la.fill = _ACTUAL_FILL
    la.font = Font(size=9)

    header_row = 5
    for c, h in enumerate(["Unit", "Activity", ""] + months, start=1):
        hc = ws.cell(row=header_row, column=c, value=h)
        hc.font = _HDR_FONT
        hc.fill = _HDR_FILL
        hc.alignment = Alignment(horizontal="center", wrap_text=True)
        hc.border = _BORDER
    row = header_row + 1

    for sec in data.get("sections", []):
        sc = ws.cell(row=row, column=1, value=sec.get("shop", ""))
        sc.font = _SHOP_FONT
        sc.fill = _SHOP_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        for c in range(2, ncols + 1):
            ws.cell(row=row, column=c).fill = _SHOP_FILL
        row += 1

        for r in sec.get("rows", []):
            top = row
            uc = ws.cell(row=top, column=1, value=r.get("unit", ""))
            uc.font = Font(bold=True, size=9)
            uc.alignment = Alignment(vertical="center", wrap_text=True)
            ac = ws.cell(row=top, column=2, value=r.get("activity", ""))
            ac.font = Font(size=9)
            ac.alignment = Alignment(vertical="center", wrap_text=True)
            ws.merge_cells(start_row=top, start_column=1, end_row=top + 1, end_column=1)
            ws.merge_cells(start_row=top, start_column=2, end_row=top + 1, end_column=2)

            for rr, (key, fill, label) in enumerate((("plan", _PLAN_FILL, "Plan"), ("actual", _ACTUAL_FILL, "Actual"))):
                rw = top + rr
                lc = ws.cell(row=rw, column=3, value=label)
                lc.font = Font(size=8, italic=True, color="5F6368")
                lc.border = _BORDER
                for ci, m in enumerate(r.get("months", []), start=4):
                    cell = ws.cell(row=rw, column=ci)
                    cell.border = _BORDER
                    if m.get(key):
                        cell.fill = fill
                for c in (1, 2):
                    ws.cell(row=rw, column=c).border = _BORDER
            row += 2

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 8
    for i in range(4, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.freeze_panes = f"D{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf_html(data: dict) -> str:
    def esc(s):
        return html.escape("" if s is None else str(s))

    months = data.get("month_labels", [])
    plant_title = data.get("plant_title") or data.get("plant", "")
    header_html = "<th class='lbl'>Unit</th><th class='lbl'>Activity</th>" + "".join(f"<th>{esc(m)}</th>" for m in months)

    sections_html = []
    for sec in data.get("sections", []):
        body_rows = []
        for r in sec.get("rows", []):
            cells = [f'<td class="lbl">{esc(r.get("unit", ""))}</td>',
                     f'<td class="lbl txt">{esc(r.get("activity", ""))}</td>']
            for m in r.get("months", []):
                plan_on = " on" if m.get("plan") else ""
                actual_on = " on" if m.get("actual") else ""
                cells.append(
                    f'<td class="cal"><div class="half plan{plan_on}"></div>'
                    f'<div class="half actual{actual_on}"></div></td>'
                )
            body_rows.append(f"<tr>{''.join(cells)}</tr>")
        sections_html.append(
            f'<div class="section-title">{esc(sec.get("shop", ""))}</div>'
            f'<table><thead><tr>{header_html}</tr></thead>'
            f'<tbody>{"".join(body_rows)}</tbody></table>'
        )

    return f"""<!doctype html>
<html><head><meta charset="utf-8">
<style>
  @page {{ size: A4 landscape; margin: 10mm 8mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: Arial, sans-serif; color: #202124; margin: 0; }}
  h1 {{ font-size: 14pt; margin: 0 0 2px 0; }}
  .subtitle {{ font-size: 9pt; color: #5f6368; margin: 0 0 6px 0; }}
  .legend {{ font-size: 8.5pt; margin: 0 0 8px 0; }}
  .legend .swatch {{ display: inline-block; width: 10px; height: 10px; margin: 0 4px 0 12px; vertical-align: middle; border: 1px solid #dadce0; }}
  .legend .plan {{ background: #fce8b2; }}
  .legend .actual {{ background: #cfe2ff; }}
  .section-title {{ background: #1a73e8; color: #fff; font-weight: 700; font-size: 9.5pt;
    padding: 4px 6px; margin-top: 10px; page-break-after: avoid; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 8pt; margin-bottom: 4px; table-layout: fixed; }}
  thead {{ display: table-header-group; }}
  tr {{ page-break-inside: avoid; }}
  th, td {{ border: 1px solid #dadce0; padding: 2px 4px; text-align: center; }}
  th {{ background: #e8f0fe; color: #174ea6; font-weight: 700; }}
  th.lbl, td.lbl {{ text-align: left; white-space: normal; }}
  td.lbl {{ font-weight: 600; }}
  td.txt {{ font-weight: 400; color: #5f6368; }}
  td.cal {{ padding: 0; }}
  .half {{ height: 9px; background: #ffffff; }}
  .half.plan.on {{ background: #fce8b2; }}
  .half.actual.on {{ background: #cfe2ff; }}
</style>
</head>
<body>
  <h1>Capital Repair — Plan vs Actual — {esc(plant_title)}</h1>
  <p class="subtitle">FY {esc(data.get("fy", ""))}</p>
  <p class="legend"><b>Legend:</b><span class="swatch plan"></span>Plan<span class="swatch actual"></span>Actual</p>
  {''.join(sections_html)}
</body></html>"""
