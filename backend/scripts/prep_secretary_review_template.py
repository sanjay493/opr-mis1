"""One-off: turn the user-updated Sep'26 Secretary Review deck into the
generator's template (secretary_review_templates/secretary_review.pptx).

- names every filled table/chart (see secretary_review_layout.required_shape_names)
- puts {MON}/{YTD}/{CPLY}/{YTD_PREV}/{PERIOD} placeholders in titles and headers
- adds a highlights table to the DSP/RSP/ISP plant slides (cloned from BSP's)
- normalises each breakdown table to three rows: BFs / SMS / Mills

Slide indices are used only here. Re-run only to rebuild the template:
    venv/Scripts/python.exe scripts/prep_secretary_review_template.py "<reference.pptx>"
"""

import copy
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

import secretary_review_layout as L  # noqa: E402
import secretary_review_pptx as sp  # noqa: E402

PLANT_SLIDES = {8: "BSP", 10: "DSP", 12: "RSP", 14: "BSL", 16: "ISP"}
BAR_SLIDES = {7: "SAIL", 9: "BSP", 11: "DSP", 13: "RSP", 15: "BSL", 17: "ISP"}
TREND_SLIDES = {25: "SAIL", 26: "BSP", 27: "DSP", 28: "RSP", 29: "BSL", 30: "ISP"}
BD_SLIDES = {21: "BSP", 22: "DSP", 23: "RSP", 24: "ISP"}
SINGLE_TABLE = {3: "tbl_plants_hm_cs", 4: "tbl_delay_hmcs", 5: "tbl_plants_ss_fs",
                6: "tbl_delay_fs", 19: "tbl_cr_1", 20: "tbl_cr_2"}

TITLES = {
    2: "SAIL  Performance : {MON} & {YTD}",
    3: "SAIL Plant-wise Performance  :  {MON}  &  {YTD}",
    4: "Major Breakdowns/Delays  during  {MON} & {YTD}",
    5: "SAIL Plant-wise Performance  :  {MON}  &  {YTD}",
    6: "Major Breakdowns/Delays  during  {MON} & {YTD}",
    7: "SAIL  Techno  Performance ({YTD})",
    **{i: f"{p}  Performance : {{MON}} & {{YTD}}" for i, p in PLANT_SLIDES.items()},
    **{i: f"{p}: Major Breakdowns/Delays during  {{MON}}" for i, p in BD_SLIDES.items()},
}

HEADER_SUBS = [
    (re.compile(r"Apr-(Sep|Aug)\s*[’']\s*26"), "{YTD}"),
    (re.compile(r"Apr-(Sep|Aug)\s*[’']\s*25"), "{YTD_PREV}"),
    (re.compile(r"Sep\s*[’']\s*26"), "{MON}"),
    (re.compile(r"Sep\s*[’']\s*25"), "{CPLY}"),
    (re.compile(r"H-1\s+Highlights"), "{PERIOD} Highlights"),
]


def kpi_of(title: str):
    t = sp.norm(title)
    if "coke" in t:
        return "coke"
    if "pci" in t:
        return "pci"
    if "fuel" in t:
        return "fuel"
    if "productivity" in t:
        return "bfprod"
    if "energy" in t:
        return "energy"
    if re.search(r"co\s*2", t):
        return "co2"
    return None


def tables(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]


def header_subs(text):
    for rx, rep in HEADER_SUBS:
        text = rx.sub(rep, text)
    return text


def set_title(slide, idx):
    for sh in slide.shapes:
        if sh.has_text_frame and re.search(r"performance|breakdowns", sp.norm(sh.text_frame.text)):
            sp.set_text_lines(sh.text_frame, [TITLES[idx]])
            sh.name = f"title_{idx}"
            return
    raise SystemExit(f"slide {idx}: title not found")


def split_main_and_hl(slide):
    tbls = tables(slide)
    main = [t for t in tbls if len(t.table.rows) > 1]
    hl = [t for t in tbls if len(t.table.rows) == 1]
    return main[0], (hl[0] if hl else None)


def clone_hl(src_hl, slide, plant_tbl, name):
    el = copy.deepcopy(src_hl._element)
    slide.shapes._spTree.insert_element_before(el, "p:extLst")
    new = tables(slide)[-1]
    new.name = name
    new.top = max(src_hl.top, plant_tbl.top + plant_tbl.height + Emu(60000))
    sp.set_text_lines(new.table.cell(0, 0).text_frame, ["Highlights"])
    sp.set_text_lines(new.table.cell(0, 1).text_frame, [""])
    return new


def normalise_bd(tbl_shape):
    t = tbl_shape.table._tbl
    trs = list(t.tr_lst)
    total_h = sum(int(tr.get("h", "0")) for tr in trs)
    proto = trs[0]
    for tr in trs:
        t.remove(tr)
    for _ in L.BD_GROUPS:
        tr = copy.deepcopy(proto)
        tr.set("h", str(total_h // len(L.BD_GROUPS)))
        t.append(tr)
    for i, (_, label) in enumerate(L.BD_GROUPS):
        sp.set_text_lines(tbl_shape.table.cell(i, 0).text_frame, [label])
        sp.set_text_lines(tbl_shape.table.cell(i, 1).text_frame, [""])


def main(src):
    prs = Presentation(src)
    slides = {i + 1: s for i, s in enumerate(prs.slides)}

    for idx in TITLES:
        set_title(slides[idx], idx)

    main_tbl, hl = split_main_and_hl(slides[2])
    main_tbl.name, hl.name = "tbl_sail", "tbl_sail_hl"

    bsp_hl = None
    for idx, plant in PLANT_SLIDES.items():
        main_tbl, hl = split_main_and_hl(slides[idx])
        main_tbl.name = f"tbl_{plant}"
        if hl is not None:
            hl.name = f"tbl_{plant}_hl"
            if plant == "BSP":
                bsp_hl = hl
    for idx, plant in PLANT_SLIDES.items():
        if f"tbl_{plant}_hl" not in sp.named_shapes(prs):
            clone_hl(bsp_hl, slides[idx], sp.named_shapes(prs)[f"tbl_{plant}"], f"tbl_{plant}_hl")

    for idx, name in SINGLE_TABLE.items():
        tables(slides[idx])[0].name = name

    for idx, plant in BD_SLIDES.items():
        t = tables(slides[idx])[0]
        t.name = f"tbl_bd_{plant}"
        normalise_bd(t)

    for mapping, prefix in ((BAR_SLIDES, "ch"), (TREND_SLIDES, "trend")):
        for idx, scope in mapping.items():
            for sh in slides[idx].shapes:
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    title = "".join(e.text or "" for e in sh.chart._chartSpace.xpath(".//c:title//a:t"))
                    key = kpi_of(title)
                    if key is None:
                        raise SystemExit(f"slide {idx}: unknown chart title {title!r}")
                    sh.name = f"{prefix}_{scope}_{key}"

    for slide in prs.slides:
        for t in tables(slide):
            for row in t.table.rows:
                for cell in row.cells:
                    sp.sub_paragraphs(cell.text_frame, header_subs)

    missing = L.required_shape_names() - set(sp.named_shapes(prs))
    if missing:
        raise SystemExit(f"missing shapes: {sorted(missing)}")
    L.TEMPLATE_PATH.parent.mkdir(exist_ok=True)
    prs.save(str(L.TEMPLATE_PATH))
    print(f"wrote {L.TEMPLATE_PATH}")


if __name__ == "__main__":
    main(sys.argv[1])
