"""One-off: copy the hand-written narrative of a Secretary Review deck that
predates the template (the Sep'26 reference) into secretary_review_text.

    venv/Scripts/python.exe scripts/import_secretary_review_texts.py --month 2026-09 --pptx "<deck.pptx>" [--dry-run]
"""

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation  # noqa: E402

import secretary_review_pptx as sp  # noqa: E402

HL_SLIDES = {2: "SAIL", 8: "BSP", 10: "DSP", 12: "RSP", 14: "BSL", 16: "ISP"}
BD_SLIDES = {21: "BSP", 22: "DSP", 23: "RSP", 24: "ISP"}
PLANTS = {"BSP", "DSP", "RSP", "BSL", "ISP"}


def _cell_text(cell):
    return "\n".join(p.text.strip() for p in cell.text_frame.paragraphs if p.text.strip())


def _tables(slide):
    return [sh.table for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]


def _group(label):
    n = sp.norm(label)
    if n.startswith("bf"):
        return "BF"
    if n.startswith("sms"):
        return "SMS"
    if n.startswith("mill"):
        return "MILL"
    return None


def extract_texts(pptx_path):
    slides = {i + 1: s for i, s in enumerate(Presentation(str(pptx_path)).slides)}
    out = {}
    for idx, scope in HL_SLIDES.items():
        for t in _tables(slides[idx]):
            if len(t.rows) == 1 and len(t.columns) >= 2:
                out[f"hl_{scope}"] = _cell_text(t.cell(0, 1))
    for idx, key in ((4, "delay_hmcs"), (6, "delay_fs")):
        out[key] = _cell_text(_tables(slides[idx])[0].cell(0, 1))
    for idx in (19, 20):
        for row in _tables(slides[idx])[0].rows:
            plant = row.cells[0].text.strip().upper()
            if plant in PLANTS:
                out[f"cr_{plant}_cur"] = _cell_text(row.cells[1])
                out[f"cr_{plant}_prev"] = _cell_text(row.cells[2])
    for idx, plant in BD_SLIDES.items():
        for row in _tables(slides[idx])[0].rows:
            g = _group(row.cells[0].text)
            if g:
                out[f"bd_{plant}_{g}"] = _cell_text(row.cells[1])
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True)
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    texts = extract_texts(a.pptx)
    for k, v in texts.items():
        print(f"--- {k}\n{v}")
    if not a.dry_run:
        import secretary_review_text as srt
        print(f"saved {srt.save_texts(a.month, texts)} blocks for {a.month}")


if __name__ == "__main__":
    main()
