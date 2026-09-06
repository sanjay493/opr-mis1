"""
Backfill: BSL "Semis Despatch" (total Slab despatch across every sale
type — Direct + Secondary/Stock Yard + IPT + Export — NOT just the export
portion, which is the pre-existing "Semis Export" item), per direct
instruction. Mirrors the same 3-source approach already used for the
other BSL despatch items:

  FY2025-26 (Apr'25-Mar'26)  <- BSL Annual Statistics 2025-26 PDF, Table
      24.6's "SLAB & THICK PLATE" page (PDF viewer page 444, doc's own
      footer "Page No. 399" — same +45 page-number offset as every other
      Annual Statistics table already used), Grand Total column (index 4
      of the row's 8 numbers: Direct, Stock Yard, IPT, Export, Grand
      Total, Defective, Total, Thick Plate — cross-checked: Direct/IPT/
      Export match the SALEABLE STEEL/EXPORT tables' own columns for the
      same months to the tonne).
  Everything else           <- excel_extractor_bsl.py's own
      extract_preview_main_products_pdf (monthly PDF, "Slab" row just
      above "Thick Plate" on the DESPATCH PERFORMANCE AT A GLANCE page)
      for every Rev/Review file on disk, plus the 2 month-end DPR files
      that cover months no monthly PDF does (2026-05, 2026-08) via the
      DPR route's own new "Semis Despatch" label search.

Same known gaps as the other BSL despatch items: no source for April
2026 or July 2026.

Usage:
  python scripts/backfill_bsl_semis_despatch.py            # dry-run
  python scripts/backfill_bsl_semis_despatch.py --apply     # writes to the live DB

Dry-run is the default on purpose — this touches the live MySQL DB
(DB_ENGINE=mysql), so the values should be reviewed before anything is
written.
"""
import sys
import os
import re
import argparse
import datetime
import calendar

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)
sys.path.insert(0, os.path.join(BACKEND_DIR, "excel_extractors"))

import pdfplumber  # noqa: E402
import openpyxl  # noqa: E402
import db  # noqa: E402
from excel_extractors.excel_extractor_bsl import (  # noqa: E402
    extract_preview_main_products_pdf, _dpr_config, _resolve_dpr_cells,
    _dpr_value_and_alert, _dpr_cum_cell, clean_val,
)

PLANT = "BSL"
ITEM = "Semis Despatch"

ANNUAL_STATS_PDF = r"C:\Users\sanja\Downloads\BSL Annual Statistics 2025-26.pdf"
SLAB_THICK_PLATE_PAGE = 443  # 0-based; viewer page 444, doc's own "Page No. 399"
_SLAB_GRAND_TOTAL_IDX = 4    # Direct, Stock Yard, IPT, Export, [Grand Total], Defective, Total, Thick Plate
_SLAB_NUM_COLS = 8

_MONTH_TO_REPORT_MONTH = {
    "APR'25": "2025-04", "MAY": "2025-05", "JUN": "2025-06", "JULY": "2025-07",
    "AUG": "2025-08", "SEP": "2025-09", "OCT": "2025-10", "NOV": "2025-11",
    "DEC": "2025-12", "JAN'26": "2026-01", "FEB": "2026-02", "MAR": "2026-03",
}
_MONTH_ROW_RE = re.compile(
    r"^(APR'25|MAY|JUN|JULY|AUG|SEP|OCT|NOV|DEC|JAN'26|FEB|MAR)\s+(.+)$"
)

MONTHLY_DIR = r"G:\My Drive\Report_format\Monthly\BSL"
MONTHEND_DIR = r"G:\My Drive\Report_format\MONTHEND\BSL"

PDF_FILES = [
    ("Rev June 2022.pdf", "2022-06"), ("Rev July 2022.pdf", "2022-07"),
    ("Rev Aug 2022.pdf", "2022-08"), ("Rev Sept 2022.pdf", "2022-09"),
    ("Rev Oct 2022.pdf", "2022-10"), ("Rev Nov 2022.pdf", "2022-11"),
    ("Rev Dec 2022.pdf", "2022-12"), ("Rev Jan 2023.pdf", "2023-01"),
    ("Rev Feb 23.pdf", "2023-02"), ("Rev March 2023.pdf", "2023-03"),
    ("Review April 2023.pdf", "2023-04"), ("Review May 2023.pdf", "2023-05"),
    ("Review June 2023.pdf", "2023-06"), ("Review July 2023.pdf", "2023-07"),
    ("Review August 2023.pdf", "2023-08"), ("Review Sept 2023.pdf", "2023-09"),
    ("Review Oct 2023.pdf", "2023-10"), ("Review Nov 2023.pdf", "2023-11"),
    ("Review Dec 2023.pdf", "2023-12"), ("Review Jan 2024.pdf", "2024-01"),
    ("Review Feb 2024.pdf", "2024-02"), ("Review March24.pdf", "2024-03"),
    ("Rev April 2024.pdf", "2024-04"), ("Rev May 2024.pdf", "2024-05"),
    ("Rev June 2024.pdf", "2024-06"), ("Rev July 2024.pdf", "2024-07"),
    ("Rev August 2024.pdf", "2024-08"), ("Rev Sept 2024.pdf", "2024-09"),
    ("Rev Oct 2024.pdf", "2024-10"), ("Rev Nov 2024.pdf", "2024-11"),
    ("Rev December 2024.pdf", "2024-12"), ("Rev Jan 2025.pdf", "2025-01"),
    ("Rev Feb 2025.pdf", "2025-02"), ("Rev March 2025.pdf", "2025-03"),
    ("Rev June26 (2).pdf", "2026-06"),
]
DPR_ONLY_FILES = [
    ("BSL-DPR31052026.xlsx", "2026-05"),
    ("DPR Mail_31082026_Rev.xlsx", "2026-08"),
]


def extract_annual_stats() -> dict:
    """{report_month: value_in_tonnes} for FY25-26 from the Annual
    Statistics PDF's SLAB & THICK PLATE page."""
    with pdfplumber.open(ANNUAL_STATS_PDF) as pdf:
        text = pdf.pages[SLAB_THICK_PLATE_PAGE].extract_text(layout=False)

    out = {}
    for line in text.splitlines():
        m = _MONTH_ROW_RE.match(line.strip())
        if not m:
            continue
        label, rest = m.groups()
        nums = re.findall(r'-?\d[\d,]*(?:\.\d+)?', rest)
        if len(nums) != _SLAB_NUM_COLS:
            continue
        vals = [float(n.replace(',', '')) for n in nums]
        out[_MONTH_TO_REPORT_MONTH[label]] = vals[_SLAB_GRAND_TOTAL_IDX]
    return out


def extract_pdf(fname: str, report_month: str):
    path = os.path.join(MONTHLY_DIR, fname)
    result = extract_preview_main_products_pdf(path, report_month)
    for r in result["production_rows"]:
        if r["item_name"] == ITEM and r["status"] == "ok" and r["value"] is not None:
            return result["month"], r["value"]  # already '000T
    return result["month"], None


def extract_dpr(fname: str):
    path = os.path.join(MONTHEND_DIR, fname)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["DPR"]

    o1_raw = ws["O1"].value
    if not isinstance(o1_raw, datetime.datetime):
        raise ValueError(f"{fname}: cell O1 isn't a date.")
    db_report_month = f"{o1_raw.year}-{str(o1_raw.month).zfill(2)}"
    days_in_month = calendar.monthrange(o1_raw.year, o1_raw.month)[1]

    defaults, _no_convert, _derived = _dpr_config()
    resolved = _resolve_dpr_cells(ws, defaults)
    cell = resolved.get(ITEM)
    if not cell:
        return db_report_month, None

    mrate = clean_val(ws[cell].value)
    cum = clean_val(ws[_dpr_cum_cell(cell)].value)
    val, _basis, _alert = _dpr_value_and_alert(mrate, cum, o1_raw.day, days_in_month, label=ITEM)
    return db_report_month, (round(val / 1000.0, 3) if val is not None else None)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true",
                        help="Actually write to the live DB. Default is a dry-run.")
    args = parser.parse_args()

    by_month = {}

    print("=== BSL Annual Statistics 2025-26 (Slab Grand Total) ===")
    for report_month, tonnes in sorted(extract_annual_stats().items()):
        stored = round(tonnes / 1000.0, 3)
        by_month[report_month] = stored
        print(f"{report_month}: {tonnes:,.0f} T")

    for fname, guessed_month in PDF_FILES:
        db_month, val = extract_pdf(fname, guessed_month)
        print(f"{db_month} ({fname}): {val if val is not None else '(no value)'}")
        if val is not None:
            by_month[db_month] = val

    for fname, expected_month in DPR_ONLY_FILES:
        db_month, val = extract_dpr(fname)
        assert db_month == expected_month, f"{fname}: expected {expected_month}, got {db_month}"
        print(f"{db_month} ({fname}): {val if val is not None else '(no value)'}")
        if val is not None:
            by_month[db_month] = val

    print(f"\n{len(by_month)} months total for {ITEM} ({PLANT}).")

    if not args.apply:
        print("Dry-run only — nothing written. Re-run with --apply to save.")
        return

    conn = db.connect()
    cur = conn.cursor()
    try:
        for report_month, stored in by_month.items():
            cur.execute("""
                INSERT INTO production_table (report_month, plant_name, item_name, month_actual)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(report_month, plant_name, item_name)
                DO UPDATE SET month_actual = excluded.month_actual
            """, (report_month, PLANT, ITEM, stored))
        conn.commit()
    finally:
        conn.close()

    print(f"Saved {len(by_month)} values for {PLANT} / {ITEM}.")


if __name__ == "__main__":
    main()
