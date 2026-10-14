"""Annexures of the quarterly Board Note, generated from the DB as PDF.

Laid out like the hand-made Q-2 annexures (Report_format/Miscellaneous/
Annexure-I/II/IV/V.pdf, whose IV and V are numbered III and IV
here) but for any FY and quarter:

  I-a  Production performance, the quarter's three months (APP / Actual /
       %Ful / CPLY / %Gr) -- production_table + production_plan_table,
       with page 4's plant rollups (board_note_common.period_sum).
  I-b  The same for the quarter and the FY-to-date ("Apr-Sep'25"); Q-1 has
       only the quarter block (its FY-to-date is the quarter).
  II   SAIL + plant techno-economic parameters (Coke Rate, CDI, BF
       Productivity, Sp. Energy) -- techno_period.build_period_report;
       targets from techno_plan_fy; best-ever annual over every earlier FY.
  III  The same three BF parameters for the new blast furnaces, furnace-wise
       (techno_bf_period) with the furnace's own techno_plan_fy target.
  IV   Product mix by mill -- page_pmix_fy_report.compute_pmix_fy_rows, read
       once from the actuals and once from the ABP.

These are stand-alone pages, not part of the layout-guarded monthly report.
Colours come from colors_config.json.
"""
from html import escape

import db
import page_techno as pt
import techno_period as tp
import techno_bf_period as tbp
import board_note_common as bnc
from page_pmix_fy_report import compute_pmix_fy_rows

ANNEXURES = ("I", "II", "III", "IV")

_MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_MONTH_FULL = ["January", "February", "March", "April", "May", "June", "July",
               "August", "September", "October", "November", "December"]


# --------------------------------------------------------------------------
# Periods and labels
# --------------------------------------------------------------------------

def _shift_year(months, n=-1):
    return [f"{int(m[:4]) + n}{m[4:]}" for m in months]


def _yy(y):
    return str(y)[2:]


def _fy_label(fy_start):
    return f"{_yy(fy_start)}-{_yy(fy_start + 1)}"


def _mon_slash(m):          # "Jul/25"
    return f"{_MON[int(m[5:7]) - 1]}/{m[2:4]}"


def _mon_full(m):           # "July'25"
    return f"{_MONTH_FULL[int(m[5:7]) - 1]}'{m[2:4]}"


def _range_label(months):   # "Apr-Sep'25"
    a, b = months[0], months[-1]
    return f"{_MON[int(a[5:7]) - 1]}-{_MON[int(b[5:7]) - 1]}'{b[2:4]}"


def _q_label(fy_start, q):  # "Q-2'25-26"
    return f"Q-{q}'{_fy_label(fy_start)}"


def periods(fy_start: int, quarter: int) -> dict:
    """Every month list and label the annexures use."""
    qm = bnc.quarter_months(fy_start, quarter)
    cum = bnc.fy_months(fy_start)[:quarter * 3]
    prev = bnc.fy_months(fy_start)[:(quarter - 1) * 3]
    long_short = {2: "H-1", 3: "9M", 4: "FY"}.get(quarter)
    return {
        "fy_start": fy_start, "quarter": quarter,
        "fy_months": bnc.fy_months(fy_start),
        "months": qm,
        "q": qm, "q_cply": _shift_year(qm),
        "cum": cum if quarter > 1 else None,
        "cum_cply": _shift_year(cum) if quarter > 1 else None,
        "prev": prev or None,
        "q_label": _q_label(fy_start, quarter),
        "q_cply_label": _q_label(fy_start - 1, quarter),
        "cum_label": _range_label(cum) if quarter > 1 else None,
        "cum_cply_label": _range_label(_shift_year(cum)) if quarter > 1 else None,
        # Annexures II/III name the FY-to-date the way the note does: H-1'25-26
        "long_label": f"{long_short}'{_fy_label(fy_start)}" if long_short else None,
        "long_cply_label": f"{long_short}'{_fy_label(fy_start - 1)}" if long_short else None,
        "prev_label": (_q_label(fy_start, 1) if quarter == 2
                       else _range_label(prev) if prev else None),
    }


# --------------------------------------------------------------------------
# Number helpers
# --------------------------------------------------------------------------

def _ful(act, plan):
    return round(act / plan * 100) if (act is not None and plan) else None


def _gr(act, cply):
    return round((act - cply) / cply * 100) if (act is not None and cply) else None


def _i(v):
    if v is None:
        return ""
    s = f"{v:.0f}"
    return "0" if s == "-0" else s


# --------------------------------------------------------------------------
# Annexure-I: production performance
# --------------------------------------------------------------------------

_P5 = ["BSP", "DSP", "RSP", "BSL", "ISP"]
ANNEX_I_ITEMS = [
    ("Hot Metal", "Hot Metal", _P5 + ["5 Plants", "VISL", "SAIL"]),
    ("Crude Steel", "Total Crude Steel", _P5 + ["5 Plants", "ASP", "SSP", "VISL", "SAIL"]),
    ("Saleable Steel", "Saleable Steel", _P5 + ["5 Plants", "ASP", "SSP", "VISL", "SAIL"]),
    ("Pig Iron", "Pig Iron", _P5 + ["5 Plants", "VISL", "SAIL"]),
    ("Finished Steel", "Finished Steel", _P5 + ["5 Plants", "ASP", "SSP", "VISL", "SAIL"]),
]


def _block(cur, plant, db_item, months, cply_months):
    plan = bnc.period_sum(cur, "plan", months, plant, db_item)
    act = bnc.period_sum(cur, "act", months, plant, db_item)
    cply = bnc.period_sum(cur, "act", cply_months, plant, db_item)
    return [plan, act, _ful(act, plan), cply, _gr(act, cply)]


def annexure_i(cur, P) -> list:
    """Two sheets: [I-a (months), I-b (quarter + FY-to-date)]."""
    blocks_a = [(_mon_slash(m), _mon_slash(_shift_year([m])[0]), [m], _shift_year([m]))
                for m in P["months"]]
    blocks_b = [(P["q_label"], P["q_cply_label"], P["q"], P["q_cply"])]
    if P["cum"]:
        blocks_b.append((P["cum_label"], P["cum_cply_label"], P["cum"], P["cum_cply"]))
    fy = _fy_label(P["fy_start"])
    sheets = []
    for tag, blocks, title in (
        ("a", blocks_a, "SAIL: Production Performance during "
         + ", ".join(_mon_full(m) for m in P["months"][:-1]) + f" and {_mon_full(P['months'][-1])}"),
        ("b", blocks_b, f"SAIL: Production Performance during {P['q_label']}"
         + (f" and {P['cum_label']}" if P["cum"] else "")),
    ):
        groups = []
        for label, db_item, plants in ANNEX_I_ITEMS:
            rows = []
            for plant in plants:
                ann = bnc.period_sum(cur, "plan", P["fy_months"], plant, db_item)
                vals = [ann]
                for _l, _c, months, cply in blocks:
                    vals += _block(cur, plant, db_item, months, cply)
                rows.append({"label": plant, "values": vals, "bold": plant in ("5 Plants", "SAIL")})
            groups.append({"title": label, "rows": rows})
        sheets.append({"tag": f"I{tag}", "title": title, "fy": fy,
                       "blocks": [(l, c) for l, c, _m, _p in blocks], "groups": groups})
    return sheets


# --------------------------------------------------------------------------
# Annexure-II: SAIL & plant techno-economic parameters
# --------------------------------------------------------------------------

# (row label, techno_period display name, decimals, better)
ANNEX_II_PARAMS = [
    ("Coke Rate (kg/thm)", "Coke Rate", 0, "low"),
    ("CDI (kg/thm)", "CDI Rate", 0, "high"),
    ("BF Productivity (t/m3/day)", "BF Productivity", 2, "high"),
    ("Specific Energy Consumption (Gcal/tcs)", "Specific Energy Consumption", 2, "low"),
]
ANNEX_II_PLANTS = [
    ("SAIL", "SAIL"), ("BSP", "BHILAI STEEL PLANT"), ("DSP", "DURGAPUR STEEL PLANT"),
    ("RSP", "ROURKELA STEEL PLANT"), ("BSL", "BOKARO STEEL PLANT"), ("ISP", "IISCO STEEL PLANT"),
]
INTL_BENCHMARKS = ("Coke Rate: 250-300 kg/thm, CDI: 200-250 kg/thm, BF productivity: "
                   "3.0-3.5 t/m<sup>3</sup>/day, specific energy consumption: 4.5-5.27 Gcal/tcs")


def _fmt(v, d):
    return "" if v is None else f"{v:.{d}f}"


def _techno_months_present():
    conn = db.connect()
    try:
        return {r[0] for r in conn.execute("SELECT DISTINCT report_month FROM techno_data").fetchall()}
    finally:
        conn.close()


def _best_ever_annual(plants, params, fy_start):
    """{(plant, param): (value, fy)} -- the best full-FY figure over every FY
    before `fy_start`. A FY counts when techno_data has all twelve of its
    months, or (for the years loaded only as an annual figure) its March,
    whose page-27 till-month cumulative is that FY's annual value."""
    present = _techno_months_present()
    full, march_only = [], []
    first = min((int(m[:4]) for m in present), default=fy_start) - 1
    for fy in range(first, fy_start):
        months = bnc.fy_months(fy)
        if all(m in present for m in months):
            full.append(fy)
        elif months[-1] in present:
            march_only.append(fy)
    names = [p[1] for p in params]
    series = {}   # (plant, param) -> {fy: value}
    if full:
        res = tp.build_period_report(plants, names,
                                     [{"label": str(fy), "months": bnc.fy_months(fy)} for fy in full])
        for sec in res["sections"]:
            for row in sec["rows"]:
                for lbl, cell in row["values"].items():
                    if cell and cell.get("value") is not None:
                        series.setdefault((row["plant"], sec["parameter"]), {})[int(lbl)] = cell["value"]
    if march_only:
        cache, dcache = {}, {}
        for pdef in tp.MAJOR_TECHNO_PARAMS:
            if pdef["display_name"] not in names:
                continue
            for plant in plants:
                for fy in march_only:
                    cell = tp._reported_cum_cell(pdef, plant, bnc.fy_months(fy), cache, dcache)
                    if cell and cell.get("value") is not None:
                        series.setdefault((plant, pdef["display_name"]), {})[fy] = cell["value"]
    better = {p[1]: p[3] for p in params}
    dec = {p[1]: p[2] for p in params}
    out = {}
    for (plant, name), by_fy in series.items():
        # Compare at the printed precision, so a tie goes to the later year.
        vals = {fy: round(v, dec[name]) for fy, v in by_fy.items() if v and v > 0}
        if not vals:
            continue
        pick = max if better[name] == "high" else min
        best = pick(vals.values())
        out[(plant, name)] = (best, max(fy for fy, v in vals.items() if v == best))
    return out


def _techno_targets(fy_start, report_month):
    fy = f"{fy_start}-{_yy(fy_start + 1)}"
    targets = {}
    for plant in _P5:
        targets[plant] = {k[1]: v for k, v in pt._get_plant_techno_plan_targets(plant, report_month).items()}
    try:
        targets["SAIL"] = {k[1]: v for k, v in pt.compute_sail_targets(fy).items()}
    except Exception:
        targets["SAIL"] = {}
    return targets


def _period_cols(P):
    cols = [(f"{_MON[int(m[5:7]) - 1]}'{m[2:4]}", [m]) for m in P["months"]]
    cols += [(P["q_label"], P["q"]), (P["q_cply_label"], P["q_cply"])]
    if P["cum"]:
        cols += [(P["long_label"], P["cum"]), (P["long_cply_label"], P["cum_cply"])]
    return cols


def annexure_ii(P) -> dict:
    plants = [p for p, _ in ANNEX_II_PLANTS]
    cols = _period_cols(P)
    res = tp.build_period_report(plants, [p[1] for p in ANNEX_II_PARAMS],
                                 [{"label": l, "months": m} for l, m in cols])
    by = {(row["plant"], sec["parameter"]): row["values"] for sec in res["sections"] for row in sec["rows"]}
    best = _best_ever_annual(plants, ANNEX_II_PARAMS, P["fy_start"])
    targets = _techno_targets(P["fy_start"], P["months"][-1])
    groups = []
    for plant, title in ANNEX_II_PLANTS:
        rows = []
        for label, name, d, _b in ANNEX_II_PARAMS:
            b = best.get((plant, name))
            vals = [f"{_fmt(b[0], d)} ({_fy_label(b[1])})" if b else "",
                    _fmt((targets.get(plant) or {}).get(name), d)]
            cells = by.get((plant, name), {})
            vals += [_fmt((cells.get(l) or {}).get("value"), d) for l, _m in cols]
            rows.append({"label": label, "values": vals})
        groups.append({"title": title, "rows": rows})
    return {"tag": "II", "fy": _fy_label(P["fy_start"]), "cols": [l for l, _ in cols],
            "title": f"SAIL: Techno-economic parameters during {P['q_label']}"
                     + (f" and {P['long_label']}" if P["cum"] else ""),
            "groups": groups}


# --------------------------------------------------------------------------
# Annexure-III: new blast furnaces
# --------------------------------------------------------------------------

ANNEX_III_FURNACES = [
    ("BF-8 at BSP", "BSP", "BF-8"),
    ("BF-5 at RSP", "RSP", "BF-5"),
    ("BF-1 at RSP", "RSP", "BF-1"),
    ("New Blast furnace at ISP", "ISP", "BF-5"),
]
ANNEX_III_PARAMS = [
    ("Coke Rate (kg/thm)", "coke_rate", 0),
    ("CDI (kg/thm)", "cdi", 0),
    ("BF Productivity (t/m3/day)", "bf_productivity", 2),
]


def annexure_iii(P) -> dict:
    fy = f"{P['fy_start']}-{_yy(P['fy_start'] + 1)}"
    cols = _period_cols(P)
    weight_cache, month_cache = {}, {}
    groups = []
    for title, plant, unit in ANNEX_III_FURNACES:
        try:
            tgt = (db.get_techno_plan(plant, fy, unit) or {}).get("data") or {}
        except Exception:
            tgt = {}
        rows = []
        for label, key, d in ANNEX_III_PARAMS:
            t = tgt.get(key)
            t = t.get("value") if isinstance(t, dict) else t
            vals = [_fmt(t, d)]
            for _l, months in cols:
                cell = tbp._range_cell(plant, unit, key, months, weight_cache, month_cache)
                vals.append(_fmt(cell.get("value"), d))
            rows.append({"label": label, "values": vals})
        groups.append({"title": title, "rows": rows})
    return {"tag": "III", "fy": _fy_label(P["fy_start"]), "cols": [l for l, _ in cols],
            "title": "SAIL: Techno-economic parameters for New Blast Furnaces",
            "subtitle": f"during {P['q_label']}" + (f" and {P['long_label']}" if P["cum"] else ""),
            "groups": groups}


# --------------------------------------------------------------------------
# Annexure-IV: product mix
# --------------------------------------------------------------------------

# Rows shown per compute_pmix_fy_rows block, in the annexure's own wording.
# label -> (shown as, always). always=False rows are printed only when they
# have a figure (mills the Q-2'25-26 annexure had no production for).
_V_ROWS = {
    "BSP": [("WRM", "WRM"), ("TMT Bars (MM)",), ("Lt Strls (MM)",), ("Rounds & Bars (BRM)",),
            ("Rail (RSM)",), ("Rail (URM)",), ("Plates",), ("Semis Total",), ("Saleable Steel",)],
    "DSP": [("Med Structurals (SM)",), ("Med Structurals (MSM)",), ("Round & bars (MM)",),
            ("Wheel & Axle",), ("Semis Total",), ("Saleable steel",)],
    "RSP": [("Plate (PM)",), ("New Plate Mill",), ("New HR Plate",), ("New HR Coils",),
            ("CRNO",), ("ERW Pipes",), ("SW Pipes",), ("Saleable Steel",)],
    "BSL": [("HR Plates",), ("HR Coils",), ("HR Sheets",), ("CR Coils (old)",), ("CR Coils (new)",),
            ("CR Sheet (old)",), ("CR Sheet (new)", None, False), ("GP/GC (old)",), ("GP/GC (new)",),
            ("Thick plate", None, False), ("Semis",), ("Saleable Steel",)],
    "ISP": [("Round & Bars(Bar Mill)",), ("Hy Structurals(USM)",), ("WRM",), ("Semis",),
            ("Saleable Steel",)],
    "SAIL": [("Semis",), ("Wire Rods",), ("Round & Bars",), ("Lt. Structurals",), ("Med. Structurals",),
             ("Hy. Structurals",), ("Rail",), ("Wheel & Axle",), ("Plate",), ("Thick plate", None, False),
             ("HR Plates",), ("HR Sheets",), ("HR Coils/Skelp",), ("CR Coils/Sheets",), ("GP/GC",),
             ("CRNO",), ("ERW pipes",), ("SW Pipes",), ("Saleable steel 5 PL",)],
    "Finished Steel": None,   # every row
}
_V_PAGE_BREAK_BEFORE = "ISP"


def _pmix_month(cur, month, table, cache):
    k = (month, table)
    if k not in cache:
        rows, block = {}, None
        for r in compute_pmix_fy_rows(cur, month, table):
            if r["kind"] == "header":
                block = r["label"]
                rows[block] = []
            else:
                rows[block].append(r)
        cache[k] = rows
    return cache[k]


def _pmix_sum(cur, months, table, cache):
    """{block: [(label, value, bold), ...]} summed over `months`."""
    out = {}
    for m in months:
        for block, rows in _pmix_month(cur, m, table, cache).items():
            acc = out.setdefault(block, [[r["label"], None, r["bold"]] for r in rows])
            for i, r in enumerate(rows):
                if r["value"] is not None:
                    acc[i][1] = (acc[i][1] or 0) + r["value"]
    return out


def annexure_iv(cur, P) -> dict:
    cache = {}
    cols = [("fy", P["fy_months"], ["plan"])]
    if P["prev"]:
        cols.append(("prev", P["prev"], ["plan", "act", "ful"]))
    for m in P["months"]:
        cols.append((m, [m], ["plan", "act", "ful"]))
    cols.append(("q", P["q"], ["plan", "act", "ful"]))
    cols.append(("q_cply", P["q_cply"], ["cply"]))
    if P["cum"]:
        cols.append(("cum", P["cum"], ["plan", "act", "ful"]))
        cols.append(("cum_cply", P["cum_cply"], ["cply"]))

    sums = {}
    for key, months, kinds in cols:
        if "plan" in kinds:
            sums[(key, "plan")] = _pmix_sum(cur, months, "production_plan_table", cache)
        if "act" in kinds or "cply" in kinds:
            sums[(key, "act")] = _pmix_sum(cur, months, "production_table", cache)

    def val(key, kind, block, i):
        rows = sums.get((key, kind), {}).get(block) or []
        return rows[i][1] if i < len(rows) else None

    ref = _pmix_sum(cur, [P["months"][0]], "production_table", cache)
    groups = []
    for block, spec in _V_ROWS.items():
        rows_ref = ref.get(block) or []
        by_label = {r[0]: i for i, r in enumerate(rows_ref)}
        wanted = ([(r[0], r[0], True) for r in rows_ref] if spec is None
                  else [(s[0], s[1] if len(s) > 1 and s[1] else s[0], s[2] if len(s) > 2 else True)
                        for s in spec])
        rows = []
        for src, shown, always in wanted:
            i = by_label.get(src)
            if i is None:
                continue
            vals = []
            for key, _m, kinds in cols:
                plan, act = val(key, "plan", block, i), val(key, "act", block, i)
                for k in kinds:
                    if k == "plan":
                        vals.append(plan)
                    elif k == "act":
                        vals.append(act)
                    elif k == "ful":
                        vals.append(_ful(act, plan))
                    elif k == "cply":
                        vals.append(act)
                if key in ("q_cply", "cum_cply"):
                    cur_key = "q" if key == "q_cply" else "cum"
                    vals.append(_gr(val(cur_key, "act", block, i), act))
            if not always and not any(v for v in vals):
                continue
            bold = rows_ref[i][2] or shown.lower().startswith(("saleable", "finished steel"))
            rows.append({"label": shown, "values": vals, "bold": bold})
        groups.append({"title": block, "rows": rows, "break": block == _V_PAGE_BREAK_BEFORE})

    fyl = f"{P['fy_start']}-{_yy(P['fy_start'] + 1)}"
    head = [(fyl, ["ABP"])]
    if P["prev"]:
        head.append((P["prev_label"], ["ABP", "Actual", "% Ful."]))
    head += [(_mon_full(m), ["ABP", "Actual", "% Ful."]) for m in P["months"]]
    head.append((P["q_label"], ["ABP", "Actual", "% Ful."]))
    head.append((P["q_cply_label"], ["Actual"]))
    head.append(("% Gr. over", [P["q_cply_label"]]))
    if P["cum"]:
        head.append((P["cum_label"], ["ABP", "Actual", "% Ful."]))
        head.append((P["cum_cply_label"], ["Actual"]))
        head.append(("% Gr. over", [P["cum_cply_label"]]))
    return {"tag": "IV", "head": head, "groups": groups,
            "title": f"SAIL Product mix performance: {P['q_label']}"
                     + (f" and {P['cum_label']}" if P["cum"] else "")}


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

def _css(colors):
    def c(k, d):
        return colors.get(k, d)
    blk = c("border_black", "#000000")
    txt = c("text_black", "#000000")
    return f"""
  @page {{ size: A4; margin: 10mm 10mm 10mm 10mm; }}
  @page land {{ size: A4 landscape; margin: 8mm 8mm 8mm 8mm; }}
  body {{ font-family: Calibri, Carlito, Arial, sans-serif; margin: 0; color: {txt};
         -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
  .sheet {{ break-after: page; }}
  .sheet:last-child {{ break-after: auto; }}
  .sheet.land {{ page: land; }}
  .tag {{ text-align: right; font-size: 9pt; font-weight: 700; }}
  h1 {{ font-size: 13pt; font-weight: 700; margin: 4px 0 2px; }}
  h1.c {{ text-align: center; margin-top: 18px; }}
  h2 {{ font-size: 12pt; font-weight: 700; text-align: center; margin: 0 0 8px; }}
  .unit {{ text-align: right; font-size: 7pt; font-weight: 700; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid {blk}; padding: 1px 3px; line-height: 1.15; white-space: nowrap; }}
  th {{ font-weight: 700; text-align: center; }}
  td {{ text-align: center; }}
  td.l {{ text-align: left; }}
  tr.b td {{ font-weight: 700; }}
  tr.g td {{ font-weight: 700; text-align: left; }}
  .t1 th, .t1 td {{ font-family: 'Arial Narrow', Arial, sans-serif; font-size: 8pt; }}
  .t1 td {{ border-top: none; border-bottom: none; }}
  .t1 tr.b td, .t1 tr.g td {{ border-top: 1px solid {blk}; border-bottom: 1px solid {blk}; }}
  .t1 tbody tr:last-child td {{ border-bottom: 1px solid {blk}; }}
  .t2 th, .t2 td {{ font-size: 10.5pt; }}
  .t2 td:first-child {{ text-align: left; }}
  .t2 th:first-child {{ text-align: left; }}
  .t2 tr.g td {{ text-align: center; padding-left: 0; }}
  .t2 tr.g td:first-child {{ text-align: center; }}
  .t5 th, .t5 td {{ font-family: 'Arial Narrow', Arial, sans-serif; font-size: 8pt; }}
  .t5 th {{ font-size: 7pt; }}
  .t5 td {{ border-top: none; border-bottom: 1px dashed {blk}; }}
  .t5 td.l {{ text-align: left; }}
  .t5 td:not(.l) {{ text-align: center; }}
  .note {{ font-size: 10.5pt; font-style: italic; margin: 2px 0 0 50px; }}
"""


def _sheet_i(s):
    nb = len(s["blocks"])
    h1 = ('<tr><th rowspan="2">Plant</th><th>APP</th>'
          + "".join(f'<th colspan="3">{escape(l)}</th><th>{escape(c)}</th><th>% Gr. over</th>'
                    for l, c in s["blocks"]) + "</tr>")
    h2 = (f"<tr><th>FY'{s['fy']}</th>"
          + "".join(f'<th>APP</th><th>Actual</th><th>% Ful.</th><th>Actual</th><th>{escape(c)}</th>'
                    for _l, c in s["blocks"]) + "</tr>")
    body = []
    for g in s["groups"]:
        body.append(f'<tr class="g"><td colspan="{2 + 5 * nb}">{escape(g["title"])}</td></tr>')
        for r in g["rows"]:
            cls = ' class="b"' if r["bold"] else ""
            body.append(f"<tr{cls}><td>{escape(r['label'])}</td>"
                        + "".join(f"<td>{_i(v)}</td>" for v in r["values"]) + "</tr>")
    return (f'<div class="sheet"><div class="tag">ANNEXURE-{s["tag"]}</div>'
            f'<h1>{escape(s["title"])}</h1><div class="unit">Unit: \'000 T</div>'
            f'<table class="t1"><thead>{h1}{h2}</thead><tbody>{"".join(body)}</tbody></table></div>')


def _br(label):
    """"Q-2'25-26" -> "Q-2<br>'25-26"; month labels ("Jul'25") stay on one line."""
    head, sep, tail = label.partition("'")
    return escape(label) if not sep or "-" not in head else f"{escape(head)}<br>'{escape(tail)}"


def _sheet_ii(s):
    cols = s["cols"]
    fy = s["fy"]
    head = ('<tr><th>Parameter</th><th>Best Ever<br>Annual</th>'
            f"<th>20{fy[:2]}-<br>{fy[3:]}<br>Target</th>"
            + "".join(f"<th>{_br(c)}</th>" for c in cols) + "</tr>")
    body = []
    n = 3 + len(cols)
    for g in s["groups"]:
        body.append(f'<tr class="g"><td>{escape(g["title"])}</td>' + "<td></td>" * (n - 1) + "</tr>")
        for r in g["rows"]:
            body.append(f"<tr><td>{escape(r['label'])}</td>"
                        + "".join(f"<td>{escape(v)}</td>" for v in r["values"]) + "</tr>")
    return (f'<div class="sheet"><div class="tag">ANNEXURE-II</div><h1 class="c">{escape(s["title"])}</h1>'
            f'<table class="t2"><thead>{head}</thead><tbody>{"".join(body)}</tbody></table>'
            f'<div class="note">International Benchmarks:<br>{INTL_BENCHMARKS}</div></div>')


def _sheet_iii(s):
    cols = s["cols"]
    fy = s["fy"]
    head = (f"<tr><th>Parameter</th><th>Target<br>FY {fy}</th>"
            + "".join(f"<th>{_br(c)}</th>" for c in cols) + "</tr>")
    body = []
    n = 2 + len(cols)
    for g in s["groups"]:
        body.append(f'<tr class="g"><td>{escape(g["title"])}</td>' + "<td></td>" * (n - 1) + "</tr>")
        for r in g["rows"]:
            body.append(f"<tr><td>{escape(r['label'])}</td>"
                        + "".join(f"<td>{escape(v)}</td>" for v in r["values"]) + "</tr>")
    return (f'<div class="sheet"><div class="tag">ANNEXURE-III</div><h1 class="c">{escape(s["title"])}<br>'
            f'{escape(s["subtitle"])}</h1>'
            f'<table class="t2"><thead>{head}</thead><tbody>{"".join(body)}</tbody></table></div>')


def _sheet_iv(s):
    head = s["head"]
    h1 = '<tr><th rowspan="2"></th>' + "".join(
        f'<th colspan="{len(subs)}">{escape(l)}</th>' for l, subs in head) + "</tr>"
    h2 = "<tr>" + "".join(f"<th>{escape(x)}</th>" for _l, subs in head for x in subs) + "</tr>"
    ncol = 1 + sum(len(subs) for _l, subs in head)
    pages, cur = [], []
    for g in s["groups"]:
        if g["break"] and cur:
            pages.append(cur)
            cur = []
        cur.append(f'<tr class="g"><td class="l" colspan="{ncol}">{escape(g["title"])}</td></tr>')
        for r in g["rows"]:
            cls = ' class="b"' if r["bold"] else ""
            cur.append(f'<tr{cls}><td class="l">{escape(r["label"])}</td>'
                       + "".join(f"<td>{_i(v)}</td>" for v in r["values"]) + "</tr>")
    pages.append(cur)
    out = []
    for body in pages:
        out.append(f'<div class="sheet land"><div class="tag">ANNEXURE-IV</div><h2>{escape(s["title"])}</h2>'
                   f"<div class=\"unit\">Unit: '000T</div>"
                   f'<table class="t5"><thead>{h1}{h2}</thead><tbody>{"".join(body)}</tbody></table></div>')
    return "".join(out)


def collect(fy: str, quarter: int, which=ANNEXURES) -> list:
    """[(annexure, sheet data), ...] for the chosen annexures, in I, II,
    III, IV order -- shared by the PDF (build) and the workbook (build_xlsx)."""
    P = periods(int(fy[:4]), quarter)
    out = []
    conn = db.connect()
    try:
        cur = conn.cursor()
        if "I" in which:
            out += [("I", s) for s in annexure_i(cur, P)]
        if "II" in which:
            out.append(("II", annexure_ii(P)))
        if "III" in which:
            out.append(("III", annexure_iii(P)))
        if "IV" in which:
            out.append(("IV", annexure_iv(cur, P)))
    finally:
        conn.close()
    return out


_HTML_SHEET = {"I": _sheet_i, "II": _sheet_ii, "III": _sheet_iii, "IV": _sheet_iv}


def build(fy: str, quarter: int, which=ANNEXURES) -> str:
    """HTML for the chosen annexures (in I, II, III, IV order)."""
    from colors_loader import load_colors_config
    parts = [_HTML_SHEET[a](s) for a, s in collect(fy, quarter, which)]
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{_css(load_colors_config())}</style>'
            f'</head><body>{"".join(parts)}</body></html>')


# --------------------------------------------------------------------------
# Excel: one worksheet per annexure sheet, same rows and headings as the PDF
# --------------------------------------------------------------------------

def _num_or_text(v):
    """Annexures II/III carry formatted strings; write real numbers where the
    string is one ("2.04" -> 2.04), keep "421 (24-25)" as text."""
    if v is None or v == "":
        return None, None
    if isinstance(v, (int, float)):
        return round(v), "0"
    try:
        f = float(v)
    except ValueError:
        return v, None
    d = len(v.split(".")[1]) if "." in v else 0
    return f, "0" if d == 0 else "0." + "0" * d


def build_xlsx(fy: str, quarter: int, which=ANNEXURES) -> bytes:
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def start(name, tag, title, ncol, subtitle=None):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=ncol, value=f"ANNEXURE-{tag}").font = bold
        ws.cell(row=1, column=ncol).alignment = Alignment(horizontal="right")
        ws.cell(row=2, column=1, value=title).font = Font(bold=True, size=13)
        r = 3
        if subtitle:
            ws.cell(row=3, column=1, value=subtitle).font = Font(bold=True, size=13)
            r = 4
        return ws, r

    def hdr(ws, r, c, text, span=1, rows=1):
        cell = ws.cell(row=r, column=c, value=text)
        cell.font, cell.alignment = bold, center
        for rr in range(r, r + rows):
            for cc in range(c, c + span):
                ws.cell(row=rr, column=cc).border = box
        if span > 1 or rows > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r + rows - 1, end_column=c + span - 1)

    def put(ws, r, c, v, is_bold=False, align=center):
        val, fmt = _num_or_text(v)
        cell = ws.cell(row=r, column=c, value=val)
        cell.border, cell.alignment = box, align
        if fmt:
            cell.number_format = fmt
        if is_bold:
            cell.font = bold

    def group_row(ws, r, title, ncol, align=left):
        for c in range(1, ncol + 1):
            ws.cell(row=r, column=c).border = box
        cell = ws.cell(row=r, column=1, value=title)
        cell.font, cell.alignment = bold, align

    def body(ws, r, groups, ncol, gal=left, lal=left):
        for g in groups:
            group_row(ws, r, g["title"], ncol, gal)
            r += 1
            for row in g["rows"]:
                b = row.get("bold", False)
                put(ws, r, 1, row["label"], b, lal)
                for i, v in enumerate(row["values"]):
                    put(ws, r, 2 + i, v, b)
                r += 1
        return r

    for tag, s in collect(fy, quarter, which):
        if tag == "I":
            nb = len(s["blocks"])
            ncol = 2 + 5 * nb
            ws, r = start(f"Annexure-{s['tag']}", s["tag"], s["title"], ncol)
            ws.cell(row=r, column=ncol, value="Unit: '000 T").font = bold
            r += 1
            hdr(ws, r, 1, "Plant", rows=2)
            hdr(ws, r, 2, "APP")
            hdr(ws, r + 1, 2, f"FY'{s['fy']}")
            c = 3
            for lbl, cply in s["blocks"]:
                hdr(ws, r, c, lbl, span=3)
                hdr(ws, r, c + 3, cply)
                hdr(ws, r, c + 4, "% Gr. over")
                for i, t in enumerate(["APP", "Actual", "% Ful.", "Actual", cply]):
                    hdr(ws, r + 1, c + i, t)
                c += 5
            first = r + 2
            r = body(ws, first, s["groups"], ncol, lal=center)
            widths = [10, 9] + [8] * (5 * nb)
        elif tag in ("II", "III"):
            fixed = (["Parameter", "Best Ever\nAnnual", f"20{s['fy'][:2]}-{s['fy'][3:]}\nTarget"]
                     if tag == "II" else ["Parameter", f"Target\nFY {s['fy']}"])
            heads = fixed + s["cols"]
            ncol = len(heads)
            ws, r = start(f"Annexure-{tag}", tag, s["title"], ncol, s.get("subtitle"))
            r += 1
            for i, h in enumerate(heads):
                hdr(ws, r, 1 + i, h)
            ws.row_dimensions[r].height = 45
            first = r + 1
            r = body(ws, first, s["groups"], ncol, gal=center)
            if tag == "II":
                ws.cell(row=r, column=1, value="International Benchmarks:").font = Font(italic=True)
                ws.cell(row=r + 1, column=1, value=INTL_BENCHMARKS.replace("<sup>3</sup>", "3")).font = Font(italic=True)
            widths = [40 if tag == "II" else 30] + ([14] if tag == "II" else []) + [10] * (ncol - (2 if tag == "II" else 1))
        else:  # IV, product mix
            ncol = 1 + sum(len(subs) for _l, subs in s["head"])
            ws, r = start("Annexure-IV", "IV", s["title"], ncol)
            ws.cell(row=r, column=ncol, value="Unit: '000T").font = bold
            r += 1
            hdr(ws, r, 1, "", rows=2)
            c = 2
            for lbl, subs in s["head"]:
                hdr(ws, r, c, lbl, span=len(subs))
                for i, t in enumerate(subs):
                    hdr(ws, r + 1, c + i, t)
                c += len(subs)
            first = r + 2
            r = body(ws, first, s["groups"], ncol)
            widths = [28] + [8] * (ncol - 1)
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.freeze_panes = ws.cell(row=first, column=2)
        ws.page_setup.orientation = "landscape" if tag == "IV" else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Shrink any sheet that would overflow its page (Annexure-IV's wide columns
# in particular) so each sheet stays one printed page.
_FIT_JS = """() => {
    const mm = 96 / 25.4;
    document.querySelectorAll('.sheet').forEach(s => {
        const land = s.classList.contains('land');
        const W = (land ? 297 - 16 : 210 - 20) * mm, H = (land ? 210 - 16 : 297 - 20) * mm;
        s.style.width = W + 'px';
        const t = s.querySelector('table');
        const r = s.getBoundingClientRect();
        const z = Math.min(1, W / Math.max(t.scrollWidth, r.width), H / r.height) * 0.995;
        if (z < 1) s.style.zoom = z;
    });
}"""


def render(html: str) -> bytes:
    """Synchronous -- call from a threadpool."""
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(viewport={"width": 1200, "height": 1000})
        page.emulate_media(media="print")
        page.set_content(html, wait_until="load")
        page.evaluate("document.fonts.ready.then(() => true)")
        page.evaluate(_FIT_JS)
        pdf = page.pdf(prefer_css_page_size=True, print_background=True)
        browser.close()
    return pdf
