"""
Excel / PDF export for the Blast Furnace Techno Report (frontend:
/reports/techno-bf-furnace) — furnace-wise custom-period/month-and-till-
month/annual data, from techno_bf_period.build_range_report /
build_direct_report's {periods, furnaces, sections} shape.

Visual style mirrors page_techno_custom_export.py's Custom Period mode
(blue header, light-blue subheader, zebra rows) for consistency with the
rest of the app; `render_pdf_bytes` is imported from
page_production_query_export.py directly, same as every other export
module here, since it has zero business logic (pure HTML-string-in,
PDF-bytes-out).
"""
import io

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import page_production_query_export as _ppqe

render_pdf_bytes = _ppqe.render_pdf_bytes

_HDR_FILL = PatternFill("solid", fgColor="1A73E8")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
_SUBHDR_FILL = PatternFill("solid", fgColor="E8F0FE")
_SUBHDR_FONT = Font(bold=True, color="174EA6", size=9)
_SECTION_FILL = PatternFill("solid", fgColor="1A73E8")
_SECTION_FONT = Font(bold=True, color="FFFFFF", size=10)
_ZEBRA_FILL = PatternFill("solid", fgColor="F8F9FA")
_THIN = Side(style="thin", color="DADCE0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _cell_str(v):
    if v is None or v == "":
        return ""
    return str(v)


def build_furnace_excel_bytes(data: dict, subtitle: str = "") -> bytes:
    periods = data.get("periods", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BF Techno Report"[:31]

    ws.cell(row=1, column=1, value="Blast Furnace Techno Report").font = Font(bold=True, size=13)
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = Font(italic=True, size=9)

    total_cols = 1 + len(periods)
    row = 4
    for sec in data.get("sections", []):
        title = f'{sec.get("parameter","")} ({sec.get("unit","")})' if sec.get("unit") else sec.get("parameter", "")
        sc = ws.cell(row=row, column=1, value=title)
        sc.font = _SECTION_FONT
        sc.fill = _SECTION_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
        for c in range(2, total_cols + 1):
            ws.cell(row=row, column=c).fill = _SECTION_FILL
        row += 1

        headers = ["Furnace"] + periods
        for c, h in enumerate(headers, start=1):
            hc = ws.cell(row=row, column=c, value=h)
            hc.font = _SUBHDR_FONT
            hc.fill = _SUBHDR_FILL
            hc.alignment = Alignment(horizontal="center")
            hc.border = _BORDER
        row += 1

        for idx, r in enumerate(sec.get("rows", [])):
            fill = _ZEBRA_FILL if idx % 2 == 1 else None
            font = Font(size=9)
            fc = ws.cell(row=row, column=1, value=r.get("furnace", ""))
            fc.font = font
            fc.border = _BORDER
            if fill:
                fc.fill = fill
            for c, label in enumerate(periods, start=2):
                cell_data = (r.get("values") or {}).get(label) or {}
                disp = _cell_str(cell_data.get("display"))
                if cell_data.get("method_used") == "average" and cell_data.get("warnings"):
                    disp = f"{disp}*" if disp else disp
                vc = ws.cell(row=row, column=c, value=disp)
                vc.font = font
                vc.border = _BORDER
                vc.alignment = Alignment(horizontal="right")
                if fill:
                    vc.fill = fill
            row += 1
        row += 1  # blank spacer row between sections

    ws.column_dimensions["A"].width = 16
    for i in range(2, total_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_furnace_pdf_html(data: dict, subtitle: str = "") -> str:
    periods = data.get("periods", [])
    header_html = "<th>Furnace</th>" + "".join(f"<th>{p}</th>" for p in periods)

    sections_html = []
    for sec in data.get("sections", []):
        body_rows = []
        for r in sec.get("rows", []):
            cells = [f'<td class="furnace">{r.get("furnace","")}</td>']
            for label in periods:
                cell_data = (r.get("values") or {}).get(label) or {}
                disp = cell_data.get("display") or "—"
                if cell_data.get("method_used") == "average" and cell_data.get("warnings"):
                    disp = f"{disp}*"
                cells.append(f"<td>{disp}</td>")
            body_rows.append(f'<tr>{"".join(cells)}</tr>')
        title = f'{sec.get("parameter","")} ({sec.get("unit","")})' if sec.get("unit") else sec.get("parameter", "")
        sections_html.append(
            f'<div class="section-title">{title}</div>'
            f'<table><thead><tr>{header_html}</tr></thead>'
            f'<tbody>{"".join(body_rows)}</tbody></table>'
        )

    return f"""<!doctype html>
<html><head><meta charset="utf-8">
<style>
  @page {{ size: A4 landscape; margin: 12mm 10mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: Arial, sans-serif; color: #202124; margin: 0; }}
  h1 {{ font-size: 14pt; margin: 0 0 2px 0; }}
  .subtitle {{ font-size: 9pt; color: #5f6368; margin: 0 0 8px 0; }}
  .section-title {{ background: #1a73e8; color: #fff; font-weight: 700; font-size: 9.5pt;
    padding: 4px 6px; margin-top: 10px; page-break-after: avoid; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 8pt; margin-bottom: 4px; }}
  thead {{ display: table-header-group; }}
  tr {{ page-break-inside: avoid; }}
  th, td {{ border: 1px solid #dadce0; padding: 3px 6px; text-align: right; white-space: nowrap; }}
  th {{ background: #e8f0fe; color: #174ea6; font-weight: 700; }}
  td.furnace {{ text-align: left; font-weight: 700; }}
  tr:nth-child(even) td {{ background: #f8f9fa; }}
</style>
</head>
<body>
  <h1>Blast Furnace Techno Report</h1>
  {f'<p class="subtitle">{subtitle}</p>' if subtitle else ''}
  {''.join(sections_html)}
</body></html>"""
