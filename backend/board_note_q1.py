"""Q-1 Board Note docx generator (quarter-only shape).

Fills `board_note_templates/board_note_q1_template.docx` (a verbatim copy of
`Report_format/BN_ Production Q-1_26-27.docx`) for any financial year. Same
technique as `board_note_q2.py`, but the Q-1 note has no H-1 pass, no
Finished-Steel-vs-MoU opening, no annexure-table MoU block, and its techno
table carries Specific Energy Consumption instead of Fuel Rate, with no
improvement-commentary paragraphs under it.

Index map: derived from `docx.Document(...).paragraphs` / `.tables` for this
exact template file. If the template is hand-edited, re-derive these.

Known limitations (documented, not bugs):
  - DSP and ISP have no "Highlights:" section in this template, and neither
    do ASP/SSP/VISL, so their best-ever bullets are not written. There is no
    slot to host them without restructuring the document.
  - BSL's first highlight paragraph (a static SMS-1 "new grade" note with no
    DB source) is treated as a bullet slot and overwritten, the same way the
    other static bullets are. Enter such notes in the manual additional
    highlights box.
  - Plant headings, the "Submitted for consideration" block and the
    signatory are left static, as in the Q-2 adapter.
"""

import io

import docx

import board_note_common as bnc
import board_note_docx_utils as bdu
import board_note_manual_text as bnm
import board_note_q2 as bq2
import db
import page_techno as pt
import techno_period as tp

TEMPLATE_PATH = "board_note_templates/board_note_q1_template.docx"

# -- Title / intro -------------------------------------------------------------
P_TITLE = 3                  # "PRODUCTION PERFORMANCE DURING QUARTER-1’26-27 (APRIL’26, MAY’26, JUNE’26) "
P_INTRO = 6                  # "SAIL and Plant-wise Production performance during Q-1’26-27 w.r.t. APP"
P_PLANT_SECTION_TITLE = 26   # "Plant wise: Production performance during Q-1’26-27 "

# -- Tables (index into docx.Document.tables) ---------------------------------
TABLE_SAIL = 0               # ABP / Actual / %Ful / CPLY / %Gr, rows: HM, CS, SS, FS
TABLE_TECHNO = 9             # Parameters | Plant | Target | Q-1cur | Q-1cply
TABLE_ANNEXURE_INDEX = 10

# -- SAIL -------------------------------------------------------------------
P_SAIL_HEADING = 8           # "SAIL: Production performance during Q-1’26-27 "
P_SAIL_LABEL = 10            # "Q-1’26-27:"
P_SAIL_SUMMARY = 11          # "Hot Metal production during Q-1’26-27 was ... "
P_SAIL_HIGHLIGHTS_HEADING = 15
P_SAIL_HIGHLIGHTS_LEADIN = 16  # "SAIL achieved best ever April & June month production for following:"

# -- Techno section ----------------------------------------------------------
P_TECHNO_SECTION_TITLE = 252   # "SAIL and Plant-wise Major Techno-economic parameters performance during Q-1’26-27:"

TECHNO_PLANTS = ["BSP", "DSP", "RSP", "BSL", "ISP", "SAIL"]
TECHNO_PARAMS = ["Coke Rate", "CDI Rate", "BF Productivity", "Specific Energy Consumption"]

# -- Item-row configs for fill_period_table (row_idx, display, db_item, add_conv) --
SAIL_ITEMS = bq2.SAIL_ITEMS
BIG5_ITEMS = [
    (2, "Hot Metal", "Hot Metal", False),
    (3, "Crude Steel", "Total Crude Steel", False),
    (4, "Saleable Steel", "Saleable Steel", False),
]
SPECIAL2_ITEMS = [
    (2, "Crude Steel", "Total Crude Steel", False),
    (3, "Saleable Steel", "Saleable Steel", False),
]
VISL_ITEMS = [
    (2, "Saleable Steel", "Saleable Steel", False),
]

# -- Per-plant index map: unit table, unit label + summary paragraphs --------
PLANT_CONFIG = {
    "BSP": dict(table=1, label=29, summary=30, item_rows=BIG5_ITEMS, kind="big"),
    "DSP": dict(table=2, label=69, summary=70, item_rows=BIG5_ITEMS, kind="big"),
    "RSP": dict(table=3, label=97, summary=98, item_rows=BIG5_ITEMS, kind="big"),
    "BSL": dict(table=4, label=133, summary=134, item_rows=BIG5_ITEMS, kind="big"),
    "ISP": dict(table=5, label=187, summary=188, item_rows=BIG5_ITEMS, kind="big"),
    "ASP": dict(table=6, label=211, summary=212, item_rows=SPECIAL2_ITEMS, kind="small"),
    "SSP": dict(table=7, label=225, summary=226, item_rows=SPECIAL2_ITEMS, kind="small"),
    "VISL": dict(table=8, label=243, summary=244, item_rows=VISL_ITEMS, kind="visl"),
}
PLANT_ORDER = ["BSP", "DSP", "RSP", "BSL", "ISP", "ASP", "SSP", "VISL"]

# -- Highlights bullet slots (every non-empty paragraph under the lead-in) ---
# DSP, ISP, ASP, SSP and VISL have no "Highlights:" section in this template.
HIGHLIGHTS_SLOTS = {
    ("SAIL", "q1"): [17],
    ("BSP", "q1"): [36, 37, 38, 39, 40, 41, 42, 43, 44, 45],
    ("RSP", "q1"): [104, 105, 106, 107],
    ("BSL", "q1"): [138, 140, 141, 142, 143, 144, 145, 146, 147, 148, 149, 150, 151, 152],
}
# Empty spacer paragraphs inside a highlights block that would otherwise sit
# between two generated bullets and show as a blank line in the middle of the
# list. Removed before the bullets are filled.
HIGHLIGHTS_BLANKS_REMOVE = {
    ("BSL", "q1"): [139],
}
# First lead-in paragraph, rewritten to the Q-1 wording where the template's
# own lead-in is a month-level claim ("April & June").
LEADIN_TEXT_FIX = {
    ("SAIL", "q1"): (P_SAIL_HIGHLIGHTS_LEADIN, "SAIL achieved best ever Q-1 production for following:"),
}
HIGHLIGHTS_HEADING = {
    ("SAIL", "q1"): P_SAIL_HIGHLIGHTS_HEADING, ("BSP", "q1"): 34,
    ("RSP", "q1"): 102, ("BSL", "q1"): 137,
}
HIGHLIGHTS_LEADIN = {
    ("SAIL", "q1"): P_SAIL_HIGHLIGHTS_LEADIN, ("BSP", "q1"): 35,
    ("RSP", "q1"): 103, ("BSL", "q1"): None,   # BSL has no lead-in sentence in the template
}

# -- "Production Performance during Q-1’26-27:" why blocks (one per unit) ---
WHY_HEADING = {
    "BSP": 47, "DSP": 78, "RSP": 109, "BSL": 154, "ISP": 192,
    "ASP": 215, "SSP": 230, "VISL": 248,
}
WHY_SLOTS = {
    "BSP": [49, 50, 51, 52, 53, 55, 56, 57, 58, 60, 61, 63, 64, 65],
    "DSP": [80, 82, 83, 85, 86, 88, 89, 90, 91, 92],
    "RSP": [111, 112, 113, 114, 115, 116, 119, 120, 121, 123, 124, 125, 126, 127, 128],
    "BSL": [156, 157, 159, 160, 162, 163, 164, 165, 166, 167, 168, 169, 170, 171, 172, 173,
            176, 177, 178, 179, 180, 181, 182, 183],
    "ISP": [194, 195, 196, 197, 199, 200, 202, 203],
    "ASP": [217, 218, 219, 221],
    "SSP": [232, 233, 235, 236, 237, 238, 239],
    "VISL": [250],
}


def _title_text(fy_start):
    yy = str(fy_start)[2:]
    zz = str(fy_start + 1)[2:]
    m1, m2, m3 = (bq2._month_name(m) for m in bnc.quarter_months(fy_start, 1))
    return (f"PRODUCTION PERFORMANCE DURING QUARTER-1’{yy}-{zz} "
            f"({m1}’{yy}, {m2}’{yy}, {m3}’{yy}) ")


def _fmt_techno_target(param, v):
    if v is None:
        return None
    return f"{v:.2f}" if param in ("BF Productivity", "Specific Energy Consumption") else str(round(v))


def _set_techno_header(table, fy_start, q1_label, q1_cply_label):
    yy = str(fy_start)[2:]
    zz = str(fy_start + 1)[2:]
    bdu.set_cell(table.rows[0].cells[2], f"Target\n{yy}-{zz}")
    bdu.set_cell(table.rows[1].cells[2], f"{yy}-{zz}")
    for r in (0, 1):
        bdu.set_cell(table.rows[r].cells[3], q1_label)
        bdu.set_cell(table.rows[r].cells[4], q1_cply_label)


def generate(fy: str) -> bytes:
    fy_start = int(fy[:4])

    Q1_CUR = bnc.quarter_months(fy_start, 1)
    Q1_CPLY = bnc.quarter_months(fy_start - 1, 1)

    q1_label = bnc.quarter_label(fy_start, 1)
    q1_cply_label = bq2._cply_label_nl(bnc.quarter_label, fy_start, 1)
    yy = str(fy_start)[2:]

    manual = bnm.get_manual_text(fy, 1)

    conn = db.connect()
    cur = conn.cursor()
    try:
        doc = docx.Document(TEMPLATE_PATH)
        tables = doc.tables
        paras = doc.paragraphs

        def para(i):
            return paras[i]

        def apply_highlights(plant, bullets):
            """Fill a unit's highlights bullet slots, then remove the
            "Highlights:" heading and lead-in outright if the bullet list is
            empty, so no dangling heading is left. No-ops for units with no
            highlights section in this template."""
            key = (plant, "q1")
            slots_idx = HIGHLIGHTS_SLOTS.get(key)
            if not slots_idx:
                return
            fix = LEADIN_TEXT_FIX.get(key)
            if fix:
                idx, text = fix
                bdu.set_para(para(idx), text)
            for blank_idx in HIGHLIGHTS_BLANKS_REMOVE.get(key, []):
                bdu.remove_paragraph(para(blank_idx))
            bdu.fill_variable_bullets([para(i) for i in slots_idx], bullets)
            if not bullets:
                heading_idx = HIGHLIGHTS_HEADING.get(key)
                leadin_idx = HIGHLIGHTS_LEADIN.get(key)
                if heading_idx is not None:
                    bdu.remove_paragraph(para(heading_idx))
                if leadin_idx is not None:
                    bdu.remove_paragraph(para(leadin_idx))

        # ---- tables ----------------------------------------------------
        bq2.set_abp_header(tables[TABLE_SAIL], fy_start)
        bq2.set_period_header(tables[TABLE_SAIL], q1_label, q1_cply_label)
        bq2.fill_period_table(cur, tables[TABLE_SAIL], "SAIL", SAIL_ITEMS, Q1_CUR, Q1_CPLY, fy_start)

        for plant in PLANT_ORDER:
            cfg = PLANT_CONFIG[plant]
            bq2.set_abp_header(tables[cfg["table"]], fy_start)
            bq2.set_period_header(tables[cfg["table"]], q1_label, q1_cply_label)
            bq2.fill_period_table(cur, tables[cfg["table"]], plant, cfg["item_rows"], Q1_CUR, Q1_CPLY, fy_start)

        # ---- paragraph text -------------------------------------------
        bdu.set_para(para(P_TITLE), _title_text(fy_start))
        bdu.set_para(para(P_INTRO), f"SAIL and Plant-wise Production performance during {q1_label} w.r.t. ABP")

        # SAIL heading, label, summary, highlights
        bdu.set_para(para(P_SAIL_HEADING), f"SAIL: Production performance during {q1_label} ")
        bdu.set_para(para(P_SAIL_LABEL), f"{q1_label}:")
        bdu.set_para(para(P_SAIL_SUMMARY),
                     bq2.plant_summary_big(cur, "SAIL", Q1_CUR, Q1_CPLY, fy_start).replace("{P}", q1_label))
        sail_bullets = bnc.best_ever_bullets(cur, "SAIL", lambda f: bnc.quarter_months(f, 1), fy_start, "q1")
        sail_bullets += bnm.additional_highlight_lines(manual, "SAIL")
        apply_highlights("SAIL", sail_bullets)

        bdu.set_para(para(P_PLANT_SECTION_TITLE),
                     f"Plant wise: Production performance during {q1_label} w.r.t. ABP")

        # per-plant label/summary, highlights, why narrative
        for plant in PLANT_ORDER:
            cfg = PLANT_CONFIG[plant]
            summary_fn = bq2._SUMMARY_FN[cfg["kind"]]

            bdu.set_para(para(cfg["label"]), f"{q1_label}:")
            bdu.set_para(para(cfg["summary"]),
                         summary_fn(cur, plant, Q1_CUR, Q1_CPLY, fy_start).replace("{P}", q1_label))

            bullets = bnc.best_ever_bullets(cur, plant, lambda f: bnc.quarter_months(f, 1), fy_start, "q1")
            bullets += bnm.additional_highlight_lines(manual, plant)
            apply_highlights(plant, bullets)

            why_heading_idx = WHY_HEADING[plant]
            bdu.set_para(para(why_heading_idx), f"Production Performance during {q1_label}:")
            placeholder = f"[Add production narrative for {plant} here]"
            lines = bnm.why_narrative_lines(manual, plant, placeholder)
            bdu.fill_variable_bullets([para(i) for i in WHY_SLOTS[plant]], lines)

        # ---- techno section --------------------------------------------
        periods = [{"label": "Q1cur", "months": Q1_CUR}, {"label": "Q1cply", "months": Q1_CPLY}]
        techno_res = tp.build_period_report(TECHNO_PLANTS, TECHNO_PARAMS, periods)
        techno_by_param = {s["parameter"]: s for s in techno_res["sections"]}

        report_month = Q1_CUR[-1]
        targets = {}
        for plant in ["BSP", "DSP", "RSP", "BSL", "ISP"]:
            raw = pt._get_plant_techno_plan_targets(plant, report_month)
            targets[plant] = {k[1]: v for k, v in raw.items()}
        sail_targets_raw = pt.compute_sail_targets(fy)
        targets["SAIL"] = {k[1]: v for k, v in sail_targets_raw.items()}

        t9 = tables[TABLE_TECHNO]
        _set_techno_header(t9, fy_start, q1_label, q1_cply_label)

        ri = 2
        for param in TECHNO_PARAMS:
            sec = techno_by_param[param]
            row_by_plant = {r["plant"]: r["values"] for r in sec["rows"]}
            for plant in TECHNO_PLANTS:
                cells = t9.rows[ri].cells
                vals = row_by_plant.get(plant, {})

                def disp(label):
                    d = vals.get(label)
                    return d["display"] if d else ""

                tgt = _fmt_techno_target(param, (targets.get(plant) or {}).get(param))
                bdu.set_cell(cells[2], tgt if tgt is not None else "")
                bdu.set_cell(cells[3], disp("Q1cur"))
                bdu.set_cell(cells[4], disp("Q1cply"))
                ri += 1

        bdu.set_para(
            para(P_TECHNO_SECTION_TITLE),
            f"SAIL and Plant-wise Major Techno-economic parameters performance during {q1_label}:"
        )

        # ---- annexure index (table 10) -----------------------------------
        t10 = tables[TABLE_ANNEXURE_INDEX]
        m1, m2, m3 = (bq2._month_name_title(m) for m in Q1_CUR)
        bdu.set_cell(t10.rows[0].cells[3],
                     f"SAIL: Production performance during {m1}’{yy}, {m2}’{yy} and {m3}’{yy}.")
        bdu.set_cell(t10.rows[1].cells[3], f"SAIL: Production performance during {q1_label}")
        bdu.set_cell(t10.rows[2].cells[3],
                     f"SAIL: Techno-economic parameters during {q1_label}, best ever achieved, norms and benchmarks")
        bdu.set_cell(t10.rows[3].cells[3],
                     f"SAIL: Production performance :New Facilities during \n{q1_label}")
        bdu.set_cell(t10.rows[4].cells[3],
                     f"SAIL: Techno-economic parameters for New Blast Furnaces during {q1_label}")
        bdu.set_cell(t10.rows[5].cells[3], f"SAIL: Product mix performance during {q1_label}")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    finally:
        conn.close()
